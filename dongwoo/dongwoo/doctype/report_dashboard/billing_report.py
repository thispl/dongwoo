import frappe
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from six import BytesIO
from datetime import datetime, timedelta
import calendar
from frappe.utils import getdate
from datetime import datetime

@frappe.whitelist()
def download():
    filename = 'Billing Statement'
    test = build_xlsx_response(filename)

def format_currency(value):
    if value is None:
        return "0"
    
    number_str = str(int(value))
    
    if len(number_str) > 3:
        last_three = number_str[-3:]  
        other_digits = number_str[:-3] 
        
        formatted_other = []
        while len(other_digits) > 2:
            formatted_other.append(other_digits[-2:])
            other_digits = other_digits[:-2]
        if other_digits:
            formatted_other.append(other_digits)        
        formatted_other.reverse()
        formatted_number = ','.join(formatted_other) + ',' + last_three
    else:
        formatted_number = number_str
    
    return formatted_number


    

def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    # frappe.log_error("Download Payroll Report Args", args)
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    from datetime import datetime
    att_date = datetime.strptime(args.start_date, "%Y-%m-%d")
    att_month = att_date.strftime("%b")  
    att_year = att_date.strftime("%y")   
    # Execute the SQL query to fetch employee details along with corresponding salary slip names
    if args.contractor:
        employee_details = frappe.db.sql("""
            SELECT e.bank_ac_no AS 'AC No', e.employee, e.first_name AS 'Employee Name', s.employee_name AS 'Salary Slip Ename', e.mop, e.employee_type AS 'Employee Type',e.revised_gross AS 'Revised',
            e.designation AS 'Designation', e.employee_type AS 'Etype', e.department AS 'Department', e.uan_number AS "UAN NO", e.esi_number AS "Esi", e.date_of_joining AS 'Date of Joining',
            s.payment_days AS 'Payment Days',s.present_days_ AS 'Pre',s.week_off_count_ AS 'WOFF',s.leave_without_pay AS 'LOP',s.absent_days AS 'ABS', s.total_working_days AS 'Days in Month',
            e.basic AS 'Basic Pay', e.house_rent_allowance AS 'HRA', e.medical_allowance AS 'Medical Allowance', e.conveyance_allowance AS 'Conveyance',
            e.education_allowance AS 'Education Allowance', e.leave_and_travel_allowance AS 'LTA', e.dress_allowance AS 'Dress Allowance', e.gross_pay AS 'Fixed Gross', 
            s.name AS 'Name', s.total_deduction AS 'Total Deduction',e.basic_da AS 'Basic',e.hra_contractor_ AS 'hra'
        FROM `tabEmployee` e
        INNER JOIN `tabSalary Slip` s ON e.name = s.employee AND s.employee_type='Contract Employee'
        INNER JOIN `tabDepartment` d ON e.department = d.name
        INNER JOIN `tabDesignation` de ON de.name = e.designation
        WHERE e.name = s.employee and s.start_date <= %s AND s.end_date >= %s AND e.contractor = %s
        ORDER BY d.order_value, e.employee_type, de.order, e.date_of_joining
    """, (args.start_date, args.end_date,args.contractor), as_dict=True)
    else:
        employee_details = frappe.db.sql("""
        SELECT e.bank_ac_no AS 'AC No', e.employee, e.first_name AS 'Employee Name', s.employee_name AS 'Salary Slip Ename', e.mop, e.employee_type AS 'Employee Type',e.revised_gross AS 'Revised',
            e.designation AS 'Designation', e.employee_type AS 'Etype', e.department AS 'Department', e.uan_number AS "UAN NO", e.esi_number AS "Esi", e.date_of_joining AS 'Date of Joining',
            s.payment_days AS 'Payment Days',s.present_days_ AS 'Pre',s.week_off_count_ AS 'WOFF',s.leave_without_pay AS 'LOP',s.absent_days AS 'ABS', s.total_working_days AS 'Days in Month',
            e.basic AS 'Basic Pay', e.house_rent_allowance AS 'HRA', e.medical_allowance AS 'Medical Allowance', e.conveyance_allowance AS 'Conveyance',
            e.education_allowance AS 'Education Allowance', e.leave_and_travel_allowance AS 'LTA', e.dress_allowance AS 'Dress Allowance', e.gross_pay AS 'Fixed Gross', 
            s.name AS 'Name', s.total_deduction AS 'Total Deduction',e.basic_da AS 'Basic',e.hra_contractor_ AS 'hra'
        FROM `tabEmployee` e
        INNER JOIN `tabSalary Slip` s ON e.name = s.employee AND s.employee_type='Contract Employee'
        INNER JOIN `tabDepartment` d ON e.department = d.name
        INNER JOIN `tabDesignation` de ON de.name = e.designation
        WHERE e.name = s.employee and s.start_date <= %s AND s.end_date >= %s
        ORDER BY d.order_value, e.employee_type, de.order, e.date_of_joining
    """, (args.start_date, args.end_date), as_dict=True)
    # frappe.log_error("Download Payroll Report Args", employee_details)
    from datetime import datetime, timedelta
    import calendar
    start_date = datetime.strptime(args.start_date, '%Y-%m-%d')
    curr_month = start_date.strftime('%b') + ' - ' + str(start_date.strftime('%y'))
    first_day_current_month = start_date.replace(day=1)
    last_day_current_month = (first_day_current_month.replace(month=start_date.month % 12 + 1, day=1) - timedelta(days=1))
    prev_month_date = first_day_current_month - timedelta(days=1)
    first_day_prev_month = prev_month_date.replace(day=1)
    last_day_prev_month = (first_day_prev_month.replace(month=prev_month_date.month % 12 + 1, day=1) - timedelta(days=1))
    prev_prev_month_date = first_day_prev_month - timedelta(days=1)
    first_day_prev_prev_month = prev_prev_month_date.replace(day=1)
    last_day_prev_prev_month = (first_day_prev_prev_month.replace(month=prev_prev_month_date.month % 12 + 1, day=1) - timedelta(days=1))
    prev_month = prev_month_date.strftime('%b') + ' - ' + str(prev_month_date.strftime('%y'))
    prev_prev_month = prev_prev_month_date.strftime('%b') + ' - ' + str(prev_prev_month_date.strftime('%y'))
    prev_start = first_day_prev_month.strftime('%Y-%m-%d')
    prev_end = last_day_prev_month.strftime('%Y-%m-%d')
    prev_prev_start = first_day_prev_prev_month.strftime('%Y-%m-%d')
    prev_prev_end = last_day_prev_prev_month.strftime('%Y-%m-%d')

    ws = wb.create_sheet(sheet_name, 0)
    ws.append([''])
    ws.append([''])
    ws.append(['','HT WAGE SUMMARY -' ''  +str(att_month)+ '-' +str(att_year)])
    ws.append(['','','','','','',prev_prev_month,prev_month,curr_month])
    tot_ma_amt=0
    prev_tot_ma_amt=0
    prev_prev_tot_ma_amt=0
    tot1=0
    tot2=0
    tot3=0
    fest1=0
    fest2=0
    fest3=0
    amnt=0
    esiot=0
    ptotesi=0
    pptotesi=0
    totesi=0
    att_bon=0
    tot_fest1=0
    tot_fest2=0
    tot_fest3=0
    overall_esi=0
    prev_tot_esi=0
    prev_prev_tot_esi=0
    overall_ot=0
    prev_tot_ot=0
    prev_prev_tot_ot=0
    overall_attb=0
    prev_tot_attb=0
    prev_prev_tot_attb=0
    overall_fall=0
    prev_tot_fall=0
    prev_prev_tot_fall=0
    fmd_prev_prev_tot_fall=0
    fmd_prev_tot_fall=0
    fmd_overall_fall=0
    fmd_overall_attb=0
    fmd_prev_tot_attb=0
    fmd_prev_prev_tot_attb=0
    fmd_prev_prev_tot_ot=0
    fmd_overall_ot=0
    fmd_prev_tot_ot=0
    fmd_prev_prev_tot_esi=0
    fmd_prev_tot_esi=0
    fmd_overall_esi=0
    fmd_tot_fest1=0
    fmd_tot_fest2=0
    fmd_tot_fest3=0
    fmd_tot1=0
    fmd_tot2=0
    fmd_tot3=0
    fmd_tot_ma_amt=0
    fmd_prev_tot_ma_amt=0
    fmd_prev_prev_tot_ma_amt=0 
    tot_pp_hpmt=0
    tot_p_hpmt=0
    tot_c_hpmt=0
    for row in employee_details: 
        # For row 1 in table 1
        wage = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Wage"},"amount") or 0.0
        actpf = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Actual PF"},"amount") or 0.0
        overall_tot=float(wage)+float(actpf)
        
        pprevslip = frappe.db.get_value("Salary Slip", {"employee": row['employee'] ,"start_date":prev_start,"end_date":prev_end},["name"]) 
        prev_wage=frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Wage"},"amount") or 0.0
        prev_actpf = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Actual PF"},"amount") or 0.0
        prev_overall_tot=float(prev_wage)+float(prev_actpf)
        
        prev_prevslip = frappe.db.get_value("Salary Slip", {"employee": row['employee'] ,"start_date":prev_prev_start,"end_date":prev_prev_end},["name"]) 
        prev_prev_wage=frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Wage"},"amount") or 0.0
        prev_prev_actpf = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Actual PF"},"amount") or 0.0
        Prev_prev_overall_tot=float(prev_prev_wage)+float(prev_prev_actpf)
        
        if row['Department']=="FMD":
            fmd_tot_ma_amt+= overall_tot
            fmd_prev_tot_ma_amt+= prev_overall_tot
            fmd_prev_prev_tot_ma_amt+= Prev_prev_overall_tot
        else:
            prev_prev_tot_ma_amt+= Prev_prev_overall_tot
            prev_tot_ma_amt+= prev_overall_tot
            tot_ma_amt+= overall_tot
        # For row 2 in table 1
        revised=((row['Revised']*row['Pre'])*0.6*0.13)
        if revised < 1170:
            revised=revised
        else:
            revised=1170
        curr=(wage+revised)*0.1

        if pprevslip:
            prev=(prev_wage+revised)*0.1
        else:
            prev=0
        prev_prev=0
        if prev_prevslip:
            prev_prev=(prev_prev_wage+revised)*0.1
        else:
            prev_prev=0
        fest = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Festival Allowance"},["amount"]) or 0.0
        fest=float(fest)
        prev_fest = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Festival Allowance"},["amount"]) or 0.0
        prev_fest=float(prev_fest)
        prev_prev_fest = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Festival Allowance"},["amount"]) or 0.0
        prev_prev_fest=float(prev_prev_fest)
        if fest>0:
            fest=0
        else:
            fest=row['Revised']
        if pprevslip:
            if prev_fest>0:
                prev_fest=0
            else:
                prev_fest=row['Revised']
        else:
            prev_fest=0
        if prev_prevslip:
            if prev_prev_fest>0:
                prev_prev_fest=0
            else:
                prev_prev_fest=row['Revised']
        else:
            prev_prev_fest=0
        tot_1=curr-(fest)*0.1
        if tot_1 < 0:
            tot_1=0
        tot_2=prev-(prev_fest)*0.1
        if tot_2 < 0:
            tot_2=0
        tot_3=prev_prev-(prev_prev_fest)*0.1
        if tot_3 < 0:
            tot_3=0
        
        if row['Department']=="FMD":
            fmd_tot1+=tot_1
            fmd_tot2+=tot_2
            fmd_tot3+=tot_3
        else:
            tot1+=tot_1
            tot2+=tot_2
            tot3+=tot_3
        # row 3 for table 1
        curr_fall = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Festival Allowance"},"amount") or 0.0
        curr_fall=float(curr_fall)
        previous_fall = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Festival Allowance"},["amount"]) or 0.0
        previous_fall=float(previous_fall)
        prev_previous_fall = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Festival Allowance"},["amount"]) or 0.0
        prev_previous_fall=float(prev_previous_fall)
        if curr_fall>0:
            fest1=(row['Revised'])*0.1
        else:
            fest1=0
        if previous_fall>0:
            if pprevslip:
                fest2=(row['Revised'])*0.1
                # if fest2<0:
                #     fest2=0
            else:
                fest2=0
        else:
            fest2=0
        if prev_previous_fall>0:
            if prev_prevslip:
                fest3=(row['Revised'])*0.1
                # if fest3 < 0 :
                #     fest3=0
            else:
                fest3=0
        else:
            fest3=0
        
        if row['Department']=="FMD":
            fmd_tot_fest1+=fest1
            fmd_tot_fest2+=fest2
            fmd_tot_fest3+=fest3
        else:
            tot_fest1+=fest1
            tot_fest2+=fest2
            tot_fest3+=fest3
        # row 4 for table 1
        esi = row['Revised']*(row['Days in Month']-int(row['WOFF']))
        if esi < 21000:
            if esi<21000:
                amnt = (float(row['Basic']) + float(row['hra'])) * float(row['Pre'])
                esiot=((float(row['Basic'])+float(row['hra']))%8)*float(row['Pre'])
                att_bon = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Attendance Bonus"},"amount") or 0.0
                att_bon=float(att_bon)
                totesi=(amnt+esiot+att_bon)*0.0325
            else:
                totesi=0
        else:
            totesi=0
        
        if row['Department']=="FMD":
            fmd_overall_esi+=totesi
        else:
            overall_esi+=totesi
        if pprevslip:
            pre=frappe.db.get_value('Salary Slip',{'name':pprevslip},['present_days_']) or 0.0
            pre=float(pre)
            week=frappe.db.get_value('Salary Slip',{'name':pprevslip},['week_off_count_']) or 0.
            week=float(week)
            pay=frappe.db.get_value('Salary Slip',{'name':pprevslip},['payment_days'])  or 0.0
            pay=float(pay)
            prevesi = row['Revised']*(pay-float(week))
            if prevesi < 21000:
                if prevesi<21000:
                    amnt=(float(row['Basic'])+float(row['hra']))*float(pre)
                    esiot=((float(row['Basic'])+float(row['hra']))%8)*float(pre)
                    patt_bon = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Attendance Bonus"},"amount") or 0.0
                    patt_bon=float(patt_bon)
                    ptotesi=(amnt+esiot+patt_bon)*0.0325
                else:
                    ptotesi=0
            else:
                ptotesi=0
        else:
            ptotesi=0
        
        if row['Department']=="FMD":
            fmd_prev_tot_esi+= ptotesi
        else:
            prev_tot_esi+= ptotesi
        if prev_prevslip:
            ppre=frappe.db.get_value('Salary Slip',{'name':prev_prevslip},['present_days_']) or 0.0
            ppre=float(ppre)
            pweek=frappe.db.get_value('Salary Slip',{'name':prev_prevslip},['week_off_count_']) or 0.0
            pweek=float(pweek)
            ppay=frappe.db.get_value('Salary Slip',{'name':prev_prevslip},['payment_days']) or 0.0
            ppay=float(ppay)
            pprev_esi = row['Revised']*(ppay-pweek)
            if pprev_esi < 21000:
                if pprev_esi<21000:
                    amnt=(float(row['Basic'])+float(row['hra']))*float(ppre)
                    esiot=((float(row['Basic'])+float(row['hra']))%8)*float(ppre)
                    patt_bon = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Attendance Bonus"},"amount") or 0.0
                    patt_bon=float(patt_bon)
                    pptotesi=(amnt+esiot+patt_bon)*0.0325
                else:
                    pptotesi=0
            else:
                pptotesi=0
        else:
            pptotesi=0
        
        if row['Department']=="FMD":
            fmd_prev_prev_tot_esi+= pptotesi
        else:
            prev_prev_tot_esi+= pptotesi
        # row 5 for table 1
        ot = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Overtime"},"amount") or 0.0
        ot=float(ot)
        if row['Department']=="FMD":
            fmd_overall_ot+=ot
        else:
            overall_ot+=ot
        prev_ot = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Overtime"},"amount") or 0.0
        prev_ot=float(prev_ot)
        if row['Department']=="FMD":
            fmd_prev_tot_ot+= prev_ot
        else:
            prev_tot_ot+= prev_ot
        prev_prev_ot = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Overtime"},"amount") or 0.0
        prev_prev_ot=float(prev_prev_ot)
        if row['Department']=="FMD":
            fmd_prev_prev_tot_ot+= prev_prev_ot
        else:
            prev_prev_tot_ot+= prev_prev_ot
        # row 6 for table 1
        attb = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Attendance Bonus"},"amount") or 0.0
        attb=float(attb)
        if row['Department']=="FMD":
            fmd_overall_attb+=attb
        else:
            overall_attb+=attb
        prev_attb = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Attendance Bonus"},"amount") or 0.0
        prev_attb=float(prev_attb)
        if row['Department']=="FMD":
            fmd_prev_tot_attb+= prev_attb
        else:
            prev_tot_attb+= prev_attb
        prev_prev_attb = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Attendance Bonus"},"amount") or 0.0
        prev_prev_attb=float(prev_prev_attb)
        if row['Department']=="FMD":
            fmd_prev_prev_tot_attb+= prev_prev_attb
        else:
            prev_prev_tot_attb+= prev_prev_attb
        # row 7 for table 1
        fall = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Festival Allowance"},"amount") or 0.0
        fall=float(fall)
        if row['Department']=="FMD":
            fmd_overall_fall+=fall
        else:
            overall_fall+=fall
        prev_fall = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Festival Allowance"},"amount") or 0.0
        prev_fall=float(prev_fall)
        if row['Department']=="FMD":
            fmd_prev_tot_fall+= prev_fall
        else:
            prev_tot_fall+= prev_fall
        prev_prev_fall = frappe.db.get_value("Salary Detail", {"parent": prev_prevslip , "salary_component": "Festival Allowance"},"amount") or 0.0
        prev_prev_fall=float(prev_prev_fall)
        if row['Department']=="FMD":
            fmd_prev_prev_tot_fall+= prev_prev_fall
        else:
            prev_prev_tot_fall+= prev_prev_fall
    ws.append(['', 'TOTAL MANDAYS AMOUNT', '', '', '', '', format_currency(round(prev_prev_tot_ma_amt)), format_currency(round(prev_tot_ma_amt)), format_currency(round(tot_ma_amt))])
    ws.append(['', 'REGULAR SERVICE CHARGE @10%(+)', '', '', '', '', format_currency(round(tot3)), format_currency(round(tot2)), format_currency(round(tot1))])
    ws.append(['', 'FESTIVAL SERVICE CHARGE @10%(+)', '', '', '', '', format_currency(round(tot_fest3)), format_currency(round(tot_fest2)), format_currency(round(tot_fest1))])
    ws.append(['', 'ESI EMPLOYER CONTRIBUTION', '', '', '', '', format_currency(round(prev_prev_tot_esi)), format_currency(round(prev_tot_esi)), format_currency(round(overall_esi))])
    ws.append(['', 'TOTAL OT AMOUNT', '', '', '', '', format_currency(round(prev_prev_tot_ot)), format_currency(round(prev_tot_ot)), format_currency(round(overall_ot))])
    ws.append(['', 'TOTAL ATTENDANCE BONUS', '', '', '', '', format_currency(round(prev_prev_tot_attb)), format_currency(round(prev_tot_attb)), format_currency(round(overall_attb))])
    ws.append(['', 'TOTAL FESTIVAL ALLOWANCE', '', '', '', '', format_currency(round(prev_prev_tot_fall)), format_currency(round(prev_tot_fall)), format_currency(round(overall_fall))])


    tot_curr=tot_ma_amt+tot_fest1+tot1+overall_esi+overall_ot+overall_attb+overall_fall
    tot_p=prev_tot_ma_amt+tot2+tot_fest2+prev_tot_esi+prev_tot_ot+prev_tot_attb+prev_tot_fall
    tot_pp=prev_prev_tot_ma_amt+tot3+tot_fest3+prev_prev_tot_esi+prev_prev_tot_ot+prev_prev_tot_attb+prev_prev_tot_fall
    fmd_tot_1=fmd_prev_prev_tot_ma_amt+fmd_tot3+fmd_tot_fest3+fmd_prev_prev_tot_esi+fmd_prev_prev_tot_ot+fmd_prev_prev_tot_attb+fmd_prev_prev_tot_fall
    fmd_tot_2=fmd_prev_tot_ma_amt+fmd_tot2+fmd_tot_fest2+fmd_prev_tot_esi+fmd_prev_tot_ot+fmd_prev_tot_attb+fmd_prev_tot_fall
    fmd_tot_3=fmd_tot_ma_amt+fmd_tot1+fmd_tot_fest1+fmd_overall_esi+fmd_overall_ot+fmd_overall_attb+fmd_overall_fall
    ws.append(['', 'TOTAL', '', '', '', '', format_currency(round(tot_pp)), format_currency(round(tot_p)), format_currency(round(tot_curr))])
    ws.append(['', 'Less: OT on HJ Scope (0*50%)', '', '', '', '', '-', '-', '-'])
    ws.append(['', 'Add: Wrongly debited last month', '', '', '', '', '-', '-', '-'])
    ws.append(['', 'TOTAL', '', '', '', '', format_currency(round(tot_pp)), format_currency(round(tot_p)), format_currency(round(tot_curr))])

    ws.append([''])
    ws.append([''])
    ws.append(['','FMD WAGE SUMMARY -' ''  +str(att_month)+ '-' +str(att_year)])
    ws.append(['','','','','','',prev_prev_month,prev_month,curr_month])
    ws.append(['', 'TOTAL MANDAYS AMOUNT', '', '', '', '', format_currency(round(fmd_prev_prev_tot_ma_amt)), format_currency(round(fmd_prev_tot_ma_amt)), format_currency(round(fmd_tot_ma_amt))])
    ws.append(['', 'REGULAR SERVICE CHARGE @10%(+)', '', '', '', '', format_currency(round(fmd_tot3)), format_currency(round(fmd_tot2)), format_currency(round(fmd_tot1))])
    ws.append(['', 'FESTIVAL SERVICE CHARGE @10%(+)', '', '', '', '', format_currency(round(fmd_tot_fest3)), format_currency(round(fmd_tot_fest2)), format_currency(round(fmd_tot_fest1))])
    ws.append(['', 'ESI EMPLOYER CONTRIBUTION', '', '', '', '', format_currency(round(fmd_prev_prev_tot_esi)), format_currency(round(fmd_prev_tot_esi)), format_currency(round(fmd_overall_esi))])
    ws.append(['', 'TOTAL OT AMOUNT', '', '', '', '', format_currency(round(fmd_prev_prev_tot_ot)), format_currency(round(fmd_prev_tot_ot)), format_currency(round(fmd_overall_ot))])
    ws.append(['', 'TOTAL ATTENDANCE BONUS', '', '', '', '', format_currency(round(fmd_prev_prev_tot_attb)), format_currency(round(fmd_prev_tot_attb)), format_currency(round(fmd_overall_attb))])
    ws.append(['', 'TOTAL FESTIVAL ALLOWANCE', '', '', '', '', format_currency(round(fmd_prev_prev_tot_fall)), format_currency(round(fmd_prev_tot_fall)), format_currency(round(fmd_overall_fall))])

    ws.append(['', 'TOTAL', '', '', '', '', format_currency(round(fmd_tot_1)), format_currency(round(fmd_tot_2)), format_currency(round(fmd_tot_3))])
    ws.append(['', 'Less: if any', '', '', '', '', '-', '-', '-'])
    ws.append(['', 'TOTAL', '', '', '', '', format_currency(round(fmd_tot_1)), format_currency(round(fmd_tot_2)), format_currency(round(fmd_tot_3))])

    ws.append([''])
    ws.append([''])
    heat_diff=tot_curr-tot_p
    fmd_diff=fmd_tot_3-fmd_tot_2
    pp_gst18=(tot_pp+fmd_tot_1)*0.18
    p_gst18=(fmd_tot_2+tot_p)*0.18
    curr_gst18=(fmd_tot_3+tot_curr)*0.18
    heat_fmd_gst=(heat_diff+fmd_diff)*0.18
    ws.append(['', 'PARTICULARS', '', '', '', prev_prev_month, prev_month, curr_month, 'DIFFERENCE'])
    ws.append(['', 'HEAT TREATMENT', '', '', '', format_currency(round(tot_pp)), format_currency(round(tot_p)), format_currency(round(tot_curr)), format_currency(round(heat_diff))])
    ws.append(['', 'FMD', '', '', '', format_currency(round(fmd_tot_1)), format_currency(round(fmd_tot_2)), format_currency(round(fmd_tot_3)), format_currency(round(fmd_diff))])
    ws.append(['', 'GRAND TOTAL', '', '', '', format_currency(round(tot_pp + fmd_tot_1)), format_currency(round(fmd_tot_2 + tot_p)), format_currency(round(fmd_tot_3 + tot_curr)), format_currency(round(heat_diff + fmd_diff))])
    ws.append(['', 'GST - 18%', '', '', '', format_currency(round(pp_gst18)), format_currency(round(p_gst18)), format_currency(round(curr_gst18)), format_currency(round(heat_fmd_gst))])
    ws.append(['', 'GRAND TOTAL', '', '', '', format_currency(round(pp_gst18 + tot_pp + fmd_tot_1)), format_currency(round(p_gst18 + fmd_tot_2 + tot_p)), format_currency(round(curr_gst18 + fmd_tot_3 + tot_curr)), format_currency(round(heat_fmd_gst + heat_diff + fmd_diff))])

    ws.append([''])
    ws.append([''])
    ws.append(['','MANPOWER DETAILS','','','','',prev_prev_month,prev_month,curr_month]) 
    dept=['Production','D.PMT','FMD','QC','H.PMT','HK & Garden','PE','ETP','Maintenance']  
    for d in dept:
        if args.contractor:
            c_hpmt = frappe.db.count("Salary Slip", {"department": d ,"start_date":args.start_date,"end_date":args.end_date,'employee_type':'Contract Employee','contractor':args.contractor}) 
            p_hpmt = frappe.db.count("Salary Slip", {"department": d ,"start_date":prev_start,"end_date":prev_end,'employee_type':'Contract Employee','contractor':args.contractor}) 
            pp_hpmt = frappe.db.count("Salary Slip", {"department": d ,"start_date":prev_prev_start,"end_date":prev_prev_end,'employee_type':'Contract Employee','contractor':args.contractor})
        else:
            c_hpmt = frappe.db.count("Salary Slip", {"department": d ,"start_date":args.start_date,"end_date":args.end_date,'employee_type':'Contract Employee'}) 
            p_hpmt = frappe.db.count("Salary Slip", {"department": d ,"start_date":prev_start,"end_date":prev_end,'employee_type':'Contract Employee'}) 
            pp_hpmt = frappe.db.count("Salary Slip", {"department": d ,"start_date":prev_prev_start,"end_date":prev_prev_end,'employee_type':'Contract Employee'}) 
        ws.append(['',d,'','','','',pp_hpmt,p_hpmt,c_hpmt])
        tot_pp_hpmt+=pp_hpmt
        tot_p_hpmt+=p_hpmt
        tot_c_hpmt+=c_hpmt
    ws.append(['','Total','','','','',tot_pp_hpmt,tot_p_hpmt,tot_c_hpmt])
    for i in range(1,51):
        ws.row_dimensions[i].height=20  
    widths=['F','G','H','I']
    for i in widths:
        ws.column_dimensions[i].width=11  
    widths=['D','E']
    for i in widths:
        ws.column_dimensions[i].width=4   
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=9)
    ws.merge_cells(start_row=40, start_column=2, end_row=40, end_column=6)
    ws.merge_cells(start_row=18, start_column=2, end_row=18, end_column=9)
    for i in range(4,16):
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    for i in range(19,30):
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    for i in range(32,38):
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
    for i in range(41,51):	
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    align_center = Alignment(horizontal='center',vertical='center')
    border = Border(left=Side(border_style='thin', color='000000'),
                 right=Side(border_style='thin', color='000000'),
                 top=Side(border_style='thin', color='000000'),
                 bottom=Side(border_style='thin', color='000000'))
    cent_align=[3,4,18,19,32,40,50]
    for n in cent_align:
        for rows in ws.iter_rows(min_row=n, max_row=n, min_col=2, max_col=9):
            for cell in rows:
                cell.alignment = align_center
    for rows in ws.iter_rows(min_row=3, max_row=15, min_col=2, max_col=9):
        for cell in rows:
            cell.border = border
    for rows in ws.iter_rows(min_row=18, max_row=29, min_col=2, max_col=9):
        for cell in rows:
            cell.border = border
    for rows in ws.iter_rows(min_row=32, max_row=37, min_col=2, max_col=9):
        for cell in rows:
            cell.border = border
    for rows in ws.iter_rows(min_row=40, max_row=50, min_col=2, max_col=9):
        for cell in rows:
            cell.border = border
    green_fill = PatternFill(start_color='bde394', end_color='bde394', fill_type='solid')
    for row in ws.iter_rows(min_row=5, max_row=15, min_col=9, max_col=9):
        for cell in row:
            cell.fill = green_fill
    for row in ws.iter_rows(min_row=20, max_row=29, min_col=9, max_col=9):
        for cell in row:
            cell.fill = green_fill
    for row in ws.iter_rows(min_row=33, max_row=37, min_col=9, max_col=9):
        for cell in row:
            cell.fill = green_fill
    yellow_fill = PatternFill(start_color='eff57a', end_color='eff57a', fill_type='solid')
    for row in ws.iter_rows(min_row=6, max_row=7, min_col=2, max_col=8):
        for cell in row:
            cell.fill = yellow_fill
    for row in ws.iter_rows(min_row=13, max_row=14, min_col=2, max_col=9):
        for cell in row:
            cell.fill = yellow_fill
    for row in ws.iter_rows(min_row=21, max_row=22, min_col=2, max_col=8):
        for cell in row:
            cell.fill = yellow_fill
    for row in ws.iter_rows(min_row=28, max_row=28, min_col=2, max_col=9):
        for cell in row:
            cell.fill = yellow_fill
   
    bold_lines=[3,4,12,13,14,15,18,19,27,28,29,32,35,37,40,50]
    for i in bold_lines:
        for row in ws.iter_rows(min_row=i, max_row=i, min_col=2, max_col=9):
            for cell in row:
                cell.font = Font(bold=True)
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'
    
