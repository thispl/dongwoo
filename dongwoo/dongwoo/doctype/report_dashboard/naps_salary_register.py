import frappe
from frappe.utils.background_jobs import enqueue
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment,PatternFill
from six import BytesIO
from frappe.utils import getdate, now_datetime
from openpyxl.utils import get_column_letter
from datetime import timedelta
from datetime import datetime
from openpyxl.styles import Border, Side
@frappe.whitelist()
def download(start_date=None, end_date=None):
    if start_date:
        start_date = getdate(start_date)
    if end_date:
        end_date = getdate(end_date)

    month_year = start_date.strftime("%B %Y")  
    month_year_text = start_date.strftime("%B - %Y").upper()
    filename = f"NAPS Salary Statement - {month_year}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "NAPS Salary Statement" 

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )


    header_text = "DONGWOO SURFACETECH (INDIA) PVT LTD"
    ws.merge_cells("B2:E2")   
    ws["B2"] = header_text
    ws["B2"].font = Font(bold=True, size=14,underline="single")
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")

    header_text1 = "DAILY ATTN. INFORMATION SYSTEM"
    ws.merge_cells("B3:E3")   
    ws["B3"] = header_text1
    ws["B3"].font = Font(size=12)
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center")

    header_text = "MONTH"
    ws.merge_cells("B4:C4")   
    ws["B4"] = header_text
    ws["B4"].font = Font(size=12)
    ws["B4"].alignment = Alignment(horizontal="left", vertical="center")

    header_text = month_year_text
    ws.merge_cells("D4:E4")   
    ws["D4"] = header_text
    ws["D4"].font = Font(size=12)
    ws["D4"].alignment = Alignment(horizontal="left", vertical="center")

    header_text = "NAPS TRAINEES"
    ws.merge_cells("B6:D6")
    cell = ws["B6"]
    cell.value = header_text
    cell.font = Font(size=12)
    cell.alignment = Alignment(horizontal="center", vertical="center")

    cell.fill = PatternFill(
        start_color="D8BFD8",  
        end_color="D3D3D3",    
        fill_type="solid"
    )

    start_row = 7
    headers = ["S. NO", "DEPT", "EMP NO.", "NAMES", "DOJ",""]

    col_start = 2
    holiday_fill = PatternFill(start_color="FFD580", end_color="FFD580", fill_type="solid")
    sunday_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    blue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
    pink_fill = PatternFill(start_color="FFC0CB", end_color="FFC0CB", fill_type="solid")
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
        
    for i, header in enumerate(headers):
        col_letter = get_column_letter(col_start + i)
        ws.merge_cells(f"{col_letter}{start_row}:{col_letter}{start_row+1}")
        cell = ws[f"{col_letter}{start_row}"]
        cell.value = header
        cell.font = Font(size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = yellow_fill
        cell.border = thin_border
    ws["H2"] = "A"       
    ws["H3"] = "B"       
    ws["H4"] = "C"
    ws["I2"] = "1st Shift"
    ws["I3"] = "2nd Shift"
    ws["I4"] = "3rd Shift"

    ws["K2"] = "CL"
    ws["K2"].fill = green_fill
    ws["L2"] = "Casual Leave"

    ws["K3"] = "AB"
    ws["K3"].fill = pink_fill
    ws["L3"] = "Absent"

    ws["K4"] = "H"
    ws["K4"].fill = holiday_fill
    ws["L4"] = "Holiday"

    ws["N2"] = "FH"
    ws["N2"].fill = blue_fill
    ws["O2"] = "Festival Holiday (if work, double wages)"

    ws["N3"] = "NH"
    ws["N3"].fill = blue_fill
    ws["O3"] = "National Holiday (if work, double wages)"

    current_col = col_start + len(headers)
    date_row = 7
    day_row = 8

    current_date = start_date
    while current_date <= end_date:
        col_letter = get_column_letter(current_col)

        is_sunday = current_date.weekday() == 6  

        fill_color = sunday_fill if is_sunday else yellow_fill

        date_cell = ws[f"{col_letter}{date_row}"]
        date_cell.value = current_date.day
        date_cell.font = Font(size=11)
        date_cell.alignment = Alignment(horizontal="center", vertical="center")
        date_cell.fill = fill_color
        date_cell.border = thin_border

        day_cell = ws[f"{col_letter}{day_row}"]
        day_cell.value = current_date.strftime("%a")  
        day_cell.font = Font(size=11)
        day_cell.alignment = Alignment(horizontal="center", vertical="center")
        day_cell.fill = fill_color
        day_cell.border = thin_border

        current_date += timedelta(days=1)
        current_col += 1

    extra_headers = [
        "Actual Days", "Total Present", "OT HRS", "FH / NH",
        "TOTAL LEAVE", "TOTAL ABSENT", "Holidays", "Total Paid days",
        "Fixed Gross", "Shift Allowance", "ATTN. BONUS", "FESTIVAL ALLOWANCE",
        "OT Amount", "Deductions", "EARNED GROSS", "EARNED STIPEND",
        "SERVICE CHARGE", "INSURANCE", "TOTAL AMOUNT"
    ]

    for header in extra_headers:
        col_letter = get_column_letter(current_col)
        ws.merge_cells(f"{col_letter}{date_row}:{col_letter}{day_row}") 
        cell = ws[f"{col_letter}{date_row}"]
        cell.value = header
        cell.font = Font(size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = yellow_fill
        cell.border = thin_border
        current_col += 1

    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=col_start, max_col=current_col-1):
        for cell in row:
            cell.border = thin_border
    
            
    conditions = []
    params = {}

    if start_date:
        conditions.append("ss.start_date >= %(start)s")
        params["start"] = start_date

    if end_date:
        conditions.append("ss.end_date <= %(end)s")
        params["end"] = end_date

    date_filter = " AND " + " AND ".join(conditions) if conditions else ""       
	
    params.update({
		"report_month": start_date.month,
		"report_year": start_date.year
	})
    
    employees = frappe.db.sql(f"""
		SELECT DISTINCT
			emp.name AS emp_no,
			emp.employee_name,
			emp.department,
			emp.date_of_joining AS doj,
			emp.holiday_list,
			emp.gross_pay,
			dep.order_for_naps,
			ss.custom_no_of_2_shift,
			ss.custom_no_of_3_shift
		FROM `tabSalary Slip` ss
		INNER JOIN `tabEmployee` emp ON emp.name = ss.employee
		LEFT JOIN `tabDepartment` dep ON emp.department = dep.name
		WHERE ss.start_date >= %(start)s
		AND ss.end_date <= %(end)s
		AND emp.employee_type = 'NAPS'
		ORDER BY dep.order_for_naps ASC, emp.date_of_joining ASC
	""", params, as_dict=True)

    # attendances = frappe.db.sql(f"""
    #     SELECT employee, attendance_date, status, shift , leave_type , overtime_hours ,b_allowance, c_allowance
    #     FROM `tabAttendance`
    #     WHERE attendance_date BETWEEN %(start)s AND %(end)s
    #     AND employee IN %(employees)s
    # """, {
    #     "start": start_date,
    #     "end": end_date,
    #     "employees": [emp.emp_no for emp in employees]
    # }, as_dict=True)


    employee_list = [emp.emp_no for emp in employees]
    if not employee_list:
        attendances = []
    else:
        placeholders = ', '.join(['%s'] * len(employee_list))
        attendances = frappe.db.sql(
            f"""
            SELECT employee, attendance_date, status, shift, leave_type, overtime_hours, b_allowance, c_allowance
            FROM `tabAttendance`
            WHERE attendance_date BETWEEN %s AND %s
            AND employee IN ({placeholders})
            """,
            [start_date, end_date] + employee_list,
            as_dict=True
        )

    attendance_map = {}
    for att in attendances:
        key = (att.employee, att.attendance_date)
        attendance_map[key] = att
    # salary_slip = frappe.db.get_value(
    #     "Salary Slip",
    #     {
    #         "employee": employees.emp_no,
    #         "start_date": start_date,
    #         "end_date": end_date
    #     },
    #     ["custom_no_of_2_shift", "custom_no_of_3_shift"],
    #     as_dict=True
    # )
    shift_records = frappe.db.sql("""
        SELECT employee, start_date, shift_type
        FROM `tabShift Assignment`
        WHERE start_date BETWEEN %(start)s AND %(end)s
        AND docstatus = 1
    """, {
        "start": start_date,
        "end": end_date
    }, as_dict=True)

    shift_map = {}
    for s in shift_records:
        key = (s.employee, s.start_date)
        shift_map[key] = s.shift_type

    holiday_lists = list(set(emp.holiday_list for emp in employees if emp.holiday_list))
    holiday_records = frappe.db.sql("""
        SELECT parent as holiday_list, holiday_date, festival_holiday, national_holiday
        FROM `tabHoliday`
        WHERE parent IN %(lists)s
    """, {"lists": holiday_lists}, as_dict=True)

    holiday_map = {}
    for h in holiday_records:
        if h['holiday_list'] not in holiday_map:
            holiday_map[h['holiday_list']] = {}
        holiday_map[h['holiday_list']][h['holiday_date']] = {
            "festival": h['festival_holiday'],
            "national": h['national_holiday']
        }

    dept_totals = {
        "total_days": 0,
        "present": 0,
        "ot_total": 0,
        "fh_nh": 0,
        "cl": 0,
        "ab": 0,
        "holiday": 0,
        "paid_days": 0,
        "fixed_gross": 0,
        "shift_allowance": 0,
        "att_bonus": 0,
        "festival_allowance": 0,
        "ot_amount": 0,
        "deductions": 0,
        "earned_gross": 0,
        "earned_stipend": 0,
        "service_charge": 0,
        "insurance": 0,
        "total_amount": 0,
    }

    

    def write_dept_total_row(dept_start_row, dept_end_row):
        nonlocal data_start_row
        ws.merge_cells(start_row=data_start_row, start_column=2, end_row=data_start_row, end_column=6)
        total_cell = ws.cell(row=data_start_row, column=2, value=" SUB TOTAL")
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        total_cell.font = Font(bold=True, size=12)

        for col in range(2, last_date_col + 1):
            c = ws.cell(row=data_start_row, column=col)
            c.fill = yellow_fill
            c.border = thin_border

        extra_start = last_date_col + 1
        extra_end = extra_start + 18

        totals = {header: 0 for header in dept_totals.keys()}
        for row in range(dept_start_row, dept_end_row + 1):
            totals["total_days"] += ws.cell(row=row, column=extra_start).value or 0
            totals["present"] += ws.cell(row=row, column=extra_start + 1).value or 0
            totals["ot_total"] += ws.cell(row=row, column=extra_start + 2).value or 0
            totals["fh_nh"] += ws.cell(row=row, column=extra_start + 3).value or 0
            totals["cl"] += ws.cell(row=row, column=extra_start + 4).value or 0
            totals["ab"] += ws.cell(row=row, column=extra_start + 5).value or 0
            totals["holiday"] += ws.cell(row=row, column=extra_start + 6).value or 0
            totals["paid_days"] += ws.cell(row=row, column=extra_start + 7).value or 0
            totals["fixed_gross"] += ws.cell(row=row, column=extra_start + 8).value or 0
            totals["shift_allowance"] += ws.cell(row=row, column=extra_start + 9).value or 0
            totals["att_bonus"] += ws.cell(row=row, column=extra_start + 10).value or 0
            totals["festival_allowance"] += ws.cell(row=row, column=extra_start + 11).value or 0
            totals["ot_amount"] += ws.cell(row=row, column=extra_start + 12).value or 0
            totals["deductions"] += ws.cell(row=row, column=extra_start + 13).value or 0
            totals["earned_gross"] += ws.cell(row=row, column=extra_start + 14).value or 0
            totals["earned_stipend"] += ws.cell(row=row, column=extra_start + 15).value or 0
            totals["service_charge"] += ws.cell(row=row, column=extra_start + 16).value or 0
            totals["insurance"] += ws.cell(row=row, column=extra_start + 17).value or 0
            totals["total_amount"] += ws.cell(row=row, column=extra_start + 18).value or 0

        for i, key in enumerate(totals.keys()):
            cell = ws.cell(row=data_start_row, column=extra_start + i, value=totals[key])
            cell.fill = yellow_fill
            cell.border = thin_border
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if key in ["fixed_gross", "shift_allowance", "att_bonus", "festival_allowance",
                    "ot_amount", "deductions", "earned_gross", "earned_stipend",
                    "service_charge", "insurance", "total_amount"]:
                cell.number_format = '#,##0'

        data_start_row += 1

    data_start_row = day_row + 1
    dept_start_row = data_start_row
    data_start_row = day_row + 1
    serial = 1
    current_department = None
    for emp in employees:
        if current_department and emp.department != current_department:
            dept_end_row = data_start_row - 1  
            write_dept_total_row(dept_start_row, dept_end_row)
            dept_start_row = data_start_row
        current_department = emp.department
        ws.merge_cells(start_row=data_start_row, start_column=2, end_row=data_start_row+1, end_column=2)  
        ws.merge_cells(start_row=data_start_row, start_column=3, end_row=data_start_row+1, end_column=3)  
        ws.merge_cells(start_row=data_start_row, start_column=4, end_row=data_start_row+1, end_column=4)  
        ws.merge_cells(start_row=data_start_row, start_column=5, end_row=data_start_row+1, end_column=5)  
        ws.merge_cells(start_row=data_start_row, start_column=6, end_row=data_start_row+1, end_column=6)  

        ws.cell(row=data_start_row, column=2, value=serial)
        ws.cell(row=data_start_row, column=3, value=emp.department)
        ws.cell(row=data_start_row, column=4, value=emp.emp_no)
        ws.cell(row=data_start_row, column=5, value=emp.employee_name)
        doj_value = emp.doj.strftime("%d-%b-%y") if emp.doj else ""
        ws.cell(row=data_start_row, column=6, value=doj_value)

        duty_cell = ws.cell(row=data_start_row, column=7, value="DUTY")
        ot_cell   = ws.cell(row=data_start_row+1, column=7, value="OT")

        for cell in [duty_cell, ot_cell]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(size=11)
            cell.border = thin_border  

        ws.column_dimensions[get_column_letter(7)].width = 10

        for c in range(2, 7):
            for r in range(data_start_row, data_start_row+2):
                cell = ws.cell(row=r, column=c)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c == 5:  
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    ws.column_dimensions["E"].width = 25
                if c == 3:  
                    ws.column_dimensions["C"].width = 20
                if c == 4 :  
                    ws.column_dimensions["D"].width = 15  
                if c == 6 :  
                    ws.column_dimensions["F"].width = 15       
                cell.font = Font(size=11)
                cell.border = thin_border
        last_date_col = 8 + (end_date - start_date).days
        pink_fill = PatternFill(start_color="FFC0CB", end_color="FFC0CB", fill_type="solid")
        green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
        emp_holidays = holiday_map.get(emp.holiday_list, {})
        for c in range(8, last_date_col + 1):
            current_day = start_date + timedelta(days=c - 8)
            duty_cell = ws.cell(row=data_start_row, column=c)
            ot_cell = ws.cell(row=data_start_row+1, column=c)

            att = attendance_map.get((emp.emp_no, current_day))
            shift = shift_map.get((emp.emp_no, current_day))
            holiday_flags = emp_holidays.get(current_day, {})
            b_shift_count = 0
            c_shift_count = 0
            if holiday_flags.get("festival") :
                duty_cell.value = "FH"
                duty_cell.fill = blue_fill
            elif holiday_flags.get("national"):
                duty_cell.value = "NH"
                duty_cell.fill = blue_fill    
            elif shift == "WW":
                duty_cell.value = "H"
                duty_cell.fill = holiday_fill  
            elif att:
                if att.status == "Present":
                    duty_cell.value = att.shift or ''
                elif att.status == "Half Day":
                    duty_cell.value = att.shift or ''
                elif att.status == "Absent":
                    duty_cell.value = "AB"
                    duty_cell.fill = pink_fill
                elif att.status == "On Leave" and att.leave_type == "Casual Leave (CL)":
                    duty_cell.value = "CL"
                    duty_cell.fill = green_fill  
                if att.shift == "WW":
                    duty_cell.value = "H"
                    duty_cell.fill = holiday_fill 

            else:
                duty_cell.value = ""

            
            duty_cell.alignment = Alignment(horizontal="center", vertical="center")
            duty_cell.font = Font(size=11)
            duty_cell.border = thin_border

            if att and getattr(att, "overtime_hours", 0):
                ot_cell.value = att.overtime_hours * 2
            else:
                ot_cell.value = ""
            ot_cell.alignment = Alignment(horizontal="center", vertical="center")
            ot_cell.font = Font(size=11)
            ot_cell.border = thin_border
            extra_header_start_col = 8 + (end_date - start_date).days + 1  
            extra_header_end_col = current_col - 1
            for c in range(extra_header_start_col, extra_header_end_col + 1):
                ws.merge_cells(
                    start_row=data_start_row,
                    start_column=c,
                    end_row=data_start_row + 1,
                    end_column=c
                )
            present_count = 0
            fh_nh_count = 0
            cl_count = 0
            ab_count = 0
            holiday_count = 0   
            ot_total = 0
            festival_allowance = 0
            for c in range(8, last_date_col + 1):
                duty_val = ws.cell(row=data_start_row, column=c).value
                if duty_val in ["A", "B", "C"]:
                    present_count += 1
                    if duty_val=="B":
                        b_shift_count +=1
                    if duty_val=="C":
                        c_shift_count +=1
                elif duty_val in ["FH", "NH"]:
                    fh_nh_count += 1
                    if duty_val == "FH":
                        festival_allowance += 300
                    else:  
                        festival_allowance += 200
                elif duty_val == "CL":
                    cl_count += 1
                elif duty_val == "AB":
                    ab_count += 1
                elif duty_val == "H" or duty_val == "WW":
                    holiday_count += 1
                ot_cell = ws.cell(row=data_start_row + 1, column=c)
                if ot_cell.value:
                    ot_total += float(ot_cell.value) 
            total_paid_days = present_count + fh_nh_count + holiday_count + cl_count
            total_days = (end_date - start_date).days + 1   
            fixed_gross = getattr(emp, "gross_pay", 0)
            # b_shift_count = emp.custom_no_of_2_shift or 0
            # c_shift_count = emp.custom_no_of_3_shift or 0
            shift_allowance = (b_shift_count * 20) + (c_shift_count * 40)
            if total_paid_days == total_days:
                if cl_count == 0:
                    att_bonus = 800
                else:
                    att_bonus = 500
            else:
                att_bonus = 0
            # festival_allowance = fh_nh_count * 200
            if ot_total:
                ot_amount = round((fixed_gross / total_days / 8) * ot_total)
            else:
                ot_amount = 0
            deductions = 0
            earned_gross = round((fixed_gross / total_days) * total_paid_days)
            earned_stipend = shift_allowance + att_bonus + festival_allowance + ot_amount + earned_gross - deductions
            service_charge = 750
            insurance = 150
            total_amount = earned_stipend + service_charge + insurance
            extra_start = last_date_col + 1

            ws.cell(row=data_start_row, column=extra_start, value=total_days)  
            ws.cell(row=data_start_row, column=extra_start + 1, value=present_count)  
            ws.cell(row=data_start_row, column=extra_start + 2, value=ot_total)       
            ws.cell(row=data_start_row, column=extra_start + 3, value=fh_nh_count)    
            ws.cell(row=data_start_row, column=extra_start + 4, value=cl_count)      
            ws.cell(row=data_start_row, column=extra_start + 5, value=ab_count)      
            ws.cell(row=data_start_row, column=extra_start + 6, value=holiday_count)  
            ws.cell(row=data_start_row, column=extra_start + 7, value=total_paid_days)
            ws.cell(row=data_start_row, column=extra_start + 8,   value=fixed_gross).number_format = '#,##0'        
            ws.cell(row=data_start_row, column=extra_start + 9,   value=shift_allowance)    
            ws.cell(row=data_start_row, column=extra_start + 10,  value=att_bonus)          
            ws.cell(row=data_start_row, column=extra_start + 11,  value=festival_allowance)
            ws.cell(row=data_start_row, column=extra_start + 12, value=ot_amount).number_format = '#,##0'     
            ws.cell(row=data_start_row, column=extra_start + 13, value=deductions)       
            ws.cell(row=data_start_row, column=extra_start + 14, value=earned_gross).number_format = '#,##0'
            ws.cell(row=data_start_row, column=extra_start + 15, value=earned_stipend).number_format = '#,##0'   
            ws.cell(row=data_start_row, column=extra_start + 16, value=service_charge)   
            ws.cell(row=data_start_row, column=extra_start + 17, value=insurance)        
            ws.cell(row=data_start_row, column=extra_start + 18, value=total_amount).number_format = '#,##0'
            
            dept_totals["total_days"] += total_days
            dept_totals["present"] += present_count
            dept_totals["ot_total"] += ot_total
            dept_totals["fh_nh"] += fh_nh_count
            dept_totals["cl"] += cl_count
            dept_totals["ab"] += ab_count
            dept_totals["holiday"] += holiday_count
            dept_totals["paid_days"] += total_paid_days
            dept_totals["fixed_gross"] += fixed_gross
            dept_totals["shift_allowance"] += shift_allowance
            dept_totals["att_bonus"] += att_bonus
            dept_totals["festival_allowance"] += festival_allowance
            dept_totals["ot_amount"] += ot_amount
            dept_totals["deductions"] += deductions
            dept_totals["earned_gross"] += earned_gross
            dept_totals["earned_stipend"] += earned_stipend
            dept_totals["service_charge"] += service_charge
            dept_totals["insurance"] += insurance
            dept_totals["total_amount"] += total_amount

            for col in range(extra_start, extra_start+19):   
                cell = ws.cell(row=data_start_row, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center"  ) 
        serial += 1
        data_start_row += 2

        

        for row in ws.iter_rows(min_row=day_row+1, max_row=data_start_row-1, min_col=2, max_col=4):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(size=11)
                cell.border = thin_border

    if current_department:
        dept_end_row = data_start_row - 1
        write_dept_total_row(dept_start_row, dept_end_row)

    def write_final_total_row(start_row, end_row):
        nonlocal data_start_row, last_date_col
        light_green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

        ws.merge_cells(start_row=data_start_row, start_column=2, end_row=data_start_row, end_column=6)
        total_cell = ws.cell(row=data_start_row, column=2, value="TOTAL")
        total_cell.font = Font(bold=True, size=12)
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for col in range(2, 8):
            ws.cell(row=data_start_row, column=col).fill = light_green_fill
            ws.cell(row=data_start_row, column=col).border = thin_border

        extra_start = last_date_col + 1
        numeric_columns = [
            "total_days", "present", "ot_total", "fh_nh", "cl", "ab", "holiday",
            "paid_days", "fixed_gross", "shift_allowance", "att_bonus",
            "festival_allowance", "ot_amount", "deductions", "earned_gross",
            "earned_stipend", "service_charge", "insurance", "total_amount"
        ]

        totals = {key: 0 for key in numeric_columns}

        for row in range(start_row, end_row + 1):
            for i, key in enumerate(numeric_columns):
                value = ws.cell(row=row, column=extra_start + i).value or 0
                totals[key] += value

        for i, key in enumerate(numeric_columns):
            col_index = extra_start + i
            value = totals[key] / 2
            cell = ws.cell(row=data_start_row, column=col_index, value=value)
            cell.fill = light_green_fill
            cell.border = thin_border
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if key in ["fixed_gross", "shift_allowance", "att_bonus", "festival_allowance",
                    "ot_amount", "deductions", "earned_gross", "earned_stipend",
                    "service_charge", "insurance", "total_amount"]:
                cell.number_format = '#,##0'

        for col in range(8, extra_start):
            cell = ws.cell(row=data_start_row, column=col)
            cell.fill = light_green_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        data_start_row += 1

    first_data_row = day_row + 1  
    last_data_row = data_start_row - 1  
    write_final_total_row(first_data_row, last_data_row)
    

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)

    frappe.response['filename'] = filename
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'


@frappe.whitelist()
def check_slip_exists(start_date, end_date):
    exists = frappe.db.exists("Salary Slip", {
        "employee_type": "NAPS",
        "start_date": start_date,
        "end_date": end_date,
        "docstatus": ("!=",2)
    })
    return {"exists": bool(exists)}