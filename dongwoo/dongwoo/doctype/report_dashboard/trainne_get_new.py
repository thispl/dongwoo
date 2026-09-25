import frappe
from frappe.utils.background_jobs import enqueue
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from six import BytesIO
from datetime import datetime
from dateutil.relativedelta import relativedelta
from dateutil import parser

@frappe.whitelist()
def download(start_date=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee List"

    # Colors
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    pink_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")      # Light Pink
    biscuit_fill = PatternFill(start_color="F5DEB3", end_color="F5DEB3", fill_type="solid")  # Biscuit/Wheat
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")    # Light Green

    # Border
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    start_row = 4

    # --- Left side fixed headers (10 cols, 2 rows merge) ---
    headers = ["S.No", "Dep", "Emp No", "Name", "Acco No", "MOP",
               "Designation", "ESI No", "DOJ", "No. of Days Paid"]

    for col, header in enumerate(headers, start=1):
        ws.merge_cells(start_row=start_row, start_column=col,
                       end_row=start_row+1, end_column=col)
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = yellow_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # --- Stipend Section (11 cols) ---
    stipend_start_col = len(headers) + 1
    stipend_end_col = stipend_start_col + 10

    ws.merge_cells(start_row=start_row, start_column=stipend_start_col,
                   end_row=start_row, end_column=stipend_end_col)
    st_title = ws.cell(row=start_row, column=stipend_start_col, value="STIPEND CALCULATION")
    st_title.fill = yellow_fill
    st_title.font = Font(bold=True)
    st_title.alignment = Alignment(horizontal="center", vertical="center")
    st_title.border = thin_border

    # Month names
    if start_date:
        try:
            base_date = datetime.strptime(start_date, "%Y-%m-%d")
        except ValueError:
            base_date = parser.parse(start_date)
    else:
        base_date = datetime.today()

    prev_month_1 = (base_date - relativedelta(months=1)).strftime("%B")
    prev_month_2 = (base_date - relativedelta(months=2)).strftime("%B")

    stipend_sub_headers = [
        "Stipend", "Gross", "OT", f"OT({prev_month_1})", f"OT({prev_month_2})",
        "Fest. Allowance", "Arrear", "PAttendance Allowance",
        "Shift Allowance", "Site Allowance", "Total"
    ]

    for i, header in enumerate(stipend_sub_headers):
        cell = ws.cell(row=start_row+1, column=stipend_start_col+i, value=header)
        # Color logic
        if "OT(" in header:  
            cell.fill = green_fill     # Previous 2 months OT → Green
        elif header == "Total":
            cell.fill = biscuit_fill   # Total → Biscuit
        else:
            cell.fill = yellow_fill

        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # --- Deduction Section (6 cols) ---
    deduction_start_col = stipend_end_col + 1
    deduction_end_col = deduction_start_col + 5

    ws.merge_cells(start_row=start_row, start_column=deduction_start_col,
                   end_row=start_row, end_column=deduction_end_col)
    ded_title = ws.cell(row=start_row, column=deduction_start_col, value="DEDUCTION")
    ded_title.fill = yellow_fill
    ded_title.font = Font(bold=True)
    ded_title.alignment = Alignment(horizontal="center", vertical="center")
    ded_title.border = thin_border

    deduction_sub_headers = ["LWF", "Advance", "PF", "ESI", "Total Deduction", "LOP"]

    for i, header in enumerate(deduction_sub_headers):
        cell = ws.cell(row=start_row+1, column=deduction_start_col+i, value=header)
        if "Total" in header:
            cell.fill = biscuit_fill   # Total Deduction → Biscuit
        else:
            cell.fill = yellow_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # --- Net Total Section (2 rows merged, 1 col) ---
    net_total_col = deduction_end_col + 1
    ws.merge_cells(start_row=start_row, start_column=net_total_col,
                   end_row=start_row+1, end_column=net_total_col)
    nt_cell = ws.cell(row=start_row, column=net_total_col, value="Net Total")
    nt_cell.fill = pink_fill     # Light Pink
    nt_cell.font = Font(bold=True)
    nt_cell.alignment = Alignment(horizontal="center", vertical="center")
    nt_cell.border = thin_border

    # --- Return file ---
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.filename = "Employee_List.xlsx"
    frappe.local.response.type = "download"
