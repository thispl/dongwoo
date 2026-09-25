import frappe
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from six import BytesIO
from datetime import datetime, timedelta
import calendar
from frappe.utils import getdate

@frappe.whitelist()
def download():
    filename = 'Monthly Salary Report'
    test = build_xlsx_response(filename)

def format_currency(value):
    if value is None:
        return "0"
    
    number_str = str(int(value))
    
    if len(number_str) > 3:
        last_three = number_str[-3:]  
        other_digits = number_str[:-3] 
        
        formatted_other = []
        while len(other_digits) > 2:
            formatted_other.append(other_digits[-2:])
            other_digits = other_digits[:-2]
        if other_digits:
            formatted_other.append(other_digits)        
        formatted_other.reverse()
        formatted_number = ','.join(formatted_other) + ',' + last_three
    else:
        formatted_number = number_str
    
    return formatted_number


    

def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()

    att_date = getdate(args.start_date)
    att_month = att_date.strftime("%B")
    att_year = att_date.year
    # Execute the SQL query to fetch employee details along with corresponding salary slip names
    employee_details = frappe.db.sql("""
    SELECT e.bank_ac_no AS 'AC No', e.employee, e.first_name AS 'Employee Name', s.employee_name AS 'Salary Slip Ename', e.mop, e.employee_type AS 'Employee Type',
           e.designation AS 'Designation', e.employee_type AS 'Etype', e.department AS 'Department', e.uan_number AS "UAN NO", e.esi_number AS "Esi", e.date_of_joining AS 'Date of Joining',
           s.payment_days AS 'Payment Days', s.leave_without_pay AS 'LOP',s.absent_days AS 'ABS', s.total_working_days AS 'Days in Month',
           e.basic AS 'Basic Pay', e.house_rent_allowance AS 'HRA', e.medical_allowance AS 'Medical Allowance', e.conveyance_allowance AS 'Conveyance',
           e.education_allowance AS 'Education Allowance', e.leave_and_travel_allowance AS 'LTA', e.dress_allowance AS 'Dress Allowance', e.gross_pay AS 'Fixed Gross', 
           s.name AS 'Name', s.total_deduction AS 'Total Deduction'
    FROM `tabEmployee` e
    INNER JOIN `tabSalary Slip` s ON e.name = s.employee
    INNER JOIN `tabDepartment` d ON e.department = d.name
    INNER JOIN `tabDesignation` de ON de.name = e.designation
    WHERE e.name = s.employee and s.start_date <= %s AND s.end_date >= %s
    ORDER BY d.order_value, e.employee_type, de.order, e.date_of_joining
""", (args.start_date, args.end_date), as_dict=True)# frappe.errprint(args.start_date)
    # frappe.errprint(args.end_date)
    start_date = datetime.strptime(args.start_date, '%Y-%m-%d')
    current_year = start_date.year
    current_month = start_date.month
    if current_month == 1:
        prev_month = 12
        prev_year = current_year - 1
    else:
        prev_month = current_month - 1
        prev_year = current_year
    if prev_month == 1:
        prev_prev_month = 12
        prev_prev_year = prev_year - 1
    else:
        prev_prev_month = prev_month - 1
        prev_prev_year = prev_year
    prev_month_name = calendar.month_name[prev_month]
    next_month_name = calendar.month_name[prev_prev_month]
    pmonth= 'OT'+'('+prev_month_name+')'
    nmonth= 'OT'+'('+next_month_name+')'
    if start_date.month == 1:
        prev_month = 12
        prev_year = start_date.year - 1
    else:
        prev_month = start_date.month - 1
        prev_year = start_date.year
    
    prev_month_start = datetime(prev_year, prev_month, 1)
    prev_month_end = (prev_month_start + timedelta(days=31)).replace(day=1) - timedelta(days=1)    
    if prev_month == 1:
        prev_prev_month = 12
        prev_prev_year = prev_year - 1
    else:
        prev_prev_month = prev_month - 1
        prev_prev_year = prev_year
    prev_prev_month_start = datetime(prev_prev_year, prev_prev_month, 1)
    prev_prev_month_end = (prev_prev_month_start + timedelta(days=31)).replace(day=1) - timedelta(days=1)
    # prev_month_start, prev_month_end, prev_prev_month_start, prev_prev_month_end
    ws = wb.create_sheet(sheet_name, 0)
    ws.append(['DWSI SALARY STATEMENT FOR THE MONTH OF' ''  +str(att_month)+ '-' +str(att_year)])
    
    ws.append(['S.NO','DEPT','EMP NO.','NAMES','A/c No','MOP','DESIGNATION','PF UAN NO','ESI NO','DATE OF JOINING','FIXED GROSS','','','','','','','','DAYS','','','SALARY CALCULATION','','','','','','','','','','','','','','','','','','','','DEDUCTION','','','','','','','','','','','TAKE HOME'])

    ws.append(['S.NO','DEPT','EMP NO.','NAMES','A/c No','MOP','DESIGNATION','PF UAN NO','ESI NO','DATE OF JOINING','Basic Pay','House rent Allowance','Medical Allowance','Convey.Allow','Edu.Allow','Leave Travel Allow','Dress Allow','Fixed Gross','LOP Days','No.of.Days Paid','Days in a Month','Basic Pay','House rent Allowance','Medical Allowance','Convey.Allow','Edu.Allow','Leave Travel Allow','Dress Allow','Gross','OT',pmonth,nmonth,'Bus Fare','Fest.Allow','Arrears','Attendance Bonus','Supervisor Allow','Performance Allow','Shift Allow','Special Duty Allow','Total','PF(12%)','VPF','TOTAL PF DED.','LWF','IT','Prof.Tax','Advance','Miscellenous Ded','ESI','Total Deduction','LOP'+'ABS'])
    
    
    serial_number = 1
    previous_department = None
    prev_category=None
    pre_category="Staff"
    department_total = 0
    total=0
    sub=0
    count=0
    count_1=0
    total1=0
    total_pf_dec=0
    takehome=0
    
    grand_total_basic=0
    grand_total_hra=0
    grand_total_med=0
    grand_total_convey=0
    grand_total_edu=0
    grand_total_lta=0
    grand_total_dress=0
    grand_total_gross=0
    grand_total_lop=0
    grand_total_days=0
    grand_total_month=0
    grand_total_ebasic=0
    grand_total_ehra=0
    grand_total_emed=0
    grand_total_elwf=0
    grand_total_econvey=0
    grand_total_eedu=0
    grand_total_elta=0
    grand_total_edress=0
    grand_total_ot=0
    grand_total_egross=0
    grand_total_bus=0
    grand_total_fes=0
    grand_total_arrear=0
    grand_total_att=0
    grand_total_supervisor=0
    grand_total_shift=0
    grand_total_spec=0
    grand_total_etotal=0
    grand_total_pf=0
    grand_total_vpf=0
    grand_total_totalpf=0
    grand_total_it=0
    grand_total_mis=0
    grand_total_pt=0
    grand_total_pot=0
    grand_total_ppot=0
    grand_total_advance=0
    grand_total_esi=0
    grand_total_totaldec=0
    grand_total_dlop=0
    grand_total_takehome=0
    grand_total_desig=0
    wgrand_total_basic=0
    wgrand_total_hra=0
    wgrand_total_med=0
    wgrand_total_convey=0
    wgrand_total_edu=0
    wgrand_total_lta=0
    wgrand_total_dress=0
    wgrand_total_gross=0
    wgrand_total_lop=0
    wgrand_total_days=0
    wgrand_total_month=0
    wgrand_total_ebasic=0
    wgrand_total_elwf=0
    wgrand_total_ehra=0
    wgrand_total_emed=0
    wgrand_total_econvey=0
    wgrand_total_eedu=0
    wgrand_total_elta=0
    wgrand_total_edress=0
    wgrand_total_ot=0
    wgrand_total_egross=0
    wgrand_total_bus=0
    wgrand_total_fes=0
    wgrand_total_arrear=0
    wgrand_total_att=0
    wgrand_total_supervisor=0
    wgrand_total_shift=0
    wgrand_total_spec=0
    wgrand_total_etotal=0
    wgrand_total_pf=0
    wgrand_total_vpf=0
    wgrand_total_totalpf= 0
    wgrand_total_it=0
    wgrand_total_mis=0
    wgrand_total_pt=0
    wgrand_total_pot=0
    wgrand_total_ppot=0
    wgrand_total_advance=0
    wgrand_total_esi=0
    wgrand_total_totaldec=0
    wgrand_total_dlop=0
    wgrand_total_takehome=0
    wgrand_total_desig=0
    grand_total_desigw=0
    grand_total_desigs=0
    grand_total_desigt=0
    sgrand_total_basic=0
    sgrand_total_hra=0
    sgrand_total_med=0
    sgrand_total_convey=0
    sgrand_total_edu=0
    sgrand_total_lta=0
    sgrand_total_dress=0
    sgrand_total_gross=0
    sgrand_total_lop=0
    sgrand_total_days=0
    sgrand_total_month=0
    sgrand_total_ebasic=0
    sgrand_total_elwf=0
    sgrand_total_ehra=0
    sgrand_total_emed=0
    sgrand_total_econvey=0
    sgrand_total_eedu=0
    sgrand_total_elta=0
    sgrand_total_edress=0
    sgrand_total_ot=0
    sgrand_total_egross=0
    sgrand_total_bus=0
    sgrand_total_fes=0
    sgrand_total_arrear=0
    sgrand_total_att=0
    sgrand_total_supervisor=0
    sgrand_total_shift=0
    sgrand_total_spec=0
    sgrand_total_etotal=0
    sgrand_total_pf=0
    sgrand_total_vpf=0
    sgrand_total_totalpf= 0
    sgrand_total_it=0
    sgrand_total_mis=0
    sgrand_total_pt=0
    sgrand_total_pot=0
    sgrand_total_ppot=0
    sgrand_total_advance=0
    sgrand_total_esi=0
    sgrand_total_totaldec=0
    sgrand_total_dlop=0
    sgrand_total_takehome=0
    sgrand_total_desig=0
    tgrand_total_basic=0
    tgrand_total_hra=0
    tgrand_total_med=0
    tgrand_total_convey=0
    tgrand_total_edu=0
    tgrand_total_lta=0
    tgrand_total_dress=0
    tgrand_total_gross=0
    tgrand_total_lop=0
    tgrand_total_days=0
    tgrand_total_month=0
    tgrand_total_ebasic=0
    tgrand_total_elwf=0
    tgrand_total_ehra=0
    tgrand_total_emed=0
    tgrand_total_econvey=0
    tgrand_total_eedu=0
    tgrand_total_elta=0
    tgrand_total_edress=0
    tgrand_total_ot=0
    tgrand_total_egross=0
    tgrand_total_bus=0
    tgrand_total_fes=0
    tgrand_total_arrear=0
    tgrand_total_att=0
    tgrand_total_supervisor=0
    tgrand_total_shift=0
    tgrand_total_spec=0
    tgrand_total_etotal=0
    tgrand_total_pf=0
    tgrand_total_vpf=0
    tgrand_total_totalpf= 0
    tgrand_total_it=0
    tgrand_total_mis=0
    tgrand_total_pt=0
    tgrand_total_pot=0
    tgrand_total_ppot=0
    tgrand_total_advance=0
    tgrand_total_esi=0
    tgrand_total_totaldec=0
    tgrand_total_dlop=0
    tgrand_total_takehome=0
    tgrand_total_desig=0
    tot_emp=0
    employees_by_category = {
    'Staff': [],
    'Worker': [],
    'D . Trainee': []
    }

    for row in employee_details: 
        employee_type = row['Etype']
        if employee_type in employees_by_category:
            employees_by_category[employee_type].append(row)
        category_order = ['Worker']
    for category in category_order:
        for row in employees_by_category[category]:
            tot_emp+=1
            frappe.errprint(row['employee'])
            prevslip = frappe.db.get_value("Salary Slip", {"employee": row['employee'] ,"start_date":prev_month_start,"end_date":prev_month_end},["name"]) 
            if prevslip:
                prevot = frappe.db.get_value("Salary Detail", {"parent": prevslip , "salary_component": "Overtime"},"amount") or 0.0
            else:
                prevot = 0.0
            pprevslip = frappe.db.get_value("Salary Slip", {"employee": row['employee'] ,"start_date":prev_prev_month_start,"end_date":prev_prev_month_end},["name"]) 
            if pprevslip:
                pprevot = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Overtime"},"amount") or 0.0
            else:
                pprevot = 0.0
            ebasic = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Basic"},"amount") or 0.0
            ehra = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "House Rent Allowance"},"amount") or 0.0
            emedical= frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Medical Allowance"},"amount") or 0.0
            econveyance = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Conveyance Allowance"},"amount")or 0.0
            eeducation = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Education Allowance"},"amount") or 0.0
            elta = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Leave and Travel Allowance"},"amount") or 0.0
            edress = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Dress Allowance"},"amount") or 0.0   
            eot = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Overtime"},"amount") or 0.0
            ebus = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Bus Fare"},"amount") or 0.0
            efes = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Festival Allowance"},"amount") or 0.0
            earrear = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Arrear"},"amount") or 0.0
            eattendance = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Attendance Bonus"},"amount") or 0.0
            esupervisor = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Supervisor Allowance"},"amount") or 0.0
            eshift = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Shift Allowance"},"amount") or 0.0
            esda = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Special Duty Allowance"},"amount") or 0.0
            elwf = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Labour Welfare  Fund"},"amount") or 0.0
            dprovident = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Provident Fund"},"amount") or 0.0
            dvpf = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Voluntary Provident Fund"},"amount") or 0.0
            dit = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Income Tax"},"amount") or 0.0
            dpt = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Professional Tax"},"amount") or 0.0
            dadvance = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Advance"},"amount") or 0.0
            desi = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Employee State Insurance"},"amount") or 0.0		
            dlop = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Loss Of Pay"},"amount") or 0.0
            mis = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Miscellenous"},"amount") or 0.0
            egross = ebasic+ehra+emedical+econveyance+eeducation+elta+edress
            total = egross +ebus+efes+earrear+eattendance+esupervisor+eshift+esda+eot
            takehome= total - row['Total Deduction']
            
            if row['Department'] == previous_department:
                print(previous_department)
            else:
                if previous_department and row['Etype']==pre_category:
                    ws.append(['','SUBTOTAL', 'SUBTOTAL','', '', *sub, '', '','', '',format_currency(sub_total_basic), format_currency(sub_total_hra),format_currency(sub_total_med), format_currency(sub_total_convey),format_currency(sub_total_edu), format_currency(sub_total_lta), format_currency(sub_total_dress), format_currency(sub_total_gross), 
                    sub_total_lop, sub_total_days, sub_total_month, format_currency(sub_total_ebasic), format_currency(sub_total_ehra),format_currency(sub_total_emed), format_currency(sub_total_econvey), format_currency(sub_total_eedu), format_currency(sub_total_elta), format_currency(sub_total_edress),
                    format_currency(sub_total_egross),format_currency(sub_total_ot),format_currency(sub_total_pot),format_currency(sub_total_ppot), format_currency(sub_total_bus), format_currency(sub_total_fes), format_currency(sub_total_arrear), format_currency(sub_total_att), format_currency(sub_total_supervisor),0,format_currency(sub_total_shift),format_currency(sub_total_spec), format_currency(sub_total_etotal), 
                    format_currency(sub_total_pf), format_currency(sub_total_vpf), format_currency(sub_total_totalpf),format_currency(sub_total_elwf), format_currency(sub_total_it), format_currency(sub_total_pt), format_currency(sub_total_advance),format_currency(sub_total_mis),format_currency(sub_total_esi),format_currency(sub_total_totaldec), format_currency(sub_total_dlop),format_currency(sub_total_takehome)])
                count=0
                pre_category=row['Etype']
                department_name = row['Department']
                previous_department = row['Department']
                sub_total_basic=0
                sub_total_hra = 0
                sub_total_med=0
                sub_total_convey=0
                sub_total_edu=0
                sub_total_lta=0
                sub_total_dress=0
                sub_total_gross=0
                sub_total_lop=0
                sub_total_elwf=0
                sub_total_days=0
                sub_total_month=0
                sub_total_ebasic=0
                sub_total_ehra=0
                sub_total_emed=0
                sub_total_econvey=0
                sub_total_eedu=0
                sub_total_elta=0
                sub_total_edress=0
                sub_total_ot=0
                sub_total_egross=0
                sub_total_bus=0
                sub_total_fes=0
                sub_total_arrear=0
                sub_total_att=0
                sub_total_supervisor=0
                sub_total_shift=0
                sub_total_spec=0
                sub_total_etotal=0
                sub_total_pf=0
                sub_total_vpf=0
                sub_total_totalpf=0
                sub_total_it=0
                sub_total_pot=0
                sub_total_ppot=0
                sub_total_mis=0
                sub_total_pt=0
                sub_total_advance=0
                sub_total_esi=0
                sub_total_totaldec=0
                sub_total_dlop=0
                sub_total_takehome=0

            total_pf_dec = dprovident+dvpf
            sub_total_basic+= row['Basic Pay']  
            sub_total_hra+=row['HRA']
            sub_total_med+=row['Medical Allowance']
            sub_total_convey+=row['Conveyance']
            sub_total_edu+=row['Education Allowance']
            sub_total_lta+=row['LTA']
            sub_total_dress+=row['Dress Allowance']
            sub_total_gross+=row['Fixed Gross']
            sub_total_lop+=row['LOP']+row['ABS']
            sub_total_days+=row['Payment Days']
            sub_total_month+=row['Days in Month']
            sub_total_ebasic+=ebasic
            sub_total_ehra+=ehra
            sub_total_emed+=emedical
            sub_total_econvey+=econveyance
            sub_total_eedu+=eeducation
            sub_total_elta+=elta
            sub_total_edress+=edress
            sub_total_ot+=eot
            sub_total_elwf+=elwf
            sub_total_egross+=egross
            sub_total_bus+=ebus
            sub_total_fes+=efes
            sub_total_arrear+=earrear
            sub_total_att+=eattendance
            sub_total_supervisor+=esupervisor
            sub_total_shift+=eshift
            sub_total_spec+=esda
            sub_total_etotal+=total
            sub_total_pf+=dprovident
            sub_total_vpf+=dvpf
            sub_total_totalpf+= total_pf_dec
            sub_total_it+=dit
            sub_total_mis+=mis
            sub_total_pt+=dpt
            sub_total_pot+=prevot
            sub_total_ppot+=pprevot
            sub_total_advance+=dadvance
            sub_total_esi+=desi
            sub_total_totaldec+=row['Total Deduction']
            sub_total_lop+=dlop
            sub_total_takehome+=takehome
            
            
            grand_total_basic+=row['Basic Pay']
            grand_total_hra+=row['HRA']
            grand_total_med+=row['Medical Allowance']
            grand_total_convey+=row['Conveyance']
            grand_total_edu+=row['Education Allowance']
            grand_total_lta+=row['LTA']
            grand_total_dress+=row['Dress Allowance']
            grand_total_gross+=row['Fixed Gross']
            grand_total_lop+=row['LOP']+row['ABS']
            grand_total_days+=row['Payment Days']
            grand_total_month+=row['Days in Month']
            grand_total_ebasic+=ebasic
            grand_total_ehra+=ehra
            grand_total_elwf+=elwf
            grand_total_emed+=emedical
            grand_total_econvey+=econveyance
            grand_total_eedu+=eeducation
            grand_total_elta+=elta
            grand_total_edress+=edress
            grand_total_ot+=eot
            grand_total_egross+=egross
            grand_total_bus+=ebus
            grand_total_fes+=efes
            grand_total_arrear+=earrear
            grand_total_att+=eattendance
            grand_total_supervisor+=esupervisor
            grand_total_shift+=eshift
            grand_total_spec+=esda
            grand_total_etotal+=total
            grand_total_pf+=dprovident
            grand_total_vpf+=dvpf
            grand_total_totalpf+= total_pf_dec
            grand_total_it+=dit
            grand_total_mis+=mis
            grand_total_pt+=dpt
            grand_total_pot+=prevot
            grand_total_ppot+=pprevot
            grand_total_advance+=dadvance
            grand_total_esi+=desi
            grand_total_totaldec+=row['Total Deduction']
            grand_total_dlop+=dlop
            grand_total_takehome+=takehome
            grand_total_desig=serial_number
            if row['Etype'] == 'Worker':
                grand_total_desigw+=1
                wgrand_total_basic+=row['Basic Pay']
                wgrand_total_hra+=row['HRA']
                wgrand_total_med+=row['Medical Allowance']
                wgrand_total_convey+=row['Conveyance']
                wgrand_total_edu+=row['Education Allowance']
                wgrand_total_lta+=row['LTA']
                wgrand_total_dress+=row['Dress Allowance']
                wgrand_total_gross+=row['Fixed Gross']
                wgrand_total_lop+=row['LOP']+row['ABS']
                wgrand_total_days+=row['Payment Days']
                wgrand_total_month+=row['Days in Month']
                wgrand_total_ebasic+=ebasic
                wgrand_total_ehra+=ehra
                wgrand_total_emed+=emedical
                wgrand_total_econvey+=econveyance
                wgrand_total_eedu+=eeducation
                wgrand_total_elta+=elta
                wgrand_total_elwf+=elwf
                wgrand_total_edress+=edress
                wgrand_total_ot+=eot
                wgrand_total_egross+=egross
                wgrand_total_bus+=ebus
                wgrand_total_fes+=efes
                wgrand_total_arrear+=earrear
                wgrand_total_att+=eattendance
                wgrand_total_supervisor+=esupervisor
                wgrand_total_shift+=eshift
                wgrand_total_spec+=esda
                wgrand_total_etotal+=total
                wgrand_total_pf+=dprovident
                wgrand_total_vpf+=dvpf
                wgrand_total_totalpf+= total_pf_dec
                wgrand_total_it+=dit
                wgrand_total_mis+=mis
                wgrand_total_pt+=dpt
                wgrand_total_pot+=prevot
                wgrand_total_ppot+=pprevot
                wgrand_total_advance+=dadvance
                wgrand_total_esi+=desi
                wgrand_total_totaldec+=row['Total Deduction']
                wgrand_total_dlop+=dlop
                wgrand_total_takehome+=takehome
                wgrand_total_desig=serial_number
            if row['Etype'] == 'Staff':
                grand_total_desigs+=1
                sgrand_total_basic+=row['Basic Pay']
                sgrand_total_hra+=row['HRA']
                sgrand_total_med+=row['Medical Allowance']
                sgrand_total_convey+=row['Conveyance']
                sgrand_total_edu+=row['Education Allowance']
                sgrand_total_lta+=row['LTA']
                sgrand_total_dress+=row['Dress Allowance']
                sgrand_total_gross+=row['Fixed Gross']
                sgrand_total_lop+=row['LOP']+row['ABS']
                sgrand_total_days+=row['Payment Days']
                sgrand_total_month+=row['Days in Month']
                sgrand_total_ebasic+=ebasic
                sgrand_total_ehra+=ehra
                sgrand_total_emed+=emedical
                sgrand_total_econvey+=econveyance
                sgrand_total_eedu+=eeducation
                sgrand_total_elta+=elta
                sgrand_total_elwf+=elwf
                sgrand_total_edress+=edress
                sgrand_total_ot+=eot
                sgrand_total_egross+=egross
                sgrand_total_bus+=ebus
                sgrand_total_fes+=efes
                sgrand_total_arrear+=earrear
                sgrand_total_att+=eattendance
                sgrand_total_supervisor+=esupervisor
                sgrand_total_shift+=eshift
                sgrand_total_spec+=esda
                sgrand_total_etotal+=total
                sgrand_total_pf+=dprovident
                sgrand_total_vpf+=dvpf
                sgrand_total_totalpf+= total_pf_dec
                sgrand_total_it+=dit
                sgrand_total_mis+=mis
                sgrand_total_pt+=dpt
                sgrand_total_pot+=prevot
                sgrand_total_ppot+=pprevot
                sgrand_total_advance+=dadvance
                sgrand_total_esi+=desi
                sgrand_total_totaldec+=row['Total Deduction']
                sgrand_total_dlop+=dlop
                sgrand_total_takehome+=takehome
                sgrand_total_desig=serial_number
            if row['Etype'] == 'D . Trainee':
                grand_total_desigt+=1
                tgrand_total_basic+=row['Basic Pay']
                tgrand_total_hra+=row['HRA']
                tgrand_total_med+=row['Medical Allowance']
                tgrand_total_convey+=row['Conveyance']
                tgrand_total_edu+=row['Education Allowance']
                tgrand_total_lta+=row['LTA']
                tgrand_total_dress+=row['Dress Allowance']
                tgrand_total_gross+=row['Fixed Gross']
                tgrand_total_lop+=row['LOP']+row['ABS']
                tgrand_total_days+=row['Payment Days']
                tgrand_total_month+=row['Days in Month']
                tgrand_total_ebasic+=ebasic
                tgrand_total_ehra+=ehra
                tgrand_total_emed+=emedical
                tgrand_total_econvey+=econveyance
                tgrand_total_eedu+=eeducation
                tgrand_total_elta+=elta
                tgrand_total_elwf+=elwf
                tgrand_total_edress+=edress
                tgrand_total_ot+=eot
                tgrand_total_egross+=egross
                tgrand_total_bus+=ebus
                tgrand_total_fes+=efes
                tgrand_total_arrear+=earrear
                tgrand_total_att+=eattendance
                tgrand_total_supervisor+=esupervisor
                tgrand_total_shift+=eshift
                tgrand_total_spec+=esda
                tgrand_total_etotal+=total
                tgrand_total_pf+=dprovident
                tgrand_total_vpf+=dvpf
                tgrand_total_totalpf+= total_pf_dec
                tgrand_total_it+=dit
                tgrand_total_mis+=mis
                tgrand_total_pt+=dpt
                tgrand_total_pot+=prevot
                tgrand_total_ppot+=pprevot
                tgrand_total_advance+=dadvance
                tgrand_total_esi+=desi
                tgrand_total_totaldec+=row['Total Deduction']
                tgrand_total_dlop+=dlop
                tgrand_total_takehome+=takehome
                tgrand_total_desig=serial_number
            
            if row['Etype'] == 'D . Trainee':
                ws.append([serial_number, department_name, row['employee'], row['Employee Name'],row['AC No'],row['mop'], row['Designation'], row['UAN NO'],row['Esi'],row['Date of Joining'], format_currency(row['Basic Pay']), format_currency(row['HRA']), format_currency(row['Medical Allowance']), format_currency(row['Conveyance']), format_currency(row['Education Allowance']), format_currency(row['LTA']), format_currency(row['Dress Allowance']), format_currency(row['Fixed Gross']), row['LOP']+row['ABS'], row['Payment Days'], row['Days in Month'], format_currency(ebasic), format_currency(ehra), format_currency(emedical), format_currency(econveyance), format_currency(eeducation), format_currency(elta), format_currency(edress), format_currency(egross),format_currency(eot),format_currency(prevot),format_currency(pprevot), format_currency(ebus), format_currency(efes), format_currency(earrear), format_currency(eattendance), format_currency(esupervisor),0, format_currency(eshift), format_currency(esda),format_currency(total), format_currency(dprovident), format_currency(dvpf), format_currency(total_pf_dec), format_currency(elwf),format_currency(dit), format_currency(dpt), format_currency(dadvance), format_currency(mis),format_currency(desi), format_currency(row['Total Deduction']),format_currency(dlop),format_currency(takehome)])
            else:
                ws.append([serial_number, department_name, row['employee'], row['Employee Name'],row['AC No'],row['mop'], row['Designation'], row['UAN NO'],'-',row['Date of Joining'], format_currency(row['Basic Pay']), format_currency(row['HRA']), format_currency(row['Medical Allowance']), format_currency(row['Conveyance']), format_currency(row['Education Allowance']), format_currency(row['LTA']), format_currency(row['Dress Allowance']), format_currency(row['Fixed Gross']), row['LOP']+row['ABS'], row['Payment Days'], row['Days in Month'], format_currency(ebasic), format_currency(ehra), format_currency(emedical), format_currency(econveyance), format_currency(eeducation), format_currency(elta), format_currency(edress), format_currency(egross),format_currency(eot),format_currency(prevot),format_currency(pprevot), format_currency(ebus), format_currency(efes), format_currency(earrear), format_currency(eattendance), format_currency(esupervisor),0, format_currency(eshift), format_currency(esda),format_currency(total), format_currency(dprovident), format_currency(dvpf), format_currency(total_pf_dec), format_currency(elwf),format_currency(dit), format_currency(dpt), format_currency(dadvance), format_currency(mis),format_currency(desi), format_currency(row['Total Deduction']),format_currency(dlop),format_currency(takehome)])

            serial_number += 1
            count+=1
            sub = [str(count) + " Members"]
        
        
        mem= [str(grand_total_desig)+ " Members"] 
        wmem= [str(grand_total_desigw)+ " Members"] 
        smem= [str(grand_total_desigs)+ " Members"] 
        tmem= [str(grand_total_desigt)+ " Members"]
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=53)
        for i in range(1,11):
            ws.merge_cells(start_row=2,start_column=i,end_row=3,end_column=i)
            i=i+1
        ws.merge_cells(start_row=2,start_column=11,end_row=2,end_column=18)
        ws.merge_cells(start_row=2,start_column=19,end_row=2,end_column=21)
        ws.merge_cells(start_row=2,start_column=22,end_row=2,end_column=41)
        ws.merge_cells(start_row=2,start_column=42,end_row=2,end_column=52)
        ws.merge_cells(start_row=2,start_column=53,end_row=3,end_column=53)
        if tot_emp > 0:
        
            ws.append(['','SUBTOTAL', 'SUBTOTAL','', '', *sub, '', '','', '',format_currency(sub_total_basic), format_currency(sub_total_hra),format_currency(sub_total_med), format_currency(sub_total_convey),format_currency(sub_total_edu), format_currency(sub_total_lta), format_currency(sub_total_dress), format_currency(sub_total_gross), 
            sub_total_lop, sub_total_days, sub_total_month, format_currency(sub_total_ebasic), format_currency(sub_total_ehra),format_currency(sub_total_emed), format_currency(sub_total_econvey), format_currency(sub_total_eedu), format_currency(sub_total_elta), format_currency(sub_total_edress),
            format_currency(sub_total_egross),format_currency(sub_total_ot),format_currency(sub_total_pot),format_currency(sub_total_ppot), format_currency(sub_total_bus), format_currency(sub_total_fes), format_currency(sub_total_arrear), format_currency(sub_total_att), format_currency(sub_total_supervisor),0,format_currency(sub_total_shift),format_currency(sub_total_spec), format_currency(sub_total_etotal), 
            format_currency(sub_total_pf), format_currency(sub_total_vpf), format_currency(sub_total_totalpf),format_currency(sub_total_elwf), format_currency(sub_total_it), format_currency(sub_total_pt), format_currency(sub_total_advance),format_currency(sub_total_mis),format_currency(sub_total_esi),format_currency(sub_total_totaldec), format_currency(sub_total_dlop),format_currency(sub_total_takehome)])
            prev_category=row['Etype']
            pre_category=row['Etype']
            tot_emp=0
            sub_total_basic=0
            sub_total_hra = 0
            sub_total_med=0
            sub_total_convey=0
            sub_total_edu=0
            sub_total_lta=0
            sub_total_dress=0
            sub_total_gross=0
            sub_total_lop=0
            sub_total_elwf=0
            sub_total_days=0
            sub_total_month=0
            sub_total_ebasic=0
            sub_total_ehra=0
            sub_total_emed=0
            sub_total_econvey=0
            sub_total_eedu=0
            sub_total_elta=0
            sub_total_edress=0
            sub_total_ot=0
            sub_total_egross=0
            sub_total_bus=0
            sub_total_fes=0
            sub_total_arrear=0
            sub_total_att=0
            sub_total_supervisor=0
            sub_total_shift=0
            sub_total_spec=0
            sub_total_etotal=0
            sub_total_pf=0
            sub_total_vpf=0
            sub_total_totalpf=0
            sub_total_it=0
            sub_total_pot=0
            sub_total_ppot=0
            sub_total_mis=0
            sub_total_pt=0
            sub_total_advance=0
            sub_total_esi=0
            sub_total_totaldec=0
            sub_total_dlop=0
            sub_total_takehome=0

            if row['Etype'] == 'Worker':
                ws.append(['', 'TOTAL FOR WORKERS','TOTAL FOR WORKERS', '','',*wmem,'','','','',format_currency(wgrand_total_basic), format_currency(wgrand_total_hra), format_currency(wgrand_total_med), format_currency(wgrand_total_convey), format_currency(wgrand_total_edu), format_currency(wgrand_total_lta), format_currency(wgrand_total_dress), format_currency(wgrand_total_gross),
                        wgrand_total_lop, wgrand_total_days, wgrand_total_month, format_currency(wgrand_total_ebasic),format_currency(wgrand_total_ehra),format_currency(wgrand_total_emed), format_currency(wgrand_total_econvey), format_currency(wgrand_total_eedu), format_currency(wgrand_total_elta), format_currency(wgrand_total_edress),
                            format_currency(wgrand_total_egross), format_currency(wgrand_total_ot),format_currency(wgrand_total_pot),format_currency(wgrand_total_ppot),format_currency(wgrand_total_bus), format_currency(wgrand_total_fes), format_currency(wgrand_total_arrear), format_currency(wgrand_total_att), format_currency(wgrand_total_supervisor),0,format_currency(wgrand_total_shift),format_currency(wgrand_total_spec), format_currency(wgrand_total_etotal), 
                            format_currency(wgrand_total_pf), format_currency(wgrand_total_vpf), format_currency(wgrand_total_totalpf), format_currency(wgrand_total_elwf),format_currency(wgrand_total_it), format_currency(wgrand_total_pt), format_currency(wgrand_total_advance),format_currency(wgrand_total_mis), format_currency(wgrand_total_esi), format_currency(wgrand_total_totaldec), format_currency(wgrand_total_dlop), format_currency(wgrand_total_takehome)])
            elif row['Etype'] == 'Staff':
                ws.append(['', 'TOTAL FOR STAFF','TOTAL FOR STAFF', '','',*smem,'','','','',format_currency(sgrand_total_basic), format_currency(sgrand_total_hra), format_currency(sgrand_total_med), format_currency(sgrand_total_convey), format_currency(sgrand_total_edu), format_currency(sgrand_total_lta), format_currency(sgrand_total_dress), format_currency(sgrand_total_gross),
                        sgrand_total_lop, sgrand_total_days, sgrand_total_month, format_currency(sgrand_total_ebasic),format_currency(sgrand_total_ehra),format_currency(sgrand_total_emed), format_currency(sgrand_total_econvey), format_currency(sgrand_total_eedu), format_currency(sgrand_total_elta), format_currency(sgrand_total_edress),
                            format_currency(sgrand_total_egross), format_currency(sgrand_total_ot),format_currency(sgrand_total_pot),format_currency(sgrand_total_ppot),format_currency(sgrand_total_bus), format_currency(sgrand_total_fes), format_currency(sgrand_total_arrear), format_currency(sgrand_total_att), format_currency(sgrand_total_supervisor),0,format_currency(sgrand_total_shift),format_currency(sgrand_total_spec), format_currency(sgrand_total_etotal), 
                            format_currency(sgrand_total_pf), format_currency(sgrand_total_vpf), format_currency(sgrand_total_totalpf), format_currency(sgrand_total_elwf),format_currency(sgrand_total_it), format_currency(sgrand_total_pt), format_currency(sgrand_total_advance),format_currency(sgrand_total_mis), format_currency(sgrand_total_esi), format_currency(sgrand_total_totaldec), format_currency(sgrand_total_dlop), format_currency(sgrand_total_takehome)])
            elif row['Etype'] == 'D . Trainee':
                ws.append(['', 'TOTAL FOR TRAINEE','TOTAL FOR TRAINEE', '','',*tmem,'','','','',format_currency(tgrand_total_basic), format_currency(tgrand_total_hra), format_currency(tgrand_total_med), format_currency(tgrand_total_convey), format_currency(tgrand_total_edu), format_currency(tgrand_total_lta), format_currency(tgrand_total_dress), format_currency(tgrand_total_gross),
                        tgrand_total_lop, tgrand_total_days, tgrand_total_month, format_currency(tgrand_total_ebasic),format_currency(tgrand_total_ehra),format_currency(tgrand_total_emed), format_currency(tgrand_total_econvey), format_currency(tgrand_total_eedu), format_currency(tgrand_total_elta), format_currency(tgrand_total_edress),
                            format_currency(tgrand_total_egross), format_currency(tgrand_total_ot),format_currency(tgrand_total_pot),format_currency(tgrand_total_ppot),format_currency(tgrand_total_bus), format_currency(tgrand_total_fes), format_currency(tgrand_total_arrear), format_currency(tgrand_total_att), format_currency(tgrand_total_supervisor),0,format_currency(tgrand_total_shift),format_currency(tgrand_total_spec), format_currency(tgrand_total_etotal), 
                            format_currency(tgrand_total_pf), format_currency(tgrand_total_vpf), format_currency(tgrand_total_totalpf), format_currency(tgrand_total_elwf),format_currency(tgrand_total_it), format_currency(tgrand_total_pt), format_currency(tgrand_total_advance),format_currency(tgrand_total_mis), format_currency(tgrand_total_esi), format_currency(tgrand_total_totaldec), format_currency(tgrand_total_dlop), format_currency(tgrand_total_takehome)])
    ws.append(['','Grand Total','Grand Total','','',*mem,'','','','',format_currency(grand_total_basic), format_currency(grand_total_hra), format_currency(grand_total_med), format_currency(grand_total_convey), format_currency(grand_total_edu), format_currency(grand_total_lta), format_currency(grand_total_dress), format_currency(grand_total_gross),
            grand_total_lop, grand_total_days, grand_total_month, format_currency(grand_total_ebasic),format_currency(grand_total_ehra),format_currency(grand_total_emed), format_currency(grand_total_econvey), format_currency(grand_total_eedu), format_currency(grand_total_elta), format_currency(grand_total_edress),
                format_currency(grand_total_egross), format_currency(grand_total_ot),format_currency(grand_total_pot),format_currency(grand_total_ppot),format_currency(grand_total_bus), format_currency(grand_total_fes), format_currency(grand_total_arrear), format_currency(grand_total_att), format_currency(grand_total_supervisor),0,format_currency(grand_total_shift),format_currency(grand_total_spec), format_currency(grand_total_etotal), 
                format_currency(grand_total_pf), format_currency(grand_total_vpf), format_currency(grand_total_totalpf), format_currency(grand_total_elwf),format_currency(grand_total_it), format_currency(grand_total_pt), format_currency(grand_total_advance),format_currency(grand_total_mis), format_currency(grand_total_esi), format_currency(grand_total_totaldec), format_currency(grand_total_dlop), format_currency(grand_total_takehome)])
    align_center = Alignment(horizontal='center',vertical='center')

    for cell in ws["1:1"]:
        cell.font = Font(bold=True,size=14)
        cell.alignment = align_center

    for cell in ws["2:2"]:
        if cell.column > 10 and cell.column != 50:
            # cell.font = Font(bold=True)
        # 	cell.alignment = align_center 
        # else:
            cell.font = Font(bold=True,size=10)
            cell.alignment = align_center     
    for cell in ws["2:2"]:
        cell.alignment = align_center 
    ws['A1'].fill = PatternFill(fgColor="FFFFFF", fill_type = "solid")

    for header in ws.iter_rows(min_row=2, max_row=3,min_col=1, max_col=52):
        for cell in header:
            cell.fill = PatternFill(fgColor='FFFF00', fill_type="solid")
    for header in ws.iter_rows(min_row=2, max_row=3,min_col=53, max_col=53):
        for cell in header:
            cell.fill = PatternFill(fgColor='FFC0CB', fill_type="solid")
    for header in ws.iter_rows(min_row=3, max_row=3,min_col=31, max_col=32):
        for cell in header:
            cell.fill = PatternFill(fgColor='3e8004', fill_type="solid")
    for header in ws.iter_rows(min_row=2, max_row=2,min_col=11, max_col=18):
        for cell in header:
            cell.fill = PatternFill(fgColor='CD853F', fill_type="solid")
            
    for header in ws.iter_rows(min_row=2, max_row=ws.max_row,min_col=19, max_col=21):
        for cell in header:
            cell.fill = PatternFill(fgColor='ADD8E6', fill_type="solid")
    
    border = Border(left=Side(border_style='thin', color='000000'),
                 right=Side(border_style='thin', color='000000'),
                 top=Side(border_style='thin', color='000000'),
                 bottom=Side(border_style='thin', color='000000'))

    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=53):
        for cell in rows:
            cell.border = border
    
    grey_fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=11, max_col=18):
        for cell in row:
            cell.fill = grey_fill
    pink_fill = PatternFill(start_color='FFC0CB', end_color='FFC0CB', fill_type='solid')
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=53, max_col=53):
        for cell in row:
            cell.fill = pink_fill
    yellow_fill = PatternFill(start_color='fff600', end_color='fff600', fill_type='solid')
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=29, max_col=29):
        for cell in row:
            cell.fill = yellow_fill
    y_fill = PatternFill(start_color='fceda2', end_color='fceda2', fill_type='solid')
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=41, max_col=41):
        for cell in row:
            cell.fill = y_fill
    yellow = PatternFill(start_color='fff600',end_color='fff600',fill_type='solid')
    green = PatternFill(start_color='adff2f',end_color='adff2f',fill_type='solid')
    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=3, max_col=3):
        for cell in rows:
            if cell.value == 'SUBTOTAL':
                for cell in ws[cell.row]:
                    cell.fill = green
                    ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=4) 
            if cell.value == 'Grand Total' or cell.value=='TOTAL FOR WORKERS' or cell.value=='TOTAL FOR TRAINEE' or cell.value=='TOTAL FOR STAFF':
                for cell in ws[cell.row]:
                    cell.fill = yellow
                    ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=4) 

    max_row = ws.max_row
    previous_department = None
    merge_start_row = None

    for row in range(1, max_row + 1):
        cell_value = ws.cell(row=row, column=2).value

        # Apply center alignment to the current cell
        ws.cell(row=row, column=2).alignment = Alignment(horizontal='center', vertical='center')

        if cell_value is None or cell_value.strip() == '':
            continue  # Skip empty rows

        if cell_value == previous_department:
            # Continue merging if the department is the same
            continue
        else:
            # If a new department is found, merge the previous group
            if merge_start_row is not None and previous_department is not None:
                ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=row - 1, end_column=2)

            # Update tracking variables for the new department
            previous_department = cell_value
            merge_start_row = row

    # Merge the last department group at the end of the loop
    if merge_start_row is not None and previous_department is not None:
        ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=max_row, end_column=2)
    right_align = Alignment(horizontal="right")
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=11, max_col=ws.max_column):
        for cell in row:
            cell.alignment = right_align
    # Handle the last group if needed
    if merge_start_row is not None and previous_department is not None:
        end_row = max_row
        ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=end_row, end_column=2)
    for cell in ws["3:3"]:
        cell.alignment = align_center
    align_left = Alignment(horizontal='left', vertical='center')
    for row in range(4, ws.max_row + 1): 
        for col in [1, 10]:
            cell = ws.cell(row=row, column=col) 
            cell.alignment = align_center
    for row in range(4, ws.max_row + 1): 
        cell = ws.cell(row=row, column=2) 
        cell.alignment = align_left
    for col in [2,4,5,7,8,10]:
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 20 
    ws.column_dimensions['C'].width = 15 
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'
    
