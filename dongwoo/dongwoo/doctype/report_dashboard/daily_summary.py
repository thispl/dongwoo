from __future__ import unicode_literals
import frappe
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.data import format_date
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import re
from frappe import _
import frappe
from frappe.model.document import Document
from datetime import date, timedelta, datetime,time
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime,today, format_date)
import math
from frappe.utils import add_months, cint, flt, getdate, time_diff_in_hours,time_diff_in_seconds
import locale
import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
import openpyxl.styles as styles

@frappe.whitelist()
def download(start_date):
    filename = 'Attendance Summary'
    args = {'start_date':start_date}
    frappe.msgprint("Report is generating in background. Kindly check after few minutes.")
    enqueue(daily_attendance_summary, queue='default', timeout=6000, event='attendance_summary',filename=filename,args=args)
    

def make_xlsx(data,args, sheet_name=None, wb=None, column_widths=None):
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
         
    ws = wb.create_sheet(sheet_name, 0)

    header_date = get_title()
    ws.append(header_date)

    header_date = get_title_1(args)
    ws.append(header_date)

    header_date = get_title_2()
    ws.append(header_date)

    header_date = get_title_3()
    ws.append(header_date)

    data=get_data_1(args)
    for d in data:
        ws.append(d)

    header_date = get_title_6()
    ws.append(header_date)

    header_date = get_title_4()
    ws.append(header_date) 

    header_date = get_title_3()
    ws.append(header_date)

    data=get_data_2(args)
    for d in data:
        ws.append(d)

    header_date = get_title_6()
    ws.append(header_date)

    header_date = get_title_5()
    ws.append(header_date)

    header_date = get_title_3()
    ws.append(header_date)

    data=get_data_3(args)
    for d in data:
        ws.append(d)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column= len(get_title_3()) )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column= len(get_title_3()) )

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column= 2 )
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column= 2 + ((len(get_title_3()) - 2)/3) )
    ws.merge_cells(start_row=3, start_column=3 + ((len(get_title_3()) - 2)/3), end_row=3, end_column= 2 + ((len(get_title_3()) - 2)/3) + ((len(get_title_3()) - 2)/3) )
    ws.merge_cells(start_row=3, start_column=3 + ((len(get_title_3()) - 2)/3) + ((len(get_title_3()) - 2)/3), end_row=3, end_column= len(get_title_3()) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 4, start_column=1, end_row=len(get_data_1(args)) + 4, end_column= 2)

    ws.merge_cells(start_row=len(get_data_1(args)) + 5, start_column=1, end_row=len(get_data_1(args)) + 5, end_column= len(get_title_3()))
    
    ws.merge_cells(start_row=len(get_data_1(args)) + 6, start_column=1, end_row=len(get_data_1(args)) + 6, end_column= 2 )
    ws.merge_cells(start_row=len(get_data_1(args)) + 6, start_column=3, end_row=len(get_data_1(args)) + 6, end_column= 2 + ((len(get_title_3()) - 2)/3) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 6, start_column=3 + ((len(get_title_3()) - 2)/3), end_row=len(get_data_1(args)) + 6, end_column= 2 + ((len(get_title_3()) - 2)/3) + ((len(get_title_3()) - 2)/3) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 6, start_column=3 + ((len(get_title_3()) - 2)/3) + ((len(get_title_3()) - 2)/3), end_row=len(get_data_1(args)) + 6, end_column= len(get_title_3()) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 7 + len(get_data_1(args)), start_column=1, end_row=len(get_data_1(args)) + 7 + len(get_data_1(args)), end_column= 2)
    
    ws.merge_cells(start_row=len(get_data_1(args)) + 8 + len(get_data_1(args)), start_column=1, end_row=len(get_data_1(args)) + 8 + len(get_data_1(args)), end_column= len(get_title_3()) )
    
    ws.merge_cells(start_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), start_column=1, end_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), end_column= 2 )
    ws.merge_cells(start_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), start_column=3, end_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), end_column= 2 + ((len(get_title_3()) - 2)/3) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), start_column=3 + ((len(get_title_3()) - 2)/3), end_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), end_column= 2 + ((len(get_title_3()) - 2)/3) + ((len(get_title_3()) - 2)/3) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), start_column=3 + ((len(get_title_3()) - 2)/3) + ((len(get_title_3()) - 2)/3), end_row=len(get_data_1(args)) + 9 + len(get_data_1(args)), end_column= len(get_title_3()) )
    ws.merge_cells(start_row=len(get_data_1(args)) + 10 + len(get_data_1(args)) + len(get_data_1(args)), start_column=1, end_row=len(get_data_1(args)) + 10 + len(get_data_1(args)) + len(get_data_1(args)), end_column= 2)
    
    align_center = Alignment(horizontal='center',vertical='center')
    left = Alignment(horizontal="left")
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin'))
    for rows in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=len(get_title_3())):
        for cell in rows:
            cell.font = Font(bold=True)
            cell.alignment = align_center
    for rows in ws.iter_rows(min_row=3, max_row=len(get_data_1(args)) + 10 + len(get_data_1(args)) + len(get_data_1(args)), min_col=1, max_col=len(get_title_3())):
        for cell in rows:
            cell.font = Font(bold=True)
            cell.alignment = align_center
    for rows in ws.iter_rows(min_row=3, min_col=1, max_row=len(get_data_1(args)) + 10 + len(get_data_1(args)) + len(get_data_1(args)), max_col= len(get_title_3()) ): 
        for cell in rows:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fgColor='DAD8D9', fill_type="solid")
            cell.alignment = align_center
    for rows in ws.iter_rows(min_row=3, min_col=3 + ((len(get_title_3()) - 2) // 3), max_row=len(get_data_1(args)) + 10 + len(get_data_1(args)) + len(get_data_1(args)) , max_col=2 + ((len(get_title_3()) - 2) // 3) + ((len(get_title_3()) - 2) // 3)):
        for cell in rows:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fgColor='FBF9FA', fill_type="solid")
            cell.alignment = align_center
    
    for rows in ws.iter_rows(min_row=1, min_col=1, max_row=len(get_data_1(args)) + 10 + len(get_data_1(args)) + len(get_data_1(args)), max_col= len(get_title_3()) ):
        for cell in rows:
            cell.font = Font(bold=True)
            cell.border = border

    for rows in ws.iter_rows(min_row=5, min_col=1, max_row= len(get_data_1(args)) + 3, max_col=2):
        for cell in rows:
            cell.alignment = left
            cell.font = Font(bold=False)
            cell.fill = PatternFill(fgColor='FBF9FA', fill_type="solid")
    
    for rows in ws.iter_rows(min_row=len(get_data_1(args)) + 8, min_col=1, max_row= len(get_data_1(args)) + 6 + len(get_data_1(args)), max_col=2):
        for cell in rows:
            cell.alignment = left
            cell.font = Font(bold=False)
            cell.fill = PatternFill(fgColor='FBF9FA', fill_type="solid")

    for rows in ws.iter_rows(min_row=len(get_data_1(args)) + 11 + len(get_data_1(args)), min_col=1, max_row= len(get_data_1(args)) + 9 + len(get_data_1(args)) + len(get_data_1(args)), max_col=2):
        for cell in rows:
            cell.alignment = left
            cell.font = Font(bold=False)
            cell.fill = PatternFill(fgColor='FBF9FA', fill_type="solid")

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 30

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def daily_attendance_summary(filename,args):
    xlsx_file = make_xlsx(filename, args)
    ret = frappe.get_doc({
            "doctype": "File",
            "attached_to_name": ' ',
            "attached_to_doctype": 'Report Dashboard',
            "attached_to_field": 'attach',
            "file_name": filename + '.xlsx',
            "is_private": 0,
            "content": xlsx_file.getvalue(),
            "decode": False
        })
    ret.save(ignore_permissions=True)
    frappe.db.commit()
    attached_file = frappe.get_doc("File", ret.name)
    frappe.db.set_value('Report Dashboard',None,'attach',attached_file.file_url)
    

@frappe.whitelist()
def get_title():
    status = []
    status = ['DongWoo Surfacetech (India) Pvt Ltd']
    return status

@frappe.whitelist()
def get_title_1(args):
    status = []
    status = ['Attendance Summary ' + args['start_date']]
    return status

@frappe.whitelist()
def get_title_2():
    row = []
    row += ["Regular"," "]
    shift_types = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
    for shift_type in shift_types:
        count = frappe.db.count("Contractor")
        row += ["Shift " + shift_type['name']]
        for _ in range(count +	6):
            row += [" "]
    return row

@frappe.whitelist()
def get_title_3():
    result = []
    result += ["Parent","Department"]
    shifts = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
    for shift in shifts:
        ec = frappe.db.sql("""SELECT name FROM `tabEmployee Type` ORDER BY `order` ASC""", as_dict=True)
        for e in ec:
            if e['name'] != "Contract Employee" and e['name'] != "Director":
                result += [e['name']]
            elif e['name'] == "Contract Employee":
                contractors = frappe.db.sql("""SELECT name FROM `tabContractor` ORDER BY `name` ASC""", as_dict=True)
                for contractor in contractors:
                    result += [contractor['name']]
        result += ["Total","Half Day",'Permission']
    return result

@frappe.whitelist()
def get_title_4():
    row = []
    row += ["Overtime"," "]
    shift_types = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
    for shift_type in shift_types:
        count = frappe.db.count("Contractor")
        count1 = frappe.db.count("Employee Type")
        row += ["Shift " + shift_type['name']]
        for _ in range(count-2 +count1):
            row += [" "]
    return row

@frappe.whitelist()
def get_title_5():
    row = []
    row += ["Total (Regular + OT)"," "]
    shift_types = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
    for shift_type in shift_types:
        count = frappe.db.count("Contractor")
        count1 = frappe.db.count("Employee Type")
        row += ["Shift " + shift_type['name']]
        for _ in range(count-2 +count1):
            row += [" "]
    return row

@frappe.whitelist()
def get_title_6():
    row = []
    row += [ ]
    shift_types = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
    for shift_type in shift_types:
        count = frappe.db.count("Contractor")
        count1 = frappe.db.count("Employee Type")
        row += [ ]
        for _ in range(count-2 +count1):
            row += [ ]
    return row



@frappe.whitelist()
def get_data_1(args):
    start_date = args['start_date']
    status = []

    # Fetch all departments and sub-departments
    departments = frappe.db.sql("""
        SELECT name, parent_department FROM `tabDepartment`
        WHERE name != "All Departments"
        ORDER BY name ASC
    """, as_dict=True)

    department_hierarchy = {}
    for dept in departments:
        parent = dept['parent_department']
        if parent!='All Departments':
            if parent:
                department_hierarchy.setdefault(parent, []).append(dept['name'])
            else:
                department_hierarchy[dept['name']] = []

    shifts = frappe.db.sql("""
        SELECT name FROM `tabShift Type`
        WHERE name != 'WW'
        ORDER BY name ASC
    """, as_dict=True)
    shift_names = [shift['name'] for shift in shifts]

    employee_types = frappe.db.sql("""
        SELECT name FROM `tabEmployee Type`
        WHERE name != 'Director'
        ORDER BY `order` ASC
    """, as_dict=True)
    employee_type_names = [e['name'] for e in employee_types]

    contractors = frappe.db.sql("""
        SELECT name FROM `tabContractor`
        ORDER BY name ASC
    """, as_dict=True)
    contractor_names = [c['name'] for c in contractors]

    attendance_records = frappe.db.sql("""
        SELECT shift, employee_type, department, contractor, status, COUNT(*) as count
        FROM `tabAttendance`
        WHERE attendance_date = %(start_date)s
        AND docstatus != 2
        GROUP BY shift, employee_type, department, contractor, status
    """, {'start_date': start_date}, as_dict=True)

    # Convert attendance records into a lookup dictionary
    attendance_dict = {}
    for record in attendance_records:
        key = (record['shift'], record['employee_type'], record['department'], record.get('contractor'))
        if key not in attendance_dict:
            attendance_dict[key] = {'Present': 0, 'Half Day': 0}
        attendance_dict[key][record['status']] = record['count']

    permission_records = frappe.db.sql("""
        SELECT shift, employee_type, department, contractor, COUNT(*) as count
        FROM `tabPermission`
        WHERE permission_date = %(start_date)s
        AND status = 'Approved'
        GROUP BY shift, employee_type, department, contractor
    """, {'start_date': start_date}, as_dict=True)

    permission_dict = {}
    for record in permission_records:
        key = (record['shift'], record['employee_type'], record['department'], record.get('contractor'))
        permission_dict[key] = record['count']

    for dept, sub_departments in department_hierarchy.items():
        for s_dept in sub_departments:
            row = [dept, s_dept]
            for shift in shift_names:
                total_present = total_half_day = total_permission = 0
                for e_type in employee_type_names:
                    if e_type == "Contract Employee":
                        for contractor in contractor_names:
                            key = (shift, e_type, s_dept, contractor)
                            present = attendance_dict.get(key, {}).get("Present", 0)
                            half_day = attendance_dict.get(key, {}).get("Half Day", 0)
                            permission = permission_dict.get(key, 0)
                            total_present += present
                            total_half_day += half_day
                            total_permission += permission
                            row.append(present if present > 0 else "-")
                    else:
                        key = (shift, e_type, s_dept, None)
                        present = attendance_dict.get(key, {}).get("Present", 0)
                        half_day = attendance_dict.get(key, {}).get("Half Day", 0)
                        permission = permission_dict.get(key, 0)
                        total_present += present
                        total_half_day += half_day
                        total_permission += permission
                        row.append(present if present > 0 else "-")

                row.extend([total_present, total_half_day, total_permission])
            status.append(row)

    row = ["Total", '']
    
    shifts = frappe.db.get_list("Shift Type", filters={"name": ["!=", "WW"]}, fields=["name"], order_by="name ASC")
    employee_types = frappe.db.get_list("Employee Type",filters={"name": ["!=", ""]}, fields=["name"], order_by="`order` ASC")

    # Fetch attendance and permission data in bulk
    attendance_records = frappe.db.sql("""
        SELECT shift, employee_type, contractor, status, COUNT(name) as count 
        FROM `tabAttendance` 
        WHERE attendance_date = %(start_date)s AND docstatus != 2 
        GROUP BY shift, employee_type, contractor, status
    """, {"start_date": args['start_date']}, as_dict=True)

    permission_records = frappe.db.sql("""
        SELECT shift, employee_type, contractor, COUNT(name) as count
        FROM `tabPermission`
        WHERE permission_date = %(start_date)s AND status = 'Approved'
        GROUP BY shift, employee_type, contractor
    """, {"start_date": args['start_date']}, as_dict=True)

    # Convert to lookup dictionaries
    attendance_dict = {}
    permission_dict = {}

    for record in attendance_records:
        key = (record["shift"], record["employee_type"], record.get("contractor"), record["status"])
        attendance_dict[key] = record["count"]

    for record in permission_records:
        key = (record["shift"], record["employee_type"], record.get("contractor"))
        permission_dict[key] = record["count"]

    # Process data
    for shift in shifts:
        s_name = shift['name']
        tot, tot_hd, tot_p = 0, 0, 0

        for employee_type in employee_types:
            e_type = employee_type['name']

            if e_type not in ["Contract Employee", "Director"]:
                present = attendance_dict.get((s_name, e_type, None, "Present"), 0)
                half_day = attendance_dict.get((s_name, e_type, None, "Half Day"), 0)
                permission = permission_dict.get((s_name, e_type, None), 0)

                tot += present
                tot_hd += half_day
                tot_p += permission

                row.append(present)

            elif e_type == "Contract Employee":
                contractors = frappe.db.get_list("Contractor", fields=["name"], order_by="name ASC")

                for contractor in contractors:
                    c_name = contractor['name']

                    present = attendance_dict.get((s_name, e_type, c_name, "Present"), 0)
                    half_day = attendance_dict.get((s_name, e_type, c_name, "Half Day"), 0)
                    permission = permission_dict.get((s_name, e_type, c_name), 0)

                    tot += present
                    tot_hd += half_day
                    tot_p += permission

                    row.append(present)

        row.append(tot)
        row.append(tot_hd)
        row.append(tot_p)

    status.append(row)
    return status


@frappe.whitelist()
def get_data_2(args):
    status = []
    
    # Pre-fetch static data
    shifts = frappe.db.sql("""
        SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC
    """, as_dict=True)

    employee_types = frappe.db.sql("""
        SELECT name FROM `tabEmployee Type` WHERE name NOT IN ('Director') ORDER BY `order` ASC
    """, as_dict=True)

    contractors = frappe.db.sql("""
        SELECT name FROM `tabContractor` ORDER BY `name` ASC
    """, as_dict=True)

    # Fetch all relevant attendance data in a single query
    attendance_data = frappe.db.sql("""
        SELECT department, shift, employee_type, contractor, overtime_hours, attendance_date
        FROM `tabAttendance`
        WHERE docstatus != 2 AND attendance_date BETWEEN %s AND %s
    """, (add_days(args['start_date'], -1), args['start_date']), as_dict=True)

    # Group attendance by keys for quick lookup
    from collections import defaultdict
    att_map = defaultdict(list)
    for att in attendance_data:
        key = (att['department'], att['shift'], att['employee_type'], att.get('contractor'))
        att_map[key].append(att['overtime_hours'])

    # Fetch departments and sub-departments
    departments = frappe.db.sql("""
        SELECT name FROM `tabDepartment` WHERE name != 'All Departments' AND is_group = 1 ORDER BY `name` ASC
    """, as_dict=True)

    sub_departments_map = defaultdict(list)
    sub_departments = frappe.db.sql("""
        SELECT name, parent_department FROM `tabDepartment`
        WHERE parent_department IS NOT NULL
    """, as_dict=True)

    for sd in sub_departments:
        sub_departments_map[sd['parent_department']].append(sd['name'])

    # Overtime thresholds by shift
    ot_thresholds = {
        "A": {"B": 15, "C": 8},
        "B": {"C": 16, "A": 8},
        "C": {"A": 15, "B": 7}
    }

    # Compute attendance data
    for department in departments:
        for sub_department in sub_departments_map.get(department['name'], []):
            row = [department['name'], sub_department]
            for shift in shifts:
                tot = 0
                
                for et in employee_types:
                    if et['name'] != "Contract Employee":
                        ot = sum(
                            sum(1 for ot_hr in att_map.get((sub_department, sft, et['name'], None), [])
                                if ot_hr >= threshold)
                            for sft, threshold in ot_thresholds[shift['name']].items()
                        )
                        tot += ot
                        row.append(ot if ot > 0 else '-')
                    else:
                        for contractor in contractors:
                            ot = sum(
                                sum(1 for ot_hr in att_map.get((sub_department, sft, et['name'], contractor['name']), [])
                                    if ot_hr >= threshold)
                                for sft, threshold in ot_thresholds[shift['name']].items()
                            )
                            tot += ot
                            row.append(ot if ot > 0 else '-')

                # Totals and placeholders
                row.append(tot)
                row.extend(['-', '-'])

            status.append(row)
    row = ["Total", ""]

    shifts = frappe.db.get_all("Shift Type", filters={"name": ["!=", "WW"]}, fields=["name"], order_by="name ASC")
    contractors = frappe.db.get_all("Contractor", fields=["name"], order_by="name ASC")
    attendance_lookup = {}
    for record in attendance_data:
        key = (record['shift'], record['employee_type'], record.get('contractor'))
        attendance_lookup.setdefault(key, []).append(record['overtime_hours'])

    # Define overtime conditions per shift
    ot_conditions = {
        "A": [("B", 15, -1), ("C", 8, -1)],
        "B": [("C", 16, -1), ("A", 8, 0)],
        "C": [("A", 15, 0), ("B", 7, 0)],
    }

    for shift in shifts:
        tot = 0
        shift_name = shift['name']
        conditions = ot_conditions.get(shift_name, [])

        for employee_type in employee_types:
            emp_type_name = employee_type['name']
            if emp_type_name != "Contract Employee":
                ot = sum(
                    sum(1 for hours in attendance_lookup.get((s, emp_type_name, None), []) if hours >= h)
                    for s, h, _ in conditions
                )
                tot += ot
                row.append(ot)
            else:
                # Process contractors
                for contractor in contractors:
                    contractor_name = contractor['name']
                    ot = sum(
                        sum(1 for hours in attendance_lookup.get((s, emp_type_name, contractor_name), []) if hours >= h)
                        for s, h, _ in conditions
                    )
                    tot += ot
                    row.append(ot)

        row.append(tot)
        row.append("-")
        row.append("-")

    status.append(row)
    return status

@frappe.whitelist()
def get_data_3(args):
    status = []
    departments = frappe.db.sql("""SELECT name FROM `tabDepartment` WHERE name != "All Departments" AND is_group = 1 ORDER BY `name` ASC""", as_dict=True)
    for department in departments:
        dept=department['name']
        sub_departments = frappe.db.sql("""SELECT name FROM `tabDepartment` WHERE `parent_department` = %(dept)s ORDER BY `name` ASC""",{"dept": dept}, as_dict=True)
        for sub_department in sub_departments:
            s_dept=sub_department['name']
            row = [dept, s_dept]
            shifts = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
            for shift in shifts:
                s_type=shift['name']
                tot = 0
                tot_hd=0
                tot_p=0
                ec = frappe.db.sql("""SELECT name FROM `tabEmployee Type` ORDER BY `order` ASC""", as_dict=True)
                for employee_type in ec:
                    etype=employee_type['name']
                    if etype != "Contract Employee" and etype != "Director":
                        c = frappe.db.count("Attendance", {
                            'attendance_date': (args['start_date']),
                            'docstatus': ('!=', '2'),
                            'shift': s_type,
                            'employee_type': etype,
                            'department': s_dept,
                            'status' : "Present"
                        }) or 0
                        tot += c
                        hd_count = frappe.db.count("Attendance", {
                            'attendance_date': (args['start_date']),
                            'docstatus': ('!=', '2'),
                            'shift': s_type,
                            'employee_type': etype,
                            'department': s_dept,
                            'status' : "Half Day"
                        }) or 0
                        tot_hd += hd_count
                        perm = frappe.db.count("Permission", {
                            'permission_date': (args['start_date']),
                            'docstatus': 1,
                            'shift': s_type,
                            'employee_type': etype,
                            'department': s_dept,
                            # 'status' : "Half Day"
                        }) or 0
                        tot_p += perm
                        
                        if s_type == "A":
                            b_shift_ot = frappe.db.count("Attendance", {
                                'attendance_date': add_days((args['start_date']),-1),
                                'docstatus': ('!=', '2'),
                                'shift': "B",
                                'employee_type': etype,
                                'department': s_dept,
                                'overtime_hours': ('>=', 15)
                            }) or 0
                            c_shift_ot = frappe.db.count("Attendance", {
                                'attendance_date': add_days((args['start_date']),-1),
                                'docstatus': ('!=', '2'),
                                'shift': "C",
                                'employee_type': etype,
                                'department': s_dept,
                                'overtime_hours': ('>=', 8)
                            }) or 0
                            ot = b_shift_ot + c_shift_ot
                        if s_type == "B":
                            c_shift_ot = frappe.db.count("Attendance", {
                                'attendance_date': add_days((args['start_date']),-1),
                                'docstatus': ('!=', '2'),
                                'shift': "C",
                                'employee_type': etype,
                                'department': s_dept,
                                'overtime_hours': ('>=', 16)
                            }) or 0
                            a_shift_ot = frappe.db.count("Attendance", {
                                'attendance_date': args['start_date'],
                                'docstatus': ('!=', '2'),
                                'shift': "A",
                                'employee_type': etype,
                                'department': s_dept,
                                'overtime_hours': ('>=', 8)
                            }) or 0
                            ot = a_shift_ot + c_shift_ot
                        if s_type == "C":
                            a_shift_ot = frappe.db.count("Attendance", {
                                'attendance_date': args['start_date'],
                                'docstatus': ('!=', '2'),
                                'shift': "A",
                                'employee_type': etype,
                                'department': s_dept,
                                'overtime_hours': ('>=', 15)
                            })
                            b_shift_ot = frappe.db.count("Attendance", {
                                'attendance_date': args['start_date'],
                                'docstatus': ('!=', '2'),
                                'shift': "B",
                                'employee_type': etype,
                                'department': s_dept,
                                'overtime_hours': ('>=', 7)
                            })
                            ot = a_shift_ot + b_shift_ot
                        tot += ot
                        if (c + ot) > 0 :
                            row.append((c + ot))
                        else:
                            row.append('-')
                        # row.append(tot_hd)
                    elif etype == "Contract Employee":
                        contractors = frappe.db.sql("""SELECT name FROM `tabContractor` ORDER BY `name` ASC""", as_dict=True)
                        for contractor in contractors:
                            c_type=contractor['name']
                            c = frappe.db.count("Attendance", {
                                'attendance_date': (args['start_date']),
                                'docstatus': ('!=', '2'),
                                'shift': s_type,
                                'employee_type': etype,
                                'department': s_dept,
                                'status' : "Present",
                                'contractor': c_type
                            }) or 0
                            tot += c
                            hd_count = frappe.db.count("Attendance", {
                                'attendance_date': (args['start_date']),
                                'docstatus': ('!=', '2'),
                                'shift': s_type,
                                'employee_type': etype,
                                'department': s_dept,
                                'status' : "Half Day",
                                'contractor': c_type
                            }) or 0
                            tot_hd += hd_count
                            perm=frappe.db.count("Permission", {
                            'permission_date': (args['start_date']),
                            'status': 'Approved',
                            'department': s_dept,
                            'shift': s_type,
                            'employee_type': etype,
                            # 'contractor': c_type,
                
                                }) or 0
                            tot_p += perm
                            if s_type == "A":
                                b_shift_ot = frappe.db.count("Attendance", {
                                    'attendance_date': add_days((args['start_date']),-1),
                                    'docstatus': ('!=', '2'),
                                    'shift': "B",
                                    'employee_type': etype,
                                    'department': s_dept,
                                    'contractor': c_type,
                                    'overtime_hours': ('>=', 15)
                                }) or 0
                                c_shift_ot = frappe.db.count("Attendance", {
                                    'attendance_date': add_days((args['start_date']),-1),
                                    'docstatus': ('!=', '2'),
                                    'shift': "C",
                                    'employee_type': etype,
                                    'department': s_dept,
                                    'contractor': c_type,
                                    'overtime_hours': ('>=', 8)
                                }) or 0
                                ot = b_shift_ot + c_shift_ot
                            if s_type == "B":
                                c_shift_ot = frappe.db.count("Attendance", {
                                    'attendance_date': add_days((args['start_date']),-1),
                                    'docstatus': ('!=', '2'),
                                    'shift': "C",
                                    'employee_type': etype,
                                    'department': s_dept,
                                    'contractor': c_type,
                                    'overtime_hours': ('>=', 16)
                                }) or 0
                                a_shift_ot = frappe.db.count("Attendance", {
                                    'attendance_date': args['start_date'],
                                    'docstatus': ('!=', '2'),
                                    'shift': "A",
                                    'employee_type': etype,
                                    'department': s_dept,
                                    'contractor': c_type,
                                    'overtime_hours': ('>=', 8)
                                }) or 0
                                ot = a_shift_ot + c_shift_ot
                            if s_type == "C":
                                a_shift_ot = frappe.db.count("Attendance", {
                                    'attendance_date': args['start_date'],
                                    'docstatus': ('!=', '2'),
                                    'shift': "A",
                                    'employee_type': etype,
                                    'department': s_dept,
                                    'contractor': c_type,
                                    'overtime_hours': ('>=', 15)
                                }) or 0
                                b_shift_ot = frappe.db.count("Attendance", {
                                    'attendance_date': args['start_date'],
                                    'docstatus': ('!=', '2'),
                                    'shift': "B",
                                    'employee_type': etype,
                                    'department': s_dept,
                                    'contractor': c_type,
                                    'overtime_hours': ('>=', 7)
                                }) or 0
                                ot = a_shift_ot + b_shift_ot
                            tot += (ot + c)
                            if (ot + c) > 0 :
                                row.append((ot + c))
                            else:
                                row.append('-')	
                            # row.append(tot_hd)	
                row.append(tot)
                row.append(tot_hd)
                row.append(tot_p)
            status.append(row)
    row = ["Total",'']
    shifts = frappe.db.sql("""SELECT name FROM `tabShift Type` WHERE `name` != 'WW' ORDER BY `name` ASC""", as_dict=True)
    for shift in shifts:
        s_name=shift['name']
        tot = 0
        tot_hd=0
        ec = frappe.db.sql("""SELECT name FROM `tabEmployee Type` ORDER BY `order` ASC""", as_dict=True)
        for employee_type in ec:
            etype=employee_type['name']
            if etype != "Contract Employee" and etype != "Director":
                c = frappe.db.count("Attendance", {
                    'attendance_date': (args['start_date']),
                    'docstatus': ('!=', '2'),
                    'shift': s_name,
                    'employee_type': etype,
                    'status' : "Present"
                }) or 0
                tot += c
                hd_count = frappe.db.count("Attendance", {
                    'attendance_date': (args['start_date']),
                    'docstatus': ('!=', '2'),
                    'shift': s_name,
                    'employee_type': etype,
                    'status' : "Half Day"
                }) or 0
                tot_hd += hd_count
                perm=frappe.db.count("Permission", {
                            'permission_date': (args['start_date']),
                            'status': 'Approved',
                            # 'department': sub_department.name,
                            'shift': s_name,
                            'employee_type': etype,
                            # 'contractor': contractor.name,
                
                        }) or 0
                tot_p += perm
                if s_name == "A":
                    b_shift_ot = frappe.db.count("Attendance", {
                        'attendance_date': add_days((args['start_date']),-1),
                        'docstatus': ('!=', '2'),
                        'shift': "B",
                        'employee_type': etype,
                        'overtime_hours': ('>=', 15)
                    }) or 0
                    c_shift_ot = frappe.db.count("Attendance", {
                        'attendance_date': add_days((args['start_date']),-1),
                        'docstatus': ('!=', '2'),
                        'shift': "C",
                        'employee_type': etype,
                        'overtime_hours': ('>=', 8)
                    }) or 0
                    ot = b_shift_ot + c_shift_ot
                if s_name == "B":
                    c_shift_ot = frappe.db.count("Attendance", {
                        'attendance_date': add_days((args['start_date']),-1),
                        'docstatus': ('!=', '2'),
                        'shift': "C",
                        'employee_type': etype,
                        'overtime_hours': ('>=', 16)
                    }) or 0
                    a_shift_ot = frappe.db.count("Attendance", {
                        'attendance_date': args['start_date'],
                        'docstatus': ('!=', '2'),
                        'shift': "A",
                        'employee_type': etype,
                        'overtime_hours': ('>=', 8)
                    }) or 0
                    ot = a_shift_ot + c_shift_ot
                if s_name == "C":
                    a_shift_ot = frappe.db.count("Attendance", {
                        'attendance_date': args['start_date'],
                        'docstatus': ('!=', '2'),
                        'shift': "A",
                        'employee_type': etype,
                        'overtime_hours': ('>=', 15)
                    })
                    b_shift_ot = frappe.db.count("Attendance", {
                        'attendance_date': args['start_date'],
                        'docstatus': ('!=', '2'),
                        'shift': "B",
                        'employee_type': etype,
                        'overtime_hours': ('>=', 7)
                    })
                    ot = a_shift_ot + b_shift_ot
                tot += ot
                if (c + ot) > 0 :
                    row.append((c + ot))
                else:
                    row.append('-')
                # row.append(tot_hd)
            elif etype == "Contract Employee":
                contractors = frappe.db.sql("""SELECT name FROM `tabContractor` ORDER BY `name` ASC""", as_dict=True)
                for contractor in contractors:
                    c_name=contractor['name']
                    c = frappe.db.count("Attendance", {
                        'attendance_date': (args['start_date']),
                        'docstatus': ('!=', '2'),
                        'shift': s_name,
                        'employee_type': etype,
                        'status' : "Present",
                        'contractor': c_name
                    }) or 0
                    tot += c
                    hd_count = frappe.db.count("Attendance", {
                        'attendance_date': (args['start_date']),
                        'docstatus': ('!=', '2'),
                        'shift': s_name,
                        'employee_type': etype,
                        'status' : "Half Day",
                        'contractor': c_name
                    }) or 0
                    tot_hd += hd_count
                    perm=frappe.db.count("Permission", {
                            'permission_date': (args['start_date']),
                            'status': 'Approved',
                            # 'department': sub_department.name,
                            'shift': s_name,
                            'employee_type': etype,
                            'contractor': c_name,
                
                        }) or 0
                    tot_p += perm
                    if s_name == "A":
                        b_shift_ot = frappe.db.count("Attendance", {
                            'attendance_date': add_days((args['start_date']),-1),
                            'docstatus': ('!=', '2'),
                            'shift': "B",
                            'employee_type': etype,
                            'contractor': c_name,
                            'overtime_hours': ('>=', 15)
                        }) or 0
                        c_shift_ot = frappe.db.count("Attendance", {
                            'attendance_date': add_days((args['start_date']),-1),
                            'docstatus': ('!=', '2'),
                            'shift': "C",
                            'employee_type': etype,
                            'contractor': c_name,
                            'overtime_hours': ('>=', 8)
                        }) or 0
                        ot = b_shift_ot + c_shift_ot
                    if s_name == "B":
                        c_shift_ot = frappe.db.count("Attendance", {
                            'attendance_date': add_days((args['start_date']),-1),
                            'docstatus': ('!=', '2'),
                            'shift': "C",
                            'employee_type': etype,
                            'contractor': c_name,
                            'overtime_hours': ('>=', 16)
                        }) or 0
                        a_shift_ot = frappe.db.count("Attendance", {
                            'attendance_date': args['start_date'],
                            'docstatus': ('!=', '2'),
                            'shift': "A",
                            'employee_type': etype,
                            'contractor': c_name,
                            'overtime_hours': ('>=', 8)
                        }) or 0
                        ot = a_shift_ot + c_shift_ot
                    if s_name == "C":
                        a_shift_ot = frappe.db.count("Attendance", {
                            'attendance_date': args['start_date'],
                            'docstatus': ('!=', '2'),
                            'shift': "A",
                            'employee_type': etype,
                            'contractor': c_name,
                            'overtime_hours': ('>=', 15)
                        }) or 0
                        b_shift_ot = frappe.db.count("Attendance", {
                            'attendance_date': args['start_date'],
                            'docstatus': ('!=', '2'),
                            'shift': "B",
                            'employee_type': etype,
                            'contractor': c_name,
                            'overtime_hours': ('>=', 7)
                        }) or 0
                        ot = a_shift_ot + b_shift_ot
                    tot += (ot + c)
                    if (ot + c) > 0 :
                        row.append((ot + c))
                    else:
                        row.append('-')
                            
        row.append(tot)
        row.append(tot_hd)
        row.append(tot_p)
    status.append(row)
    return status 

