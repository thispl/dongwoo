import calendar
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from io import BytesIO
import frappe
from frappe.utils import nowdate
from openpyxl.utils import get_column_letter
from frappe.utils import getdate
from frappe.utils.background_jobs import enqueue
from dongwoo.mark_attendance import check_nhfh,check_holiday,get_dates
from erpnext.setup.doctype.employee.employee import get_holiday_list_for_employee
from babel.numbers import format_decimal
from hrms.hr.utils import get_holiday_dates_for_employee


@frappe.whitelist()
def download(start_date,end_date,contractor):
    filename = 'CL Salary Register'
    args = {'start_date':start_date,'end_date':end_date,'contractor':contractor}
    frappe.msgprint("Report is generating in the background, kindly check after a few minutes on the same page.")
    enqueue(build_xlsx_response, queue='default', timeout=6000, event='build_xlsx_response',filename=filename,args=args)
    

def build_xlsx_response(filename,args):
    xlsx_file = make_xlsx(filename,args)
    ret = frappe.get_doc({
            "doctype": "File",
            "attached_to_name": ' ',
            "attached_to_doctype": 'Report Dashboard',
            "attached_to_field": 'attach',
            "file_name": filename + '.xlsx',
            "is_private": 0,
            "content": xlsx_file.getvalue(),
            "decode": False
        })
    ret.save(ignore_permissions=True)
    frappe.db.commit()
    attached_file = frappe.get_doc("File", ret.name)
    frappe.db.set_value('Report Dashboard',None,'attach',attached_file.file_url)

def preload_month_data(employees, start_date, end_date):
    attendance = frappe.db.sql("""
        SELECT employee, attendance_date, status, shift
        FROM `tabAttendance`
        WHERE attendance_date BETWEEN %s AND %s
        AND employee IN %s
        AND docstatus = 1
    """, (start_date, end_date, tuple(employees)), as_dict=1)

    shifts = frappe.db.sql("""
        SELECT employee, start_date, shift_type
        FROM `tabShift Assignment`
        WHERE start_date BETWEEN %s AND %s
        AND employee IN %s
    """, (start_date, end_date, tuple(employees)), as_dict=1)


    ots = frappe.db.sql("""
        SELECT employee,
            SUM(ot_hours) AS total_ot_hrs
        FROM `tabOvertime Request`
        WHERE ot_date BETWEEN %s AND %s
        AND docstatus = 1
        GROUP BY employee
    """, (start_date, end_date), as_dict=1)

    holidays = frappe.db.sql("""
        SELECT holiday_date
        FROM `tabHoliday`
        WHERE holiday_date BETWEEN %s AND %s
    """, (start_date, end_date), as_dict=1)

    return attendance, shifts, ots, holidays


def build_maps(attendance, shifts, holidays):
    att_map = {}
    shift_map = {}
    ot_map = {}
    holiday_map = {}

    for a in attendance:
        emp = a.employee
        date = a.attendance_date

        att_map.setdefault(emp, {})[date] = a

        if a.overtime_hours and a.overtime_hours > 0:
            ot_map.setdefault(emp, {})[date] = a.overtime_hours

    for s in shifts:
        shift_map.setdefault(s.employee, {})[s.date] = s.shift_type

    for h in holidays:
        holiday_map[h.employee] = holiday_map.get(h.employee, 0) + 1

    return att_map, shift_map, ot_map, holiday_map

def get_shift_plan(emp_id, start_date, end_date, shift_type):
    return frappe.db.sql("""
        SELECT COUNT(*)
        FROM `tabShift Assignment`
        WHERE employee = %s
          AND start_date BETWEEN %s AND %s
          AND TRIM(UPPER(shift_type)) = %s
          AND docstatus = 1
    """, (emp_id, start_date, end_date, shift_type.upper()))[0][0]


def get_shift_data(emp_id, start_date, end_date):
    dates = get_dates_between(start_date, end_date)
    data = []
    holiday_list = frappe.db.get_value("Employee", emp_id, "holiday_list")
    shift_schedule = frappe.db.sql(
        """
        SELECT start_date, shift_type
        FROM `tabShift Assignment`
        WHERE docstatus = 1
          AND employee = %s
          AND start_date BETWEEN %s AND %s
        ORDER BY start_date
        """,
        (emp_id, start_date, end_date),
        as_dict=1
    )
    
    # shift_map = {shift["start_date"].strftime('%Y-%m-%d'): shift["shift_type"] for shift in shift_schedule}
    shift_map = {shift["start_date"]: shift["shift_type"] for shift in shift_schedule}

    
    for date in dates:
        if date in shift_map:
            if shift_map[date] == 'A':
                s = 1
            elif shift_map[date] == 'B':
                s = 2
            elif shift_map[date] == 'C':
                s = 3
            elif shift_map[date] == 'WW':
                s = 4
            else:
                s = "-"
        else:
            is_holiday = frappe.db.exists("Holiday", {
                "holiday_date": date,
                "parent": holiday_list
            })
            s = 4 if is_holiday else "-"
        data.append(s)

    return data


# def get_shif_attend(emp, dates, att_map, shift):
#     count = 0

#     for d in dates:
#         day = att_map.get(emp, {}).get(d)

#         if not day:
#             continue

#         if day["status"] == "Present" and day["shift"] == shift:
#             count += 1

#     return count

def get_shif_attend(emp, dates, att_map, shift, holiday_list):
    count = 0

    for d in dates:
        holiday = frappe.db.get_value(
            "Holiday",
            {
                "holiday_date": d,
                "parent": holiday_list
            },
            ["national_holiday", "festival_holiday"],
            as_dict=True
        )
        if holiday:
            continue
        day = att_map.get(emp, {}).get(d)
        if not day:
            continue
        if day["status"] == "Present" and day["shift"] == shift:
            count += 1

    return count


def get_actual_attend(emp_id, start_date, end_date):
    dates = get_dates_between(start_date, end_date)
    data = []

    att_data = frappe.db.sql(
        """
        SELECT 
            attendance_date, 
            status, 
            shift
        FROM `tabAttendance`
        WHERE docstatus = 1
          AND employee = %s 
          AND attendance_date BETWEEN %s AND %s
        """,
        (emp_id, start_date, end_date),
        as_dict=True
    )

    att_map = {
        att["attendance_date"]: {
            "status": att["status"],
            "shift": att["shift"]
        }
        for att in att_data
    }

    holiday_list = frappe.db.get_value("Employee", emp_id, "holiday_list")

    for d in dates:
        day = att_map.get(d)

        if not day:
            data.append("-")
            continue

        status = day["status"]
        shift = day["shift"]
        holiday = frappe.db.get_value(
            "Holiday",
            {
                "holiday_date": d,
                "parent": holiday_list
            },
            ["national_holiday", "festival_holiday"],
            as_dict=True
        )
        if status == "Present":
            if holiday:
                if holiday.national_holiday:
                    data.append("NH")
                elif holiday.festival_holiday:
                    data.append("FH")
                else:
                    data.append("")  
                continue

            elif frappe.db.exists(
                "Shift Assignment",
                {
                    "start_date": d,
                    "employee": emp_id,
                    "shift_type": "WW"
                }
            ):
                data.append("H")
                continue

            data.append(shift or "")

        elif status == "Half Day":
            data.append("H" + (shift or ""))

        elif status == "On Leave":
            data.append("O" + (shift or ""))

        elif status == "Absent":
            if frappe.db.exists("Holiday", {"holiday_date": d, "parent": holiday_list}) or \
               frappe.db.exists("Shift Assignment", {"start_date": d, "employee": emp_id, "shift_type": "WW"}):
                data.append("H")
            else:
                data.append(" ")

        else:
            data.append("")

    return data

def get_ot_map(emp, start_date, end_date):
    ot_data = frappe.db.get_all(
        "Attendance",
        filters={
            "employee": emp,
            "attendance_date": ["between", [start_date, end_date]],
            "docstatus": 1
        },
        fields=["attendance_date", "overtime_hours"]
    )

    return {
        d.attendance_date: d.overtime_hours
        for d in ot_data
        if d.overtime_hours and d.overtime_hours > 0
    }


def get_continue_data(emp_id, start_date, end_date):
    dates = get_dates_between(start_date, end_date)
    data = []

    ot_data = frappe.db.sql(
        """
        SELECT ot_date, total_hours
        FROM `tabOvertime Request`
        WHERE employee = %s
          AND ot_date BETWEEN %s AND %s
          AND workflow_state = 'Approved'
          AND docstatus = 1
        """,
        (emp_id, start_date, end_date),
        as_dict=True
    )

    ot_map = {
        ot["ot_date"]: ot["total_hours"]
        for ot in ot_data
        if ot.get("total_hours")
    }

    att_data = frappe.db.sql(
        """
        SELECT attendance_date, overtime_hours
        FROM `tabAttendance`
        WHERE employee = %s
          AND attendance_date BETWEEN %s AND %s
          AND docstatus != 2
        """,
        (emp_id, start_date, end_date),
        as_dict=True
    )

    att_map = {
        att["attendance_date"]: att["overtime_hours"]
        for att in att_data
        if att.get("overtime_hours")
    }

    for d in dates:
        if d in ot_map:
            data.append(ot_map[d])
        elif d in att_map:
            data.append(att_map[d])
        else:
            data.append("")

    return data

def get_weekoff_and_holiday_count(emp_id, start_date, end_date):
    dates = set()

    # Week Off dates
    ww_dates = frappe.db.sql("""
        SELECT start_date
        FROM `tabShift Assignment`
        WHERE employee = %s
          AND start_date BETWEEN %s AND %s
          AND TRIM(UPPER(shift_type)) = 'WW'
    """, (emp_id, start_date, end_date), as_dict=True)

    for d in ww_dates:
        dates.add(d.start_date)

    # Holiday dates
    holiday_dates = frappe.db.sql("""
        SELECT h.holiday_date
        FROM `tabHoliday` h
        JOIN `tabHoliday List` hl ON h.parent = hl.name
        JOIN `tabEmployee` e ON e.holiday_list = hl.name
        WHERE e.name = %s
          AND h.holiday_date BETWEEN %s AND %s
    """, (emp_id, start_date, end_date), as_dict=True)

    for d in holiday_dates:
        dates.add(d.holiday_date)

    return len(dates)



def get_nal_from_map(emp, att_map):
    nal = 0

    for a in att_map.get(emp, {}).values():

        if a.shift and a.shift.strip().upper() == "WW":
            continue

        if a.status == "Absent":
            nal += 1

        elif a.status == "Half Day" and not a.leave_application:
            nal += 0.5

    return nal


def get_al_from_map(emp, att_map):
    return sum(1 for a in att_map.get(emp, {}).values() if a.status == "On Leave")

def get_half_day_from_map(emp, att_map):
    return sum(1 for a in att_map.get(emp, {}).values() if a.status == "Half Day")

def get_cont_a_from_map(emp, ot_map):
    return ot_map.get(emp, {}).get("cont_a", 0)


def get_dates_between(start_date, end_date):
    days = []
    cur = start_date
    while cur <= end_date:
        days.append(cur)
        cur += timedelta(days=1)
    return days


def make_xlsx(data,args, sheet_name=None, wb=None, column_widths=None):
    if wb is None:
        wb = Workbook()
    ws = wb.active
    start_date = getdate(args['start_date'])
    end_date = getdate(args['end_date'])
    contractor=args['contractor']
    given_year = start_date.year
    given_month = start_date.month
    total_days = calendar.monthrange(given_year, given_month)[1]
    
    blue_color = PatternFill(start_color="95b3d7", end_color="95b3d7", fill_type="solid")
    red_color = PatternFill(start_color="ff5050", end_color="ff5050", fill_type="solid")
    orange_color = PatternFill(start_color="e26b0a", end_color="e26b0a", fill_type="solid")
    teal_color = PatternFill(start_color="d8e4bc", end_color="d8e4bc", fill_type="solid")
    green_color = PatternFill(start_color="92d050", end_color="92d050", fill_type="solid")
    yellow_color = PatternFill(start_color="ffff00", end_color="ffff00", fill_type="solid")
    light_orange = PatternFill(start_color="ffeb9c", end_color="ffeb9c", fill_type="solid")
    sunday_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    orange_color = PatternFill(start_color="ffc000", end_color="ffc000", fill_type="solid")
    sky_blue = PatternFill(start_color="92cddc", end_color="92cddc", fill_type="solid")
    light_blue = PatternFill(start_color="b7dee8", end_color="b7dee8", fill_type="solid")
    purple = PatternFill(start_color="ccc0da", end_color="ccc0da", fill_type="solid")
    light_green = PatternFill(start_color="57ff57", end_color="57ff57", fill_type="solid")
    dark_green = PatternFill(start_color="00b050", end_color="00b050", fill_type="solid")
    tomato = PatternFill(start_color="ff6d4b", end_color="ff6d4b", fill_type="solid")
    bright_blue = PatternFill(start_color="00b0f0", end_color="00b0f0", fill_type="solid")
    skin = PatternFill(start_color="fcd5b4", end_color="fcd5b4", fill_type="solid")
    actual_green = PatternFill(start_color="99ff99", end_color="99ff99", fill_type="solid")
    light_pink = PatternFill(start_color="fde9d9", end_color="fde9d9", fill_type="solid")
    week_off_pink = PatternFill(start_color="ffc7ce", end_color="ffc7ce", fill_type="solid")
    schedule_green = PatternFill(start_color="7fb957", end_color="7fb957", fill_type="solid")
    actual_header_green = PatternFill(start_color="99ff66", end_color="99ff66", fill_type="solid")
    continue_green = PatternFill(start_color="e26b0a", end_color="e26b0a", fill_type="solid")
    onleave_color = PatternFill(start_color="ffc000", end_color="ffc000", fill_type="solid")
    holiday_color = PatternFill(start_color="8db4e2", end_color="8db4e2", fill_type="solid")
 
    red_font = Font(color="FF0000")
    dark_red = Font(color="9c0000")
 
    font = Font(bold=True, color="000000")
    align_center_center = Alignment(vertical="center", horizontal="center", wrap_text=True)
    alignment = Alignment(horizontal="center")
    alignment_right = Alignment(horizontal="right")
    alignment_left = Alignment(horizontal="left")
    wrap_text = Alignment(wrap_text=True)
    title_font = Font(bold=True, size=14)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    width_1 = [5, 10, 15, 15, 20, 18, 18, 18, 18] 

    if total_days == 30:
        day_col = [5] * 30
    elif total_days == 31:
        day_col = [5] * 31
    elif total_days == 29:
        day_col = [5] * 29
    elif total_days == 28:
        day_col = [5] * 28
    width_2 = [10] * 15
    width_3 = [15] * 35
 
    column_widths = width_1 + day_col + width_2 + width_3  

    for i, width in enumerate(column_widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = width

    ws.column_dimensions["D"].hidden = True
    ws.column_dimensions["G"].hidden = True
    ws.column_dimensions["H"].hidden = True
    ws["A1"].value = f"{contractor} DAILY ATTENDANCE CUM SALARY STATEMENT"
    ws["A4"].value = f"FOR THE MONTH OF {start_date.strftime('%B')} - {start_date.year}"
    ws["J2"].fill = blue_color
    ws["J3"].fill = red_color
    ws["J4"].fill = orange_color
    ws["J5"].fill = teal_color
    ws["K2"].value = "COMP OFF"
    ws["K3"].value = "WEEKLY OFF"
    ws["K4"].value = "HANI SCOPE OT"
    ws["K5"].value = "DWSI SCOPE OT"
    ws["N2"].value = "HA"
    ws["N3"].value = "HB"
    ws["N4"].value = "HC"
    ws["N5"].value = ""
    ws["O2"].value = "HALF DAY - A SHIFT"
    ws["O3"].value = "HALF DAY - B SHIFT"
    ws["O4"].value = "HALF DAY - C SHIFT"
    ws["O5"].value = ""
    ws["U1"].value = "Schedule"
    ws["U2"].value = "Worked"
    ws["U3"].value = "Continue"
    ws["U4"].value = "NAL"
    ws["U5"].value = "AL"
    ws["W1"].value = "1: A Shift, 2: B Shift, 4 : W off"
    ws["W2"].value = "1: A Shift, 2: B Shift, 4 : W off"
    ws["W3"].value = "16-'A' & 'B', 14 - 'C' Shift"
    ws["W4"].value = "XA & XB & XC :Non approver Leave & Shift"
    ws["W5"].value = "OA & OB & OC :Approver Leave & Shift"
    ws["AE1"].value = "BF"
    ws["AE2"].value = "CF"
    ws["AE3"].value = "MF"
    ws["AE4"].value = "H2"
    ws["AF1"].value = "BCF"
    ws["AF2"].value = "CCF"
    ws["AF3"].value = "MF"
    ws["AF4"].value = "H2"
    ws["AH1"].value = "BCF"
    ws["AH2"].value = "CCF"
    ws["AH3"].value = "MBF"
    ws["AH4"].value = "OC"
    ws["AH4"].fill = light_orange
    ws["AH4"].font = Font(color="885100")
    ws["AI1"].value = "PLANT 1"
    ws["AI2"].value = "DISBATCH"
    ws["AI3"].value = "RECIVING"
    ws["AI4"].value = "OTHER CUSTOMER"
    ws["AK1"].value = "FLT"
    ws["AK2"].value = "BC"
    ws["AK3"].value = "UL"
    ws["AK4"].value = "WD"
    ws["AM1"].value = "FORK LIFT"
    ws["AM2"].value = "BIN CLEAN"
    ws["AM3"].value = "Unloading"
    ws["AM4"].value = "Welder"
    ws["AN1"].value = "DOC"
    ws["AN2"].value = "MN"
    ws["AN3"].value = "FM"
    ws["BP4"].font = font
    ws["BP5"].font = font


    contractor = args.get("contractor")
    service_charge = 0

    if contractor:
        service_charge = frappe.db.get_value(
            "Contractor",
            contractor,
            "service_charges"
        ) or 0

    row = [""]
    ws.append(row)  

    row = ["S.NO", "DEPT", "NEW EMP.ID", "EMP ID", "NAME", "DOJ (YYYY-MM-DD)", "DOB (YYYY-MM-DD)", "Category", ""]
    row += [f"{day}" for day in range(1, total_days + 1)]
    row += ["Schedule", "", "", "", "", "", "Actual Attend", "", "", "", "", "", "Leave", "", ""]
    row += ["Total", "Cont (A)", "TOTAL OT HOURS", "EXISTING GROSS 03-2024", "Increase", "REVISED GROSS 04-2024", "OT RATE", "Shoe, Uniform", "Total", "Total Man Days", "Total Wage", "PF Restricted Wage Per Day", "PF Wage", "Actual PF 13%", "PF (on 60%) for commission", "PF for Service Charge (on 60%) Rev.", "BASIC+DA", "HRA", "WASHING ALLOWANCE", "Fixed Gross for ESI Calc.", "ESI WAGE", "ESI OT Wage", "ESI @ 3.25%", "OT HOURS", "Total OT Amount", "Atten. bonus", "Fest. Allow.", "Grand Total", f"Service Charge {service_charge}%", "DWSI OT", "Contractor OT", "Salary for Festival day Present", "Festival day Service Charges - 10%", "Not present but paid for Festival day", "Deduction of Service charge - 10%"]
    ws.append(row)
 
    row = [""] * 9 
    row += [calendar.day_abbr[datetime(given_year, given_month, day).weekday()] for day in range(1, total_days + 1)]
    row += ["A", "B", "C", "Wrk", "H", "Total", "A", "B", "C", "Wrk", "W'OFF", "Total", "NAL", "AL", "HALF DAY"]
    ws.append(row)
 
    for day in range(1, total_days + 1):
        weekday = datetime(given_year, given_month, day).weekday()
        column = day + 9

        if weekday == 6:
            ws.cell(row=7, column=column).fill = sunday_fill
            ws.cell(row=8, column=column).fill = sunday_fill
    
    ws.merge_cells("A7:A8")
    ws.merge_cells("B7:B8")
    ws.merge_cells("C7:C8")
    ws.merge_cells("D7:D8")
    ws.merge_cells("E7:E8")
    ws.merge_cells("F7:F8")
    ws.merge_cells("G7:G8")
    ws.merge_cells("H7:H8")
    ws.merge_cells("I7:I8")

    if total_days == 30:
        ws["BE7"].fill = orange_color
        ws["BF7"].fill = yellow_color
        ws["BH7"].fill = teal_color
        ws["BN7"].fill = sky_blue
        ws["BO7"].fill = sky_blue
        ws["BR7"].fill = light_blue
        ws["BS7"].fill = purple
        ws["BT7"].fill = purple
        ws["BU7"].fill = purple
        ws["BV7"].fill = purple
        ws["BW7"].fill = purple
        ws["BX7"].fill = purple
        ws["BY7"].fill = purple
        ws["CF7"].fill = light_green
        ws["CG7"].fill = tomato
        ws["CH7"].fill = dark_green
        ws["CI7"].fill = bright_blue
        ws["CJ7"].fill = skin
        ws["CK7"].fill = skin

        ws.merge_cells("AN7:AS7")
        ws.merge_cells("AT7:AY7")
        ws.merge_cells("AZ7:BB7")

    if total_days == 31:
        ws["BF7"].fill = orange_color
        ws["BG7"].fill = yellow_color
        ws["BI7"].fill = teal_color
        ws["BO7"].fill = sky_blue
        ws["BP7"].fill = sky_blue
        ws["BS7"].fill = light_blue
        ws["BT7"].fill = purple
        ws["BU7"].fill = purple
        ws["BV7"].fill = purple
        ws["BW7"].fill = purple
        ws["BX7"].fill = purple
        ws["BY7"].fill = purple
        ws["BZ7"].fill = purple
        ws["CG7"].fill = light_green
        ws["CH7"].fill = tomato
        ws["CI7"].fill = dark_green
        ws["CJ7"].fill = bright_blue
        ws["CK7"].fill = skin
        ws["CL7"].fill = skin

        ws.merge_cells("AO7:AT7")
        ws.merge_cells("AU7:AZ7")
        ws.merge_cells("BA7:BC7")

    ws["A7"].alignment = align_center_center
    ws["B7"].alignment = align_center_center
    ws["C7"].alignment = align_center_center
    ws["D7"].alignment = align_center_center
    ws["E7"].alignment = align_center_center
    ws["F7"].alignment = align_center_center
    ws["G7"].alignment = align_center_center
    ws["H7"].alignment = align_center_center
    
    
    if contractor:
        department = frappe.db.sql("""select distinct department from `tabSalary Slip`
                                        where salary_structure = 'Contractor SS' and 
                                        start_date = '%s' and end_date = '%s' and contractor = '%s'"""%(start_date, end_date, contractor), as_dict=1)
    else:
        department = frappe.db.sql("""select distinct department from `tabSalary Slip`
                                        where salary_structure = 'Contractor SS' and 
                                        start_date = '%s' and end_date = '%s'"""%(start_date, end_date), as_dict=1)
    start_after_days = row_start = total_row_start = total_start_after_days = 9
    end_after_days = row_end = total_row_end = total_end_after_days = row_start + 3

    grand_shift_a_plan = grand_shift_b_plan = grand_shift_c_plan = grand_week_off = grand_work_plan = grand_holidays_plan = grand_total_plan = grand_shift_a_attend = grand_shift_b_attend = grand_shift_c_attend = grand_work_attend = grand_week_off_attend = grand_total_attend = grand_nal = grand_al = grand_half_day_count = grand_total_after_attend = grand_cont_a = 0
    grand_ot_hours = grand_existing_gross = grand_increase = grand_revised_gross = grand_ot_rate = grand_shoe_uniform = grand_total_ = grand_total_mandays = grand_total_wage = grand_pf_restricted_per_day = grand_pf_wage = grand_actual_pf_13 = grand_pf_60 = grand_pf_service_charge = grand_basic = grand_hra = grand_wa = grand_fixed_gross_esi = grand_esi_wage = grand_esi_ot_wage = grand_esi_3_25 = grand_ot_hours = grand_total_ot_amount = grand_att_bonus = grand_fest_allowance = grand_grand_total = grand_service_charge_10 = grand_dwsi_ot = grand_hj_ot = grand_salary_for_fest = grand_fest_day_service_charge = grand_not_present_paid = grand_deduction_secvice_charge = 0
    
    fmd_shift_a_plan = fmd_shift_b_plan = fmd_shift_c_plan = fmd_week_off = fmd_work_plan = fmd_holidays_plan = fmd_total_plan = fmd_shift_a_attend = fmd_shift_b_attend = fmd_shift_c_attend = fmd_work_attend = fmd_week_off_attend = fmd_total_attend = fmd_nal = fmd_al = fmd_half_day_count = fmd_total_after_attend = fmd_cont_a = 0
    fmd_ot_hours = fmd_existing_gross = fmd_increase = fmd_revised_gross = fmd_ot_rate = fmd_shoe_uniform = fmd_total_ = fmd_total_mandays = fmd_total_wage = fmd_pf_restricted_per_day = fmd_pf_wage = fmd_actual_pf_13 = fmd_pf_60 = fmd_pf_service_charge = fmd_basic = fmd_hra = fmd_wa = fmd_fixed_gross_esi = fmd_esi_wage = fmd_esi_ot_wage = fmd_esi_3_25 = fmd_ot_hours = fmd_total_ot_amount = fmd_att_bonus = fmd_fest_allowance = fmd_grand_total = fmd_service_charge_10 = fmd_dwsi_ot = fmd_hj_ot = fmd_salary_for_fest = fmd_fest_day_service_charge = fmd_not_present_paid = fmd_deduction_secvice_charge = 0
    
    ht_shift_a_plan = ht_shift_b_plan = ht_shift_c_plan = ht_week_off = ht_work_plan = ht_holidays_plan = ht_total_plan = ht_shift_a_attend = ht_shift_b_attend = ht_shift_c_attend = ht_work_attend = ht_week_off_attend = ht_total_attend = ht_nal = ht_al = ht_half_day_count = ht_total_after_attend = ht_cont_a = 0
    ht_ot_hours = ht_existing_gross = ht_increase = ht_revised_gross = ht_ot_rate = ht_shoe_uniform = ht_total_ = ht_total_mandays = ht_total_wage = ht_pf_restricted_per_day = ht_pf_wage = ht_actual_pf_13 = ht_pf_60 = ht_pf_service_charge = ht_basic = ht_hra = ht_wa = ht_fixed_gross_esi = ht_esi_wage = ht_esi_ot_wage = ht_esi_3_25 = ht_ot_hours = ht_total_ot_amount = ht_att_bonus = ht_fest_allowance = ht_grand_total = ht_service_charge_10 = ht_dwsi_ot = ht_hj_ot = ht_salary_for_fest = ht_fest_day_service_charge = ht_not_present_paid = ht_deduction_secvice_charge = 0
    
    
    for dept in department:
        if contractor:
            data = frappe.db.sql(
                """
                SELECT 
                    s.department, 
                    s.employee, 
                    s.employee_name, 
                    e.date_of_joining, 
                    e.date_of_birth, 
                    e.designation,
                    e.contractor AS con,
                    s.name as slip_id,
                    CAST(e.existing_gross AS SIGNED) AS existing_gross, 
                    CAST(e.increase_amount AS SIGNED) AS increase_amount, 
                    e.revised_gross AS revised_gross, 
                    CAST(e.shoe_uniform AS SIGNED) AS shoe_uniform, 
                    s.payment_days AS payment_days,
                    s.present_days_ AS present_days,
                    e.basic_da AS basic_da, 
                    e.hra_contractor_ AS house_rent_allowance, 
                    s.ot_hours AS ot_hours, 
                    s.start_date,s.end_date,
                    s.total_working_days,
                    CAST(s.att_bonus AS SIGNED) AS att_bonus,
                    CAST(s.festival_allowance AS SIGNED) AS festival_allowance
                    
                FROM `tabSalary Slip` s 
                INNER JOIN `tabEmployee` e ON e.name = s.employee
                INNER JOIN `tabAttendance` a ON a.employee = s.employee
                WHERE 
                    s.department = '%s' 
                    AND s.start_date = '%s' 
                    AND s.end_date = '%s' 
                    AND s.salary_structure = 'Contractor SS'
                    AND s.contractor = '%s'
                GROUP BY s.name
                """ % (dept.department, start_date, end_date,contractor),
                as_dict=True,
            )
        else:
            data = frappe.db.sql(
                """
                SELECT 
                    s.department, 
                    s.employee, 
                    s.employee_name, 
                    e.date_of_joining, 
                    e.date_of_birth, 
                    e.designation,
                    e.contractor AS con,
                    s.name as slip_id,
                    CAST(e.existing_gross AS SIGNED) AS existing_gross, 
                    CAST(e.increase_amount AS SIGNED) AS increase_amount, 
                    e.revised_gross AS revised_gross, 
                    CAST(e.shoe_uniform AS SIGNED) AS shoe_uniform, 
                    s.payment_days AS payment_days,
                    s.present_days_ AS present_days,
                    e.basic_da AS basic_da, 
                    e.hra_contractor_ AS house_rent_allowance, 
                    s.ot_hours AS ot_hours, 
                    s.start_date,s.end_date,
                    s.total_working_days,
                    CAST(s.att_bonus AS SIGNED) AS att_bonus,
                    CAST(s.festival_allowance AS SIGNED) AS festival_allowance
                    
                FROM `tabSalary Slip` s 
                INNER JOIN `tabEmployee` e ON e.name = s.employee
                INNER JOIN `tabAttendance` a ON a.employee = s.employee
                WHERE 
                    s.department = '%s' 
                    AND s.start_date = '%s' 
                    AND s.end_date = '%s' 
                    AND s.salary_structure = 'Contractor SS'
                GROUP BY s.name
                """ % (dept.department, start_date, end_date),
                as_dict=True,
            )

        ind = 1

        row_start = total_row_start
        start_after_days = total_start_after_days
        row_end = total_row_end 
        end_after_days = total_end_after_days
        
        total_shift_a_plan = total_shift_b_plan = total_shift_c_plan = total_week_off = total_holidays_plan = total_work_plan = total_total_plan = total_shift_a_attend = total_shift_b_attend = total_shift_c_attend = total_work_attend = total_week_off_attend = total_total_attend = total_nal = total_al = total_half_day_count = total_total_after_attend = total_cont_a = 0
        total_ot_hours = total_existing_gross = total_increase = total_revised_gross = total_ot_rate = total_shoe_uniform = total_total = total_total_mandays = total_total_wage = total_pf_restricted_per_day = total_pf_wage = total_actual_pf_13 = total_pf_60 = total_pf_service_charge = total_basic = total_hra = total_wa = total_fixed_gross_esi = total_esi_wage = total_esi_ot_wage = total_esi_3_25 = total_ot_hours = total_total_ot_amount = total_att_bonus = total_fest_allowance = total_grand_total = total_service_charge_10 = total_salary_for_fest = total_fest_day_service_charge = total_dwsi_ot = total_hj_ot = total_not_present_paid = total_deduction_service_charge = 0
        all_employees = frappe.db.sql("""
            SELECT DISTINCT s.employee
            FROM `tabSalary Slip` s
            WHERE s.salary_structure = 'Contractor SS'
            AND s.start_date = %s
            AND s.end_date = %s
        """, (start_date, end_date), as_list=1)

        employees = [e[0] for e in all_employees]

        attendance, shifts, ots, holidays = preload_month_data(employees,start_date,end_date)

        att_map, shift_map, ot_map, holiday_map = build_maps(attendance,shifts,holidays)

        dates = get_dates_between(start_date, end_date) 
        
        for i in data:
            emp_id = i.employee
            shift_a_plan = get_shift_plan(emp_id, start_date, end_date, 'A')
            shift_b_plan = get_shift_plan(emp_id, start_date, end_date, 'B')
            shift_c_plan = get_shift_plan(emp_id, start_date, end_date, 'C')    
            week_off     = get_shift_plan(emp_id, start_date, end_date, 'WW')
            
            holidays_plan = get_weekoff_and_holiday_count(
                emp_id,
                start_date,
                end_date
            )


            work_plan = shift_a_plan + shift_b_plan + shift_c_plan
            total_plan = work_plan + holidays_plan

            holiday_list = frappe.db.get_value("Employee", emp_id, "holiday_list")
            
            # ATTENDANCE
            shift_a_attend = get_shif_attend(emp_id, dates, att_map, 'A',holiday_list)
            shift_b_attend = get_shif_attend(emp_id, dates, att_map, 'B',holiday_list)
            shift_c_attend = get_shif_attend(emp_id, dates, att_map, 'C',holiday_list)

            work_attend = shift_a_attend + shift_b_attend + shift_c_attend
            week_off_attend = holidays_plan
            total_attend = work_attend + week_off_attend

            # LEAVE / ABSENT
            nal = get_nal_from_map(emp_id, att_map)
            al = get_al_from_map(emp_id, att_map)
            half_day_count = get_half_day_from_map(emp_id, att_map)

            if half_day_count >0:
                half_day_count=int(half_day_count)/2
            total_after_attend = total_attend + nal + al + half_day_count
            cont_a = get_cont_a_from_map(emp_id, ot_map)
            existing_gross = round((i.existing_gross), 2)
            increase = round((i.increase_amount), 2)
            revised_gross = round((i.revised_gross), 2)
            ot_rate = round((i.revised_gross / 8), 2)
            total_mandays = i.present_days
        
            working_days = i.total_working_days - holidays_plan
            pf = round(15000 / working_days, 0) if working_days > 0 else 0
            pf=round(pf)
            if revised_gross < pf:
                pf_restricted_per_day = revised_gross
            if revised_gross > pf:
                pf_restricted_per_day = pf
            
            pf_wage1 = round(pf_restricted_per_day * total_mandays, 0)

            if pf_wage1 > 15000:
                pf_wage = 15000
            else:
                pf_wage = pf_wage1

            actual_pf_13 = round(pf_wage * 0.13, 0)
            

            pf_60 = actual_pf_13 * 0.6 
            contractors = frappe.get_all(
                "Contractor",
                filters={"name": i.con},
                fields=[
                    "name",
                    "from_date",
                    "to_date",
                    "service_charges",
                    "shoe_amount",
                    "minimum_wages"
                ]
            )

            basic = 0.00 
            shoe_uniform = 0.00
            service_charges = 0.00 
            if contractors:
                c = contractors[0]   
                slip_start = i.start_date
                slip_end = i.end_date
                minimum_wages = float(c.minimum_wages or 0)
                if c.from_date <= slip_start and c.to_date >= slip_end:
                    basic = minimum_wages
                else:
                    basic = 0.00
                shoe_uniform = round(float(c.shoe_amount or 0), 2)
                service_charges = float(c.service_charges or 0)   

            total = round((i.revised_gross + shoe_uniform), 2)
            
            total_wage = round((total_mandays * total), 0)     
            if pf_60 < 1170:
                pf_service_charge = pf_60
            else:
                pf_service_charge = 1170
            hra = round(((revised_gross - basic) /2), 2)
            wa = round(((revised_gross - basic) /2), 2)
            fixed_gross_esi = round(revised_gross * 26)
            att_bonus = frappe.db.get_value("Salary Detail", {"parent": i.slip_id , "salary_component": "Attendance Bonus"},"amount") or 0.0
            att_bonus=float(att_bonus)
            if fixed_gross_esi < 21000:
                esi_wage = round(((basic+hra) * i.present_days),2)
                basic=float(basic)
                esi_ot_wage = round((((basic+hra) / 8) * float(i.ot_hours)),2)
                esi_3_25 = round(((int(esi_wage) + int(esi_ot_wage) + int(att_bonus))*0.0325),2)
            else:
                esi_wage = 0
                esi_ot_wage = 0
                esi_3_25 = 0
            if i.ot_hours:
                ot_hours = i.ot_hours  
            else:
                ot_hours = 0
            fest_allowance = i.festival_allowance
            revised_gross = float(revised_gross)
            ot_hours1 = float(ot_hours)
            # ot_hours = ot_hours1*2
            ot_hours = ot_hours1
            total_ot_amount = round(((revised_gross / 8) * ot_hours), 0) 
            grand_total = fest_allowance + att_bonus + total_ot_amount + actual_pf_13 + total_wage
            service_charge_10 = round(((total_wage + pf_service_charge) * (service_charges / 100)), 0)
            if fest_allowance > 0:
                salary_for_fest = revised_gross
                fest_day_service_charge = round((salary_for_fest * 0.1), 0)
            else:
                salary_for_fest = 0
                fest_day_service_charge = 0
            
            ot_data = calculate_ot_hours(i.employee, start_date, end_date)
            hj_ot = ot_data.get("contractor_ot_hrs", 0)
            dwsi_ot = ot_data.get("dwsi_ot_hrs", 0)
            dates = get_dates_between(start_date, end_date)
            holiday_list = frappe.db.get_value("Employee", emp_id, "holiday_list")

            ans = 0  
            if holiday_list:
                for d in dates:
                    holiday = frappe.db.get_value(
                        "Holiday",
                        {
                            "holiday_date": d,
                            "parent": holiday_list
                        },
                        ["national_holiday", "festival_holiday"],
                        as_dict=True
                    )
                    if holiday and (holiday.national_holiday or holiday.festival_holiday):
                        ans = 1
                        break 
            if ans == 1:  
                if fest_allowance == 0:
                    not_present_paid = revised_gross
                else:
                    not_present_paid = 0
            else:
                not_present_paid = 0
                        

            deduction_service_charge = round((not_present_paid * 0.1), 0)

            row1=[ind, i.department, i.employee, i.employee, i.employee_name, i.date_of_joining, i.date_of_birth, i.designation, "Schedule "] 
            shift_sch = get_shift_data(i.employee, start_date, end_date)
            for j in shift_sch:
                row1.append(j)
            row1 += [shift_a_plan, shift_b_plan, shift_c_plan, work_plan, holidays_plan, total_plan, shift_a_attend, shift_b_attend, shift_c_attend, work_attend, week_off_attend, total_attend, nal, al, half_day_count, total_after_attend, cont_a]
            row1 += [ot_hours, formatted(existing_gross), formatted(increase), formatted(revised_gross), ot_rate, formatted(shoe_uniform), formatted(total), total_mandays, formatted(total_wage), formatted(pf_restricted_per_day), formatted(pf_wage), formatted(actual_pf_13), formatted(pf_60), formatted(pf_service_charge), formatted(basic), formatted(hra), formatted(wa), formatted(fixed_gross_esi), formatted(esi_wage), formatted(esi_ot_wage), formatted(esi_3_25), ot_hours, formatted(total_ot_amount), formatted(att_bonus), formatted(fest_allowance), formatted(grand_total), formatted(service_charge_10), formatted(dwsi_ot), formatted(hj_ot), formatted(salary_for_fest), formatted(fest_day_service_charge), formatted(not_present_paid), formatted(deduction_service_charge)]
            ws.append(row1)
            row2 = ["", "", "", "", "", "", "", "", "Actual  "]
            actual_data = get_actual_attend(i.employee, start_date, end_date)

            for k in actual_data:
                row2.append(k)

            ws.append(row2)

            row_idx = ws.max_row      
            start_col = 10            

            for idx, k in enumerate(actual_data):
                col_idx = start_col + idx
                cell = ws.cell(row=row_idx, column=col_idx)

                if k == "H":
                    cell.fill = light_pink
            row3 = ["", "", "", "", "", "", "", "", "Continue "]
            fill_hj = PatternFill(start_color='E26B0A', end_color='E26B0A', fill_type='solid')   # HJ OT color
            fill_dwsi = PatternFill(start_color='E26B00', end_color='E26B00', fill_type='solid') # DWSI OT color


            continue_data = get_continue_data(emp_id, start_date, end_date)
            for val in continue_data:
                row3.append(val if val != 0 else "")

            ws.append(row3) 

            row_idx = ws.max_row
            col_idx = 10 


            ws.append(["", "", "", "", "", "", "", "", "Workspot"])


            for k in range(1, 9):
                ws.merge_cells(start_row=row_start, end_row=row_end, start_column=k, end_column=k)
                cell = ws.cell(row=row_start, column=k)
            
            ind += 1
            row_start += 4
            row_end += 4
            total_row_start = row_start + 1
            total_row_end = row_end + 1
            for l in range(total_days+10, 60+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                
            for l in range(total_days+16, 20+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=l).fill = actual_green

            for l in range(total_days+22, 24+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=l).fill = yellow_color
            
            for l in range(total_days+24, 25+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=l).fill = actual_green
     
            for l in range(total_days+27, 28+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=l).fill = yellow_color

            for l in range(total_days+28, 29+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=l).fill = light_pink
                    

            for l in range(total_days+30, 31+total_days):
                ws.merge_cells(start_row=start_after_days, end_row=end_after_days, start_column=l, end_column=l)
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=l).fill = teal_color

            from_date = start_date.strftime('%Y-%m-%d')
            to_date = end_date.strftime('%Y-%m-%d')
    
            for i in holiday_col_color(from_date, to_date):
                for row in range(start_after_days, end_after_days + 1):
                    ws.cell(row=row, column=i).fill = holiday_color
     
            start_after_days += 4
            end_after_days += 4
            total_start_after_days = start_after_days + 1
            total_end_after_days = end_after_days + 1

            total_shift_a_plan += shift_a_plan
            total_shift_b_plan += shift_b_plan
            total_shift_c_plan += shift_c_plan
            total_week_off += week_off
            total_holidays_plan += holidays_plan
            total_work_plan += work_plan
            total_total_plan += total_plan
            total_shift_a_attend += shift_a_attend
            total_shift_b_attend += shift_b_attend
            total_shift_c_attend += shift_c_attend
            total_work_attend += work_attend
            total_week_off_attend += week_off_attend
            total_total_attend += total_attend
            total_nal += nal
            total_al += al
            total_half_day_count += half_day_count
            total_total_after_attend += total_after_attend
            total_cont_a += cont_a

            total_ot_hours += ot_hours
            total_existing_gross += existing_gross
            total_increase += increase
            total_revised_gross += revised_gross
            total_ot_rate += ot_rate
            total_shoe_uniform += shoe_uniform
            total_total += total
            total_total_mandays += total_mandays
            total_total_wage += total_wage
            total_pf_restricted_per_day += pf_restricted_per_day
            total_pf_wage += pf_wage
            total_actual_pf_13 += actual_pf_13
            total_pf_60 += pf_60
            total_pf_service_charge += pf_service_charge
            total_basic += basic
            total_hra += hra
            total_wa += wa
            total_fixed_gross_esi += fixed_gross_esi
            total_esi_wage += esi_wage
            total_esi_ot_wage += esi_ot_wage
            total_esi_3_25 += esi_3_25
            total_total_ot_amount += total_ot_amount
            total_att_bonus += att_bonus
            total_fest_allowance += fest_allowance
            total_grand_total += grand_total
            total_service_charge_10 += service_charge_10
            total_salary_for_fest += salary_for_fest
            total_fest_day_service_charge += fest_day_service_charge
            total_dwsi_ot += dwsi_ot
            total_hj_ot += hj_ot
            total_not_present_paid += not_present_paid
            total_deduction_service_charge += deduction_service_charge

        subtotal_row = ["Subtotal"] + [""]*(8+total_days)
        subtotal_row += [formatted(total_shift_a_plan), formatted(total_shift_b_plan), formatted(total_shift_c_plan), formatted(total_work_plan), formatted(total_holidays_plan), formatted(total_total_plan), formatted(total_shift_a_attend), formatted(total_shift_b_attend), formatted(total_shift_c_attend), formatted(total_work_attend), formatted(total_week_off_attend), formatted(total_total_attend), formatted(total_nal), formatted(total_al), formatted(total_half_day_count), formatted(total_total_after_attend), formatted(total_cont_a)]
        subtotal_row += [formatted(total_ot_hours), formatted(total_existing_gross), formatted(total_increase), formatted(total_revised_gross), formatted(total_ot_rate), formatted(total_shoe_uniform), formatted(total_total), formatted(total_total_mandays), formatted(total_total_wage), formatted(total_pf_restricted_per_day), formatted(total_pf_wage), formatted(total_actual_pf_13), formatted(total_pf_60), formatted(total_pf_service_charge), formatted(total_basic), formatted(total_hra), formatted(total_wa), formatted(total_fixed_gross_esi), formatted(total_esi_wage), formatted(total_esi_ot_wage), formatted(total_esi_3_25), formatted(total_ot_hours), formatted(total_total_ot_amount), formatted(total_att_bonus), formatted(total_fest_allowance), formatted(total_grand_total), formatted(total_service_charge_10), formatted(total_dwsi_ot), formatted(total_hj_ot), formatted(total_salary_for_fest), formatted(total_fest_day_service_charge), formatted(total_not_present_paid), formatted(total_deduction_service_charge)]
        ws.append(subtotal_row)
        ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
        cell = ws.cell(row=row_start, column=1)
        cell.font = Font(bold=True)
        for i in range(10+total_days, 60+total_days):
            cell = ws.cell(row=row_start, column=i)
            cell.font = Font(bold=True)
        from_date = start_date.strftime('%Y-%m-%d')
        to_date = end_date.strftime('%Y-%m-%d')
        for i in holiday_col_color(from_date, to_date):
            ws.cell(row=row_start, column=i).fill = holiday_color

        if dept.department == 'FMD':
            fmd_shift_a_plan += total_shift_a_plan
            fmd_shift_b_plan += total_shift_b_plan
            fmd_shift_c_plan += total_shift_c_plan
            fmd_week_off += total_week_off
            fmd_holidays_plan += total_holidays_plan
            fmd_work_plan += total_work_plan
            fmd_total_plan += total_total_plan
            fmd_shift_a_attend += total_shift_a_attend
            fmd_shift_b_attend += total_shift_b_attend
            fmd_shift_c_attend += total_shift_c_attend
            fmd_work_attend += total_work_attend
            fmd_week_off_attend += total_week_off_attend
            fmd_total_attend += total_total_attend
            fmd_nal += total_nal
            fmd_al += total_al
            fmd_half_day_count += total_half_day_count
            fmd_total_after_attend += total_total_after_attend
            fmd_cont_a += total_cont_a

            fmd_existing_gross += total_existing_gross
            fmd_increase += total_increase
            fmd_revised_gross += total_revised_gross
            fmd_ot_rate += total_ot_rate
            fmd_shoe_uniform += total_shoe_uniform
            fmd_total_ += total_total
            fmd_total_mandays += total_total_mandays
            fmd_total_wage += total_total_wage
            fmd_pf_restricted_per_day += total_pf_restricted_per_day
            fmd_pf_wage += total_pf_wage
            fmd_actual_pf_13 += total_actual_pf_13
            fmd_pf_60 += total_pf_60
            fmd_pf_service_charge += total_pf_service_charge
            fmd_basic += total_basic
            fmd_hra += total_hra
            fmd_wa += total_wa
            fmd_fixed_gross_esi += total_fixed_gross_esi
            fmd_esi_wage += total_esi_wage
            fmd_esi_ot_wage += total_esi_ot_wage
            fmd_esi_3_25 += total_esi_3_25
            fmd_ot_hours += total_ot_hours
            fmd_total_ot_amount += total_total_ot_amount
            fmd_att_bonus += total_att_bonus
            fmd_fest_allowance += total_fest_allowance
            fmd_grand_total += total_grand_total
            fmd_service_charge_10 += total_service_charge_10
            fmd_salary_for_fest += total_salary_for_fest
            fmd_fest_day_service_charge += total_fest_day_service_charge
            fmd_dwsi_ot += total_dwsi_ot
            fmd_hj_ot += total_hj_ot
            fmd_not_present_paid += total_not_present_paid
            fmd_deduction_secvice_charge += total_deduction_service_charge
        else:
            ht_shift_a_plan += total_shift_a_plan
            ht_shift_b_plan += total_shift_b_plan
            ht_shift_c_plan += total_shift_c_plan
            ht_week_off += total_week_off
            ht_holidays_plan += total_holidays_plan
            ht_work_plan += total_work_plan
            ht_total_plan += total_total_plan
            ht_shift_a_attend += total_shift_a_attend
            ht_shift_b_attend += total_shift_b_attend
            ht_shift_c_attend += total_shift_c_attend
            ht_work_attend += total_work_attend
            ht_week_off_attend += total_week_off_attend
            ht_total_attend += total_total_attend
            ht_nal += total_nal
            ht_al += total_al
            ht_half_day_count += total_half_day_count
            ht_total_after_attend += total_total_after_attend
            ht_cont_a += total_cont_a

            ht_existing_gross += total_existing_gross
            ht_increase += total_increase
            ht_revised_gross += total_revised_gross
            ht_ot_rate += total_ot_rate
            ht_shoe_uniform += total_shoe_uniform
            ht_total_ += total_total
            ht_total_mandays += total_total_mandays
            ht_total_wage += total_total_wage
            ht_pf_restricted_per_day += total_pf_restricted_per_day
            ht_pf_wage += total_pf_wage
            ht_actual_pf_13 += total_actual_pf_13
            ht_pf_60 += total_pf_60
            ht_pf_service_charge += total_pf_service_charge
            ht_basic += total_basic
            ht_hra += total_hra
            ht_wa += total_wa
            ht_fixed_gross_esi += total_fixed_gross_esi
            ht_esi_wage += total_esi_wage
            ht_esi_ot_wage += total_esi_ot_wage
            ht_esi_3_25 += total_esi_3_25
            ht_ot_hours += total_ot_hours
            ht_total_ot_amount += total_total_ot_amount
            ht_att_bonus += total_att_bonus
            ht_fest_allowance += total_fest_allowance
            ht_grand_total += total_grand_total
            ht_service_charge_10 += total_service_charge_10
            ht_salary_for_fest += total_salary_for_fest
            ht_fest_day_service_charge += total_fest_day_service_charge
            ht_dwsi_ot += total_dwsi_ot
            ht_hj_ot += total_hj_ot
            ht_not_present_paid += total_not_present_paid
            ht_deduction_secvice_charge += total_deduction_service_charge
    
        grand_shift_a_plan += total_shift_a_plan
        grand_shift_b_plan += total_shift_b_plan
        grand_shift_c_plan += total_shift_c_plan
        grand_week_off += total_week_off
        grand_holidays_plan += total_holidays_plan
        grand_work_plan += total_work_plan
        grand_total_plan += total_total_plan
        grand_shift_a_attend += total_shift_a_attend
        grand_shift_b_attend += total_shift_b_attend
        grand_shift_c_attend += total_shift_c_attend
        grand_work_attend += total_work_attend
        grand_week_off_attend += total_week_off_attend
        grand_total_attend += total_total_attend
        grand_nal += total_nal
        grand_al += total_al
        grand_half_day_count += total_half_day_count
        grand_total_after_attend += total_total_after_attend
        grand_cont_a += total_cont_a

        grand_existing_gross += total_existing_gross
        grand_increase += total_increase
        grand_revised_gross += total_revised_gross
        grand_ot_rate += total_ot_rate
        grand_shoe_uniform += total_shoe_uniform
        grand_total_ += total_total
        grand_total_mandays += total_total_mandays
        grand_total_wage += total_total_wage
        grand_pf_restricted_per_day += total_pf_restricted_per_day
        grand_pf_wage += total_pf_wage
        grand_actual_pf_13 += total_actual_pf_13
        grand_pf_60 += total_pf_60
        grand_pf_service_charge += total_pf_service_charge
        grand_basic += total_basic
        grand_hra += total_hra
        grand_wa += total_wa
        grand_fixed_gross_esi += total_fixed_gross_esi
        grand_esi_wage += total_esi_wage
        grand_esi_ot_wage += total_esi_ot_wage
        grand_esi_3_25 += total_esi_3_25
        grand_ot_hours += total_ot_hours
        grand_total_ot_amount += total_total_ot_amount
        grand_att_bonus += total_att_bonus
        grand_fest_allowance += total_fest_allowance
        grand_grand_total += total_grand_total
        grand_service_charge_10 += total_service_charge_10
        grand_salary_for_fest += total_salary_for_fest
        grand_fest_day_service_charge += total_fest_day_service_charge
        grand_dwsi_ot += total_dwsi_ot
        grand_hj_ot += total_hj_ot
        grand_not_present_paid += total_not_present_paid
        grand_deduction_secvice_charge += total_deduction_service_charge

    ws.append([""])
    row_start = row_start + 1
    ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
    ht_total = ["HT-TOTAL"] + [""] * (8+total_days)
    ht_total += [formatted(ht_shift_a_plan), formatted(ht_shift_b_plan), formatted(ht_shift_c_plan), formatted(ht_work_plan), formatted(ht_holidays_plan), formatted(ht_total_plan), formatted(ht_shift_a_attend), formatted(ht_shift_b_attend), formatted(ht_shift_c_attend), formatted(ht_work_attend), formatted(ht_week_off_attend), formatted(ht_total_attend), formatted(ht_nal), formatted(ht_al), formatted(ht_half_day_count), formatted(ht_total_after_attend), formatted(ht_cont_a)]
    ht_total += [formatted(ht_ot_hours), formatted(ht_existing_gross), formatted(ht_increase), formatted(ht_revised_gross), formatted(ht_ot_rate), formatted(ht_shoe_uniform), formatted(ht_total_), formatted(ht_total_mandays), formatted(ht_total_wage), formatted(ht_pf_restricted_per_day), formatted(ht_pf_wage), formatted(ht_actual_pf_13), formatted(ht_pf_60), formatted(ht_pf_service_charge), formatted(ht_basic), formatted(ht_hra), formatted(ht_wa), formatted(ht_fixed_gross_esi), formatted(ht_esi_wage), formatted(ht_esi_ot_wage), formatted(ht_esi_3_25), formatted(ht_ot_hours), formatted(ht_total_ot_amount), formatted(ht_att_bonus), formatted(ht_fest_allowance), formatted(ht_grand_total), formatted(ht_service_charge_10), formatted(ht_dwsi_ot), formatted(ht_hj_ot), formatted(ht_salary_for_fest), formatted(ht_fest_day_service_charge), formatted(ht_not_present_paid), formatted(ht_deduction_secvice_charge)]
    ws.append(ht_total)
    row_start = row_start + 1
    ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
    cell = ws.cell(row=row_start, column=1)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='left')
    for i in range(10+total_days, 60+total_days):
        cell = ws.cell(row=row_start, column=i)
        cell.font = Font(bold=True)
  
    ws.append([""])
    row_start = row_start + 1
    ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
    fmd_total = ["FMD-TOTAL"] + [""]*(8+total_days)
    fmd_total += [formatted(fmd_shift_a_plan), formatted(fmd_shift_b_plan), formatted(fmd_shift_c_plan), formatted(fmd_work_plan), formatted(fmd_holidays_plan), formatted(fmd_total_plan), formatted(fmd_shift_a_attend), formatted(fmd_shift_b_attend), formatted(fmd_shift_c_attend), formatted(fmd_work_attend), formatted(fmd_week_off_attend), formatted(fmd_total_attend), formatted(fmd_nal), formatted(fmd_al), formatted(fmd_half_day_count), formatted(fmd_total_after_attend), formatted(fmd_cont_a)]
    fmd_total += [formatted(fmd_ot_hours), formatted(fmd_existing_gross), formatted(fmd_increase), formatted(fmd_revised_gross), formatted(fmd_ot_rate), formatted(fmd_shoe_uniform), formatted(fmd_total_), formatted(fmd_total_mandays), formatted(fmd_total_wage), formatted(fmd_pf_restricted_per_day), formatted(fmd_pf_wage), formatted(fmd_actual_pf_13), formatted(fmd_pf_60), formatted(fmd_pf_service_charge), formatted(fmd_basic), formatted(fmd_hra), formatted(fmd_wa), formatted(fmd_fixed_gross_esi), formatted(fmd_esi_wage), formatted(fmd_esi_ot_wage), formatted(fmd_esi_3_25), formatted(fmd_ot_hours), formatted(fmd_total_ot_amount), formatted(fmd_att_bonus), formatted(fmd_fest_allowance), formatted(fmd_grand_total), formatted(fmd_service_charge_10), formatted(fmd_dwsi_ot), formatted(fmd_hj_ot), formatted(fmd_salary_for_fest), formatted(fmd_fest_day_service_charge), formatted(fmd_not_present_paid), formatted(fmd_deduction_secvice_charge)]
    ws.append(fmd_total)
    row_start = row_start + 1
    ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
    cell = ws.cell(row=row_start, column=1)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='left')
    for i in range(10+total_days, 60+total_days):
        cell = ws.cell(row=row_start, column=i)
        cell.font = Font(bold=True)
  
    ws.append([""])
    row_start = row_start + 1
    ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
    grand_total_row = ["Grand Total "] + [""]*(8+total_days)
    grand_total_row += [formatted(grand_shift_a_plan), formatted(grand_shift_b_plan), formatted(grand_shift_c_plan), formatted(grand_work_plan), formatted(grand_holidays_plan), formatted(grand_total_plan), formatted(grand_shift_a_attend), formatted(grand_shift_b_attend), formatted(grand_shift_c_attend), formatted(grand_work_attend), formatted(grand_week_off_attend), formatted(grand_total_attend), formatted(grand_nal), formatted(grand_al), formatted(grand_half_day_count), formatted(grand_total_after_attend), formatted(grand_cont_a)]
    grand_total_row += [formatted(grand_ot_hours), formatted(grand_existing_gross), formatted(grand_increase), formatted(grand_revised_gross), formatted(grand_ot_rate), formatted(grand_shoe_uniform), formatted(grand_total_), formatted(grand_total_mandays), formatted(grand_total_wage), formatted(grand_pf_restricted_per_day), formatted(grand_pf_wage), formatted(grand_actual_pf_13), formatted(grand_pf_60), formatted(grand_pf_service_charge), formatted(grand_basic), formatted(grand_hra), formatted(grand_wa), formatted(grand_fixed_gross_esi), formatted(grand_esi_wage), formatted(grand_esi_ot_wage), formatted(grand_esi_3_25), formatted(grand_ot_hours), formatted(grand_total_ot_amount), formatted(grand_att_bonus), formatted(grand_fest_allowance), formatted(grand_grand_total), formatted(grand_service_charge_10), formatted(grand_dwsi_ot), formatted(grand_hj_ot), formatted(grand_salary_for_fest), formatted(grand_fest_day_service_charge), formatted(grand_not_present_paid), formatted(grand_deduction_secvice_charge)]
    ws.append(grand_total_row)
    row_start = row_start + 1
    ws.merge_cells(start_row=row_start, end_row=row_start, start_column=1, end_column=9)
    cell = ws.cell(row=row_start, column=1)
    cell.font = Font(bold=True)
    for i in range(10+total_days, 60+total_days):
        cell = ws.cell(row=row_start, column=i)
        cell.font = Font(bold=True)
    
    for row in ws.iter_rows(): 
        for cell in row:     
            if cell.value == "Grand Total ":  
                last_row = cell.row
                break
    # Apply styles
    for row in ws.iter_rows(min_row=1, max_row=4, max_col=1):
        for cell in row:
            cell.font = font
    for row in ws.iter_rows(min_row=2, max_row=5, min_col=10, max_col=19):
        for cell in row:
            cell.border = thin_border
    for row in ws.iter_rows(min_row=1, max_row=5, min_col=21, max_col=29):
        for cell in row:
            cell.border = thin_border
    for row in ws.iter_rows(min_row=1, max_row=1, min_col=21, max_col=29):
        for cell in row:
            cell.fill = green_color
    for row in ws.iter_rows(min_row=4, max_row=5, min_col=21, max_col=29):
        for cell in row:
            cell.fill = yellow_color
    for row in ws.iter_rows(min_row=1, max_row=4, min_col=31, max_col=40):
        for cell in row:
            cell.border = thin_border
    for row in ws.iter_rows(min_row=7, max_row=8, min_col=1, max_col=total_days+59):
        for cell in row:
            cell.border = thin_border
            cell.alignment = align_center_center
    for row in ws.iter_rows(min_row=7, max_row=7, min_col=86, max_col=121):
        for cell in row:
            cell.alignment = align_center_center
    for row in ws.iter_rows(min_row=9, max_row=last_row, min_col=10, max_col=total_days+59):
        for cell in row:
            cell.alignment = align_center_center
            cell.border = thin_border
    for row in ws.iter_rows(min_row=9, max_row=last_row, min_col=1, max_col=8):
        for cell in row:
            cell.alignment = align_center_center
            cell.border = thin_border
    for row in ws.iter_rows(min_row=9, max_row=last_row, min_col=9, max_col=9):
        for cell in row:
            cell.border = thin_border
    
    for row in ws.iter_rows(min_row=8, max_row=8, min_col=total_days+16, max_col=19+total_days):
        for cell in row:
            cell.fill = actual_header_green

    for row in ws.iter_rows(min_row=8, max_row=8, min_col=total_days+22, max_col=23+total_days):
        for cell in row:
            cell.fill = yellow_color
    
    for row in ws.iter_rows(min_row=8, max_row=8, min_col=total_days+24, max_col=24+total_days):
        for cell in row:
            cell.fill = actual_green
    
    for row in range(25+total_days, 60+total_days): 
        ws.merge_cells(start_row=7, end_row=8, start_column=row, end_column=row)
    from_date = start_date.strftime('%Y-%m-%d')
    to_date = end_date.strftime('%Y-%m-%d')
    for i in holiday_col_color(from_date, to_date):
        ws.cell(row=8, column=i).fill = holiday_color
        ws.cell(row=7, column=i).fill = holiday_color

    
    for row in ws.iter_rows(): 
        for cell in row:     
            if cell.value == "Schedule ":  
                schedule_row = cell.row
                schedule_col = cell.column + 1
                for rows in ws.iter_rows(min_row=schedule_row, max_row=schedule_row, min_col=schedule_col, max_col=schedule_col+total_days-1):
                    for cells in rows:
                        cells.fill = schedule_green
                    for col_idx in holiday_col_color(from_date, to_date):
                        scell = ws.cell(row=schedule_row, column=col_idx)
                        scell.value = 4
                        scell.font = red_font
                break  
            if cell.value == "Continue ":  
                actual_row = cell.row
                actual_col = cell.column + 1 
                for rows in ws.iter_rows(min_row=actual_row, max_row=actual_row, min_col=actual_col, max_col=actual_col+total_days-1):
                    for cells in rows:
                        if cells.value:
                            cells.fill = continue_green
                break

    for row in ws.iter_rows(): 
        for cell in row:     
            if cell.value in [" X ", " XA ", " XB ", " XC "]: 
                cell.fill = light_orange
            if cell.value == " H ": 
                cell.fill = week_off_pink
                cell.font = dark_red
            if cell.value == " 4 ": 
                cell.font = red_font
            if cell.value in [" O ", " OA ", " OB ", " OC "]: 
                cell.fill = onleave_color
            if cell.value in [" HD ", " HA ", " HB ", " HC "]:
                cell.fill = purple
            
    
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def formatted(amount):
    return format_decimal(amount or 0, locale='en_IN')


def calculate_ot_hours(employee, start_date, end_date):
    return {
        "dwsi_ot_hrs": frappe.db.get_value('OT Type', {'employee': employee, 'from_date': start_date, 'to_date': end_date}, 'dwsi_ot') or 0,
        "contractor_ot_hrs": frappe.db.get_value('OT Type', {'employee': employee, 'from_date': start_date, 'to_date': end_date}, 'contractor_ot') or 0
    }


def holiday_col_color(from_date, to_date):
    holiday_dates = frappe.get_all('Holiday',
        filters={'holiday_date': ['between', [from_date, to_date]]},
        pluck='holiday_date'
    )
    return [d.day + 9 for d in holiday_dates]


