import frappe
import datetime
import math
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import re
from frappe import _
import frappe
from frappe.model.document import Document
from datetime import date, timedelta, datetime,time
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime,today, format_date)
import math
from frappe.utils import add_months, cint, flt, getdate, time_diff_in_hours,time_diff_in_seconds
import locale


import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
import openpyxl.styles as styles
from openpyxl.drawing.image import Image
import requests
import PIL
import io
import urllib3
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
from openpyxl.styles.numbers import FORMAT_PERCENTAGE
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,format_date,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime)



@frappe.whitelist()
def download():
    filename = 'D . Trainee'
    test = build_xlsx_response(filename)


@frappe.whitelist()
def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    ws = wb.create_sheet(sheet_name, 0)

    header0 = ['']
    ws.append(header0)
    header1 = ['PERSONNEL RECORD D. TRAINEES']
    ws.append(header1)    
    header2 = ['Prepared by']
    ws.append(header2)    

    header3 = ['S.No','Photo','ID','Name','Dept','Designation','DOB','DOJ','Age(yr.mth)','DOC','Qualification / Professional License','','','','','','Father Name','Contact No/ Nos','Remarks']
    ws.append(header3)  

    header4 = ['','','','','','','','','','','Regular Qualification (High rank 2 degree)','','','Additional Qualification (High rank 2 degree)','','','','','']
    ws.append(header4) 
    header5 = ['','','','','','','','','','','Degree','Institution','Completion yr.','Degree','Institution','Completion yr.','','','']
    ws.append(header5) 

    data=get_data(args)
    for row in data:
        ws.append(row)
        img_url=frappe.db.get_value("File",{'attached_to_name':row[2],"attached_to_doctype": 'Employee'},["file_url"])
        if img_url:
            image_url='http://192.168.11.20/' + img_url
            try:
                image_data = requests.get(image_url).content
                img = Image(io.BytesIO(image_data))
                img_cell = ws.cell(row=data.index(row) + 7, column=2)
                img.width = 100
                img.height = 100
                

                ws.add_image(img, img_cell.coordinate)
                img_height_pixels = img.height
                img_height_excel = img_height_pixels * 0.75
                ws.row_dimensions[data.index(row)+7].height = img_height_excel
            except Exception as e:
                print(f"Error adding image: {e}")    

    for i in range(1,11):
        ws.merge_cells(start_row=4,start_column=i,end_row=6,end_column=i)

    ws.merge_cells(start_row=4,start_column=11,end_row=4,end_column=16)
    ws.merge_cells(start_row=5,start_column=11,end_row=5,end_column=13)
    ws.merge_cells(start_row=5,start_column=14,end_row=5,end_column=16)
   
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=19)
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=19)
    ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=19)

    ws.merge_cells(start_row=4,start_column=17,end_row=6,end_column=17)
    ws.merge_cells(start_row=4,start_column=18,end_row=6,end_column=18)
    ws.merge_cells(start_row=4,start_column=19,end_row=6,end_column=19)

    ws.column_dimensions['J'].width = 15
    ws.column_dimensions['K'].width = 15
    ws.column_dimensions['L'].width = 15
    ws.column_dimensions['M'].width = 15
    ws.column_dimensions['N'].width = 15
    ws.column_dimensions['O'].width = 15

    ws.column_dimensions['P'].width = 15
    ws.column_dimensions['Q'].width = 15
    ws.column_dimensions['R'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 12
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['S'].width = 15



    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 10

    for row in range(1,6):
        ws.row_dimensions[row].height =30

    for row in range(7, ws.max_row+1):
        ws.row_dimensions[row].height = 80

    
    for header in ws.iter_rows(min_row=4, max_row=6,min_col=1, max_col=19):
        for cell in header:
            cell.fill = PatternFill(fgColor='acdbfa', fill_type="solid")
        
    border = Border(left=Side(border_style='thin', color='000000'),
        right=Side(border_style='thin', color='000000'),
        top=Side(border_style='thin', color='000000'),
        bottom=Side(border_style='thin', color='000000'))

    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1,max_col=19):
        for cell in rows:
            cell.border = border

    align_center = Alignment(horizontal='center',vertical='center') 
    for align in ws.iter_rows(min_row=1, max_row=ws.max_row,min_col=1, max_col=19):
            for cell in align:
                cell.alignment = align_center


    xlsx_file=BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

@frappe.whitelist()
def get_data(args):
    var = frappe.db.get_all('Employee', {'status': 'Active','employee_type':'D . Trainee'}, ["*"],order_by='department')
    count=0
    row=[]
    data=[]
    qual=''
    year=''
    schol=''
    post_qual=''
    post_year=''
    post_schol=''
    
    for i in var:
        chart_data = frappe.db.get_all('Organization Chart', {'employee': i.name,'employee_type':'D . Trainee'}, ["*"],order_by='department')
        age=calculate_age(i.date_of_birth)
        count+=1
        for chart in chart_data:
            qualification = frappe.get_all("Education", {"parent": chart.name}, ['*'])
            if qualification:
                for q in qualification:
                    if q.level == "Graduate" or q.level == "Under Graduate":
                        qual=q.qualification
                        year=q.year
                        schol=q.school_univ 

                    elif q.level == "Post Graduate":
                        post_qual=q.qualification
                        post_year=q.year
                        post_schol=q.school_univ 
            else:
                qual = None
                schol= None
                year= None
                post_qual = None
                post_year= None
                post_schol= None
        if not chart_data:
            qual = None
            schol= None
            year= None
            post_qual = None
            post_year= None
            post_schol= None
        row.append([ count, '',i.name,i.employee_name,i.department,i.designation,i.date_of_birth,i.date_of_joining,
            age,i.contract_end_date,qual, schol, year,post_qual, post_schol,post_year, i.father_detail_,i.personal_primary_contact_number])

    return row

@frappe.whitelist()
def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

@frappe.whitelist()
def calculate_age(date_of_birth):
    today = datetime.today()
    # age = today.year - date_of_birth.year, today.month -date_of_birth.month -((today.month, today.day) < (date_of_birth.month, date_of_birth.day))
    age_years = today.year - date_of_birth.year
    age_months = today.month - date_of_birth.month

    if age_months < 0:
        age_years -= 1
        age_months += 12    
    
    return f"{age_years} . {age_months}"

    
    
    



