from __future__ import unicode_literals
import frappe
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.data import format_date
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue

from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import re
from frappe import _
import frappe
from frappe.model.document import Document
from datetime import date, timedelta, datetime,time
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime,today, format_date)
import math
from frappe.utils import add_months, cint, flt, getdate, time_diff_in_hours,time_diff_in_seconds
import locale


import openpyxl
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
import openpyxl.styles as styles
from openpyxl.drawing.image import Image
import requests
import PIL
import io
import urllib3


@frappe.whitelist()
def download():
    filename = 'Organizational Chart'
    test = build_xlsx_response(filename)
    
def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    ws = wb.create_sheet(sheet_name, 0)
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20 
    ws.append(title(args))
    ws.append(title1(args))
    ws.append(get_col(args))
    ws.append(get_col1(args))
    ws.append(get_col2(args))
    ws.append([''])
    data= get_data(args)
    for row in data:
        ws.append(row)
        img_url=frappe.db.get_value("File",{'attached_to_name':row[6],"attached_to_doctype": 'Employee'},["file_url"])
        if img_url:
            image_url='http://192.168.11.20/' + img_url
            try:
                image_data = requests.get(image_url).content
                img = Image(io.BytesIO(image_data))
                img_cell = ws.cell(row=data.index(row) + 7, column=2)
                img.width = 100
                img.height = 100
                

                ws.add_image(img, img_cell.coordinate)
                img_height_pixels = img.height
                img_height_excel = img_height_pixels * 0.75
                ws.row_dimensions[data.index(row)+7].height = img_height_excel
            except Exception as e:
                print(f"Error adding image: {e}")
    ws.merge_cells(start_row=3, start_column=1, end_row=5, end_column=1 )
    ws.merge_cells(start_row=3, start_column=2, end_row=5, end_column=2 )
    ws.merge_cells(start_row=3, start_column=3, end_row=5, end_column=3 )
    ws.merge_cells(start_row=3, start_column=4, end_row=5, end_column=4 )
    ws.merge_cells(start_row=3, start_column=5, end_row=5, end_column=5 )
    ws.merge_cells(start_row=3, start_column=6, end_row=5, end_column=6 )
    ws.merge_cells(start_row=3, start_column=7, end_row=5, end_column=7 )
    ws.merge_cells(start_row=3, start_column=8, end_row=5, end_column=8 )
    ws.merge_cells(start_row=3, start_column=9, end_row=5, end_column=9 )
    ws.merge_cells(start_row=3, start_column=10, end_row=5, end_column=10 )
    ws.merge_cells(start_row=3, start_column=11, end_row=5, end_column=11 )
    ws.merge_cells(start_row=3, start_column=12, end_row=5, end_column=12 )
    ws.merge_cells(start_row=3, start_column=13, end_row=4, end_column=14 )
    ws.merge_cells(start_row=3, start_column=15, end_row=4, end_column=16 )
    year_data = get_yr()
    end_col = 17 + len(year_data) - 1
    ws.merge_cells(start_row=3, start_column=17, end_row=4, end_column=end_col )
    # ws.merge_cells(start_row=3, start_column=end_col+1, end_row=5, end_column=end_col+1 )
    end_col=end_col-1
    ws.merge_cells(start_row=3, start_column=end_col+2, end_row=4, end_column=end_col+len(year_data)+1 )
    ws.merge_cells(start_row=3, start_column=end_col+len(year_data)+2, end_row=4, end_column=end_col+len(year_data)+len(year_data)+1 )
    ws.merge_cells(start_row=3, start_column=end_col+len(year_data)+len(year_data)+2, end_row=4, end_column=end_col+len(year_data)+len(year_data)+len(year_data)+1 )
    width=end_col+len(year_data)+len(year_data)+len(year_data)+1
    ws.merge_cells(start_row=3, start_column=width+1, end_row=4, end_column=width+1 )
    ws.merge_cells(start_row=3, start_column=width+2, end_row=4, end_column=width+3 )
    ws.merge_cells(start_row=3, start_column=width+4, end_row=4, end_column=width+4)
    ws.merge_cells(start_row=3, start_column=width+5, end_row=4, end_column=width+ 6)
    # ws.merge_cells(start_row=3, start_column=width+6, end_row=5, end_column=width+ 6)
    ws.merge_cells(start_row=3, start_column=width+7, end_row=5, end_column=width+7 )
    ws.merge_cells(start_row=3, start_column=width+8, end_row=4, end_column=width+9 )
    ws.merge_cells(start_row=3, start_column=width+10, end_row=4, end_column=width+11 )
    ws.merge_cells(start_row=3, start_column=width+12 , end_row=4, end_column=width+18 )
    # ws.merge_cells(start_row=4, start_column=width+12, end_row=4, end_column=width+18)
    ws.merge_cells(start_row=3, start_column=width+19, end_row=5, end_column=width+19 )
    ws.merge_cells(start_row=3, start_column=width+20, end_row=5, end_column=width+20 )
    ws.merge_cells(start_row=3, start_column=width+21, end_row=5, end_column=width+21 )
    ws.merge_cells(start_row=3, start_column=width+22, end_row=5, end_column=width+22 )
    ws.merge_cells(start_row=3, start_column=width+23, end_row=5, end_column=width+23 )
    ws.merge_cells(start_row=3, start_column=width+24, end_row=5, end_column=width+24 )
    ws.merge_cells(start_row=3, start_column=width+25, end_row=5, end_column=width+25 )
    ws.merge_cells(start_row=3, start_column=width+26, end_row=5, end_column=width+26 )
    ws.merge_cells(start_row=3, start_column=width+26, end_row=5, end_column=width+26 )
    ws.merge_cells(start_row=3, start_column=width+27, end_row=5, end_column=width+27 )
    ws.merge_cells(start_row=3, start_column=width+28, end_row=5, end_column=width+28)
    ws.merge_cells(start_row=3, start_column=width+29, end_row=5, end_column=width+29 )
    ws.merge_cells(start_row=3, start_column=width+30, end_row=5, end_column=width+30 )
    ws.merge_cells(start_row=3, start_column=width+31, end_row=5, end_column=width+31 )
    to=width+32
    sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
    for s in sal:
        tot = to + 3
        ws.merge_cells(start_row=4, start_column= to, end_row=4, end_column= tot)
        to = tot + 1
    ws.merge_cells(start_row=3, start_column=to, end_row=4, end_column=to + 2 )
    ws.merge_cells(start_row=3, start_column=to+3, end_row=4, end_column=to + 12 )
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=500):
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True) 
    ws.row_dimensions[5].height=19
    for rows in ws.iter_rows(min_row=3, max_row=6, min_col=4, max_col=500):
        for cell in rows:
            if cell.value=='Overseas Training (If Any)' or cell.value=='Committee Member (If Any)' or cell.value=='Other Comments / Remarks' or  cell.value == 'Previous Employeer' or cell.value == 'Regular Qualification (High rank 2 degree)' or cell.value == 'Additional Qualification (High rank 2 degree)' or cell.value == 'License Name / Issuing Organization / Aquisition yr.' or cell.value == 'Joining Designation' or cell.value == 'Present Address' or cell.value == 'Permanent Address' or cell.value == 'Location' or cell.value == 'Official' or cell.value == 'Personal':
                column_letter = get_column_letter(cell.column)
                ws.column_dimensions[column_letter].width = 29
            if  cell.value == 'No of Evaluator' or cell.value == 'Jr.Engg' or cell.value == 'Engg' or cell.value == 'Sr.Engg' or cell.value == 'Asst Manager' or cell.value == 'Dy manager' or cell.value == 'Sr Manager' or cell.value == 'AGM' or cell.value == 'Average Score' or cell.value == 'Manager':
                column_letter = get_column_letter(cell.column)
                ws.column_dimensions[column_letter].width = 18
            if cell.value=='Color-blindness' or  cell.value=='Political Affairs' or cell.value=='Individual Character':
                column_letter = get_column_letter(cell.column)
                ws.column_dimensions[column_letter].width = 15
            if cell.value=='Health Condition- Medical test- Medical history':
                column_letter = get_column_letter(cell.column)
                ws.column_dimensions[column_letter].width = 18
    for cell in ws[3]:  
        if cell.value=='Attendance':
            sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
            length=0
            for i in sal:
                length+=1
            length=length*4
            length=length-1
            start_column = cell.column
            end_column = start_column + length  
            ws.merge_cells(start_row=3, start_column=start_column, end_row=3, end_column=end_column)
            # if cell.value == "Performance Apprisal":
            #     ws.merge_cells(start_row=3, start_column=start_column-1, end_row=6, end_column=start_column-1)
        if cell.value == "Performance Apprisal" or cell.value == 'Weighted Average':
            start_column = cell.column
            end_column = start_column + 13  
            ws.merge_cells(start_row=3, start_column=start_column, end_row=3, end_column=end_column)
            # if cell.value == "Performance Apprisal":
            # ws.merge_cells(start_row=3, start_column=end_column+1, end_row=6, end_column=end_column+1)
            for i in range(start_column, end_column + 1):
                ws.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
            # ws.merge_cells(start_row=3, start_column=start_column-1, end_row=6, end_column=start_column-1)
                # ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 19  # Set column width to 19
        if cell.value=='Address':
            start_column = cell.column
            end_column = start_column + 1 
            ws.merge_cells(start_row=3, start_column=start_column, end_row=3, end_column=end_column)
        if cell.value == 'Award / Good Work Performance (If any)':
            start_column = cell.column
            end_column = start_column + 4  
            ws.merge_cells(start_row=3, start_column=start_column, end_row=4, end_column=end_column)
            # for i in range(start_column, end_column + 1):
            #     # ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 19  # Set column width to 19
        

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column )      
    align_center = Alignment(horizontal='center',vertical='center')
    ws.column_dimensions['B'].width = 10
    border_thin = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'))
    column_width = 19
    for col in range(2, 13):  
        col_letter = get_column_letter(col)  
        ws.column_dimensions[col_letter].width = column_width
    for header in ws.iter_rows(min_row=3, max_row=6, min_col=1, max_col=end_column):
            for cell in header:
                cell.fill = PatternFill(fgColor='c7e9f0', fill_type = "solid")
                cell.font = Font(bold=True,size=10)
                cell.alignment = align_center     
                cell.border=border_thin     
    att=get_leave()
    leave=len(att)
    
    align_center = Alignment(horizontal='center',vertical='center')
    for header in ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=1):
            for cell in header:
                cell.font = Font(bold=True,size=15)
    for header in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=1):
            for cell in header:
                cell.font = Font(bold=True,size=10)
    for header in ws.iter_rows(min_row=3, max_row=6, min_col=1, max_col=end_column):
            for cell in header:
                cell.fill = PatternFill(fgColor='c7e9f0', fill_type = "solid")
                cell.font = Font(bold=True,size=10)
                cell.alignment = align_center     
                cell.border=border_thin  
    header_range = ws['A1':ws.cell(row=len(get_data(args))+6, column=+width+200+leave).coordinate]
    for row in header_range:
        for cell in row:
            cell.border = border_thin  
    max_row = ws.max_row
    ws.column_dimensions['B'].width = 16
    ws.freeze_panes = 'D1'
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def title(args):
    et = args['employee_type']
    data=["PERSONNEL RECORD ("+ et +")"]
    return data

def title1(args):
    fdate = datetime.strptime(args['start_date'],'%Y-%m-%d')
    f_date = fdate.strftime('%Y.%m.%d')
    data=["Prepared on "+ today()]
    return data

def get_yr():
    sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
    return sal
def get_leave():
    row=[]
    sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
    for i in sal:
        row.append("Attendance in "+i.name[0:4])
        row.append("")
        row.append("")
        row.append("")
    return row
def get_col(args):
    row=[]
    sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
    row += ["S.No","Photo","Name","Department","Designation","Incharge","Employee ID","DOB","DOJ","Age","Years of Exp.","Working in DWSI",'Probation Date','','Confirmed Date','',"Salary (Gross in Rupees)"]
    for i in range(len(sal) - 1):
        row.append(sal[i].name)
    row+=["Department Transference"]
    for i in range(len(sal) - 1):
        row.append(sal[i].name)
    row+=["Designation Transference"]
    for i in range(len(sal) - 1):
        row.append(sal[i].name)
    row+=["Performance Grade"]
    for i in range(len(sal) - 1):
        row.append(sal[i].name)
        
    row+=["Previous Employeer","Qualification / Professional License","","Professional License","Address","","Location","Contact No.","","Email ID","","Family Details","","","","","","","Gender","Marital","Religion","Passport","Height (cm)","Weight (Kg)","Blood Group","Eyesight","Color-blindness","Hearing","Health Condition- Medical test- Medical history","Political Affairs","Individual Character","Attendance"]
    length=0
    for i in sal:
        length+=1
    length=length*4
    for i in range(1,length):
        row.append(" ")
    row+=['Reason For Leave','','',"Disciplinary Record (If any)","","","","","","","","","","Award / Good Work Performance (If any)","","","Overseas Training (If any)",'','Performance Apprisal','','','','','','','','','','','','','','Weighted Average','','','','','','','','','','','','','']
    return row
def get_col1(args):
    row=[]
    sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
    row +=["","","","","","","","","","","","",]
    for i in sal:
        row.append("") 
    # row.append(" ")
    for i in sal:
        row.append("")
    for i in sal:
        row.append("")
    for i in sal:
        row.append("")
    row += [""," "," "," ","Present Address","Permanent Address","","","","","","Name/DOB/Occupation/Coresidence","","","","","","","","","","","","","","","","","","",""]
    for i in sal:
        sdate = getdate(i.year_start_date)
        yvalue = i.name
        yearvalue = yvalue.split('-')[0]
        # print(year)
        row.append(yearvalue)
        row.append("")
        row.append("")
        row.append("")
    row+=["","","","","","","","","","","",'','','','','','','','','','','','No of Evaluator','Jr.Engg','Engg','Engg','Sr.Engg','Sr.Engg','Asst Manager','Asst Manager','Dy manager','Manager',"Sr Manager",'AGM','Average Score','Other Comments / Remarks']
    row+=['No of Evaluator','Jr.Engg','Engg','Engg','Sr.Engg','Sr.Engg','Asst Manager','Asst Manager','Dy manager','Manager',"Sr Manager",'AGM','Average Score','Other Comments / Remarks']
    return row
    
def get_col2(args):
    row=[]
    sal=frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
    row +=["","","","","","","","","","","","","Plan","Apprd","Plan","Apprd"]
    for i in sal:
        row.append((i.name)) 
    # row.append(" ")
    for i in sal:
        row.append(i.name)
    for i in sal:
        row.append(i.name)
    for i in sal:
        # row.append(i.name)
        input_year_range = i.name
        start_year, end_year = input_year_range.split("-")
        short_start_year = start_year[2:]
        short_end_year = end_year[2:]
        result = f"{short_start_year}-{short_end_year}"
        row.append(result)
    row += ["Latest / The before of the Latest","Regular Qualification (High rank 2 degree)","Additional Qualification (High rank 2 degree)","License Name / Issuing Organization / Aquisition yr.","Present Address","Permanent Address","","Official","Personal","Official","Personal","Father","Mother","Spouse","Child (1)","Child (2)","Child (3)","Child (4)","","","","","","","","","","","","",""]
    for i in sal:
        row.append("CL")
        row.append("SL")
        row.append("EL")
        row.append("Total")
    row+=['CL','SL','EL',"1st","2nd","3rd","4th","5th","6th",'7th','8th','9th','10th',"1st","2nd","3rd","Overseas Training (If Any)","Committee Member (If Any)"]
    return row

def get_data(args):
    data = []
    department=frappe.get_all("Department",{'name':('!=','All Departments')},['*'],order_by='order_value ASC')
    ind =1
    for dep in department:
        if args['employee_type']:
            emp_supervisors = frappe.get_all("Employee",{'employee_type':args['employee_type'],'designation': 'supervisor','status':'Active','department':dep.name},['*'],order_by='date_of_joining ASC')	
            emp_others = frappe.get_all("Employee",{'employee_type':args['employee_type'],'status':'Active','department':dep.name,'designation': ('!=','supervisor')},['*'],order_by='date_of_joining ASC')	
                    
        else:
            emp_supervisors = frappe.get_all("Employee",{'status':'Active','department':dep.name,'designation': 'supervisor',},['*'],order_by='date_of_joining ASC')
            emp_others = frappe.get_all("Employee",{'status':'Active','department':dep.name,'designation': ('!=','supervisor')},['*'],order_by='date_of_joining ASC')

        for e in emp_supervisors:
            chart_data = frappe.db.get_all('Organization Chart', {'employee': e.name}, ["*"],order_by='department')
            row = [ind,'',e.employee_name,e.department,e.designation,e.incharge,e.name,e.date_of_birth,e.date_of_joining,e.age,e.service_years_in_previous_companies,e.service_years_in_dwsi]
            for chart in chart_data:
                row.extend(chart.probation_planned or '')
                row.extend(chart.probation_approved or '')
                row.extend(chart.confirmation_planned or '')
                row.extend(chart.confirmation_approved or '')
                # org=frappe.db.get_all("Organization Chart",{'employee':e.name},['*'])
            # for o in org:
                fiscal_year = frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
                for i in fiscal_year:
                    sal = frappe.db.get_value("Salary Gross", {"parent": chart.name,'periodin_yrs':i.name},['gross_salary'])
                    if sal:
                        row.extend([sal])
                    else:
                        row.extend(["-"])
                # row.append(e.joining_designation or '')
                for i in fiscal_year:
                    dep = frappe.db.get_value("Department Change", {"parent": chart.name,'period':i.name},['from_department'])
                    if dep:
                        row.extend([dep +" "+ i.name])
                    else:
                        row.extend([" "])
                for i in fiscal_year:
                    des = frappe.db.get_value("Designation Change", {"parent": chart.name,'period':i.name},['from_designation'])
                    if des:
                        row.extend([des +" "+ i.name])
                    else:
                        row.extend([" "])
                for i in fiscal_year:
                    grade = frappe.db.get_value("Performance Grade", {"parent": chart.name,'period':i.name},['grade'])
                    if des:
                        row.extend([grade])
                    else:
                        row.extend([" "])
                pre_exp = frappe.get_all("Previous Work Experience", {"parent": chart.name},['*'],order_by='from_date DESC', limit=1) 
                if pre_exp:
                    for exp in pre_exp:
                        row.append(exp.location)
                else:
                    row.append(" ")
                qualification = frappe.get_all("Education", {"parent": chart.name}, ['*'])
                regular = []
                additional = []

                if qualification:
                    for q in qualification:
                        if q.level == "Graduate" or q.level == "Under Graduate":
                            regular.append(f"{q.qualification  or ' '} {q.year or ' '} {q.school_univ or ' '}")
                            # additional.append('')

                        elif q.level == "Post Graduate":
                            # regular.append('')
                            additional.append(f"{q.qualification  or ' '} {q.year or ' '} {q.school_univ or ' '}")

                        else:
                            regular.append('')
                            additional.append('')

                    row.extend(regular)
                    row.extend(additional)
                
                else:
                    row.extend([" ", " "])

                row.append("")	
                row.append(e.current_address or '')
                row.append(e.permanent_address or '')
                row.append(e.current_location or '')
                row.append(e.company_phone_number or '')
                row.append(e.personal_primary_contact_number or '')
                row.append(e.company_email or '')
                row.append(e.personal_email or '')
                row.append(e.father_detail_ or '')
                row.append(e.mother_details or '')
                row.append(e.spouse_details_ or '')
                row.append(e.child or '')
                row.append(e.child2 or '')
                row.append(e.child3 or '')
                row.append(e.child4 or '')
                row.append(e.gender or '')
                row.append(e.marital_status or '')
                row.append(e.religion or '')
                if e.passport_number___:
                    passport="Yes"
                else:
                    passport="No"
                row.append(passport)
                if e.height !=0:
                    row.append(e.height)
                else:
                    row.append("")
                if e.weight !=0:
                    row.append(e.weight)
                else:
                    row.append("")
                row.append(e.blood_group or '')
                row.append("")
                row.append("")
                row.append("")
                if e.self_health_issues_1!="None":
                    row.append(e.self_health_issues_1 or '')
                else:
                    row.append("")
                row.append(e.political_party_interested_in_political_party or '')
                row.append("")
                for i in fiscal_year:
                    leave=frappe.get_all("Leave Details",{"parent":chart.name,"Year":i.name},["*"])
                    if leave:
                        for l in leave:
                            row.append(l.cl)
                            row.append(l.sl)
                            row.append(l.el)
                            row.append(l.total)
                    else:
                        row+=["","","",""]
                cl_reason=chart.reason_for_cl
                sl_reason=chart.reason_for_sl
                el_reason=chart.reason_for_el
                row.append(cl_reason or '')
                row.append(sl_reason or '')
                row.append(el_reason or '')
                warned= frappe.get_all("Disciplinary Action",{"parent":chart.name},['*'],limit=10)
                if warned:
                    if len(warned)!=10:
                        warn=10-len(warned)
                        wn=len(warned)+int(warn)
                        for dis in range(wn):
                            if dis < len(warned):
                                row.append("Warning Ordered on " + str(warned[dis]['order_date']) + " for " + warned[dis]['reason'] or '')
                            else:
                                row.append("")
                    else:
                        for dis in warned:
                            row+=["Warning Ordered on "+ str(dis.order_date)+" for "+dis.reason]
                    # row.extend(discipline)
                else:
                    row.append(" ")
                awards=frappe.get_all("Rewards And Recognition",{"parent":chart.name},['*'],limit=3)
                if awards:
                    if len(awards)!=3:
                        award=3-len(awards)
                        awd=len(awards)+int(award)
                        for aw in range(awd):
                            if aw < len(awards):
                                row.append(awards[aw]['name_of_reward'] +" on "+str(awards[aw]["date_and_year"])+" for "+awards[aw]['reason'])
                            else:
                                row.append("")
                    else:
                        for aw in awards:
                            row+=[aw.name_of_reward+" on "+str(aw.date_and_year)+" for "+aw.reason]
                else:
                    row.append(" ")
                

                row.append(e.overseas_training_details or '')
                row.append(e.committee_member or '')
                # row.append(" ")
                # org=frappe.db.get_all("Organization Chart",{'employee':e.name},['*'])
                # if chart:
                #     for o in org:
                row.append(chart.p1 or '')
                row.append(chart.p2 or '')
                row.append(chart.p3 or '')
                row.append(chart.p4 or '')
                row.append(chart.p5 or '')
                row.append(chart.p6 or '')
                row.append(chart.p7 or '')
                row.append(chart.p8 or '')
                row.append(chart.p9 or '')
                row.append(chart.p10 or '')
                row.append(chart.p11 or '')
                row.append(chart.p12 or '')
                row.append(chart.p13 or '')
                row.append(chart.p14 or '')
                row.append(" ")
                row.append(chart.p15 or '')
                row.append(chart.p16 or '')
                row.append(chart.p17 or '')
                row.append(chart.p18 or '')
                row.append(chart.p19 or '')
                row.append(chart.p20 or '')
                row.append(chart.p21 or '')
                row.append(chart.p22 or '')
                row.append(chart.p23 or '')
                row.append(chart.p24 or '')
                row.append(chart.p25 or '')
                row.append(chart.p26 or '')
                row.append(chart.p27 or '')
                row.append(chart.p28 or '')
                # else:
                #     row.append(e.p1 or '')
                #     row.append(e.p2 or '')
                #     row.append(e.p3 or '')
                #     row.append(e.p4 or '')
                #     row.append(e.p5 or '')
                #     row.append(e.p6 or '')
                #     row.append(e.p7 or '')
                #     row.append(e.p8 or '')
                #     row.append(e.p9 or '')
                #     row.append(e.p10 or '')
                #     row.append(e.p11 or '')
                #     row.append(e.p12 or '')
                #     row.append(e.p13 or '')
                #     row.append(e.p14 or '')
                #     row.append(" ")
                #     row.append(e.p15 or '')
                #     row.append(e.p16 or '')
                #     row.append(e.p17 or '')
                #     row.append(e.p18 or '')
                #     row.append(e.p19 or '')
                #     row.append(e.p20 or '')
                #     row.append(e.p21 or '')
                #     row.append(e.p22 or '')
                #     row.append(e.p23 or '')
                #     row.append(e.p24 or '')
                #     row.append(e.p25 or '')
                #     row.append(e.p26 or '')
                #     row.append(e.p27 or '')
                #     row.append(e.p28 or '')



            ind+=1
            
            data.append(row)
    
    
    
        for e in emp_others:
            chart_data = frappe.db.get_all('Organization Chart', {'employee': e.name}, ["*"],order_by='department')
            row = [ind,'',e.employee_name,e.department,e.designation,e.incharge,e.name,e.date_of_birth,e.date_of_joining,e.age,e.service_years_in_previous_companies,e.service_years_in_dwsi]
            for chart in chart_data:
                row.extend(chart.probation_planned or '')
                row.extend(chart.probation_approved or '')
                row.extend(chart.confirmation_planned or '')
                row.extend(chart.confirmation_approved or '')
                fiscal_year = frappe.db.get_all("Fiscal Year",["name"],order_by="name ASC")
                for i in fiscal_year:
                    sal = frappe.db.get_value("Salary Gross", {"parent": chart.name,'periodin_yrs':i.name},['gross_salary'])
                    if sal:
                        row.extend([sal])
                    else:
                        row.extend(["-"])
                # row.append(e.joining_designation or '')
                for i in fiscal_year:
                    dep = frappe.db.get_value("Department Change", {"parent": chart.name,'period':i.name},['from_department'])
                    if dep:
                        row.extend([dep +" "+ i.name])
                    else:
                        row.extend([" "])
                for i in fiscal_year:
                    des = frappe.db.get_value("Designation Change", {"parent": chart.name,'period':i.name},['from_designation'])
                    if des:
                        row.extend([des +" "+ i.name])
                    else:
                        row.extend([" "])
                for i in fiscal_year:
                    grade = frappe.db.get_value("Performance Grade", {"parent": chart.name,'period':i.name},['grade'])
                    if des:
                        row.extend([grade])
                    else:
                        row.extend([" "])
                pre_exp = frappe.get_all("Previous Work Experience", {"parent": chart.name},['*'],order_by='from_date DESC', limit=1) 
                if pre_exp:
                    for exp in pre_exp:
                        row.append(exp.location)
                else:
                    row.append(" ")
                qualification = frappe.get_all("Education", {"parent": chart.name}, ['*'])
                regular = []
                additional = []

                if qualification:
                    for q in qualification:
                        if q.level == "Graduate" or q.level == "Under Graduate":
                            regular.append(f"{q.qualification  or ' '} {q.year or ' '} {q.school_univ or ' '}")
                            # additional.append('')

                        elif q.level == "Post Graduate":
                            # regular.append('')
                            additional.append(f"{q.qualification  or ' '} {q.year or ' '} {q.school_univ or ' '}")

                        else:
                            regular.append('')
                            additional.append('')

                    row.extend(regular)
                    row.extend(additional)
                
                else:
                    row.extend([" ", " "])

                row.append("")	
                row.append(e.current_address or '')
                row.append(e.permanent_address or '')
                row.append(e.current_location or '')
                row.append(e.company_phone_number or '')
                row.append(e.personal_primary_contact_number or '')
                row.append(e.company_email or '')
                row.append(e.personal_email or '')
                row.append(e.father_detail_ or '')
                row.append(e.mother_details or '')
                row.append(e.spouse_details_ or '')
                row.append(e.child or '')
                row.append(e.child2 or '')
                row.append(e.child3 or '')
                row.append(e.child4 or '')
                row.append(e.gender or '')
                row.append(e.marital_status or '')
                row.append(e.religion or '')
                if e.passport_number___:
                    passport="Yes"
                else:
                    passport="No"
                row.append(passport)
                if e.height !=0:
                    row.append(e.height)
                else:
                    row.append("")
                if e.weight !=0:
                    row.append(e.weight)
                else:
                    row.append("")
                row.append(e.blood_group or '')
                row.append("")
                row.append("")
                row.append("")
                if e.self_health_issues_1!="None":
                    row.append(e.self_health_issues_1 or '')
                else:
                    row.append("")
                row.append(e.political_party_interested_in_political_party or '')
                row.append("")
                for i in fiscal_year:
                    leave=frappe.get_all("Leave Details",{"parent":chart.name,"Year":i.name},["*"])
                    if leave:
                        for l in leave:
                            row.append(l.cl)
                            row.append(l.sl)
                            row.append(l.el)
                            row.append(l.total)
                    else:
                        row+=["","","",""]
                cl_reason=chart.reason_for_cl
                sl_reason=chart.reason_for_sl
                el_reason=chart.reason_for_el
                row.append(cl_reason or '')
                row.append(sl_reason or '')
                row.append(el_reason or '')
                warned= frappe.get_all("Disciplinary Action",{"parent":chart.name},['*'],limit=10)
                if warned:
                    if len(warned)!=10:
                        warn=10-len(warned)
                        wn=len(warned)+int(warn)
                        for dis in range(wn):
                            if dis < len(warned):
                                row.append("Warning Ordered on " + str(warned[dis]['order_date']) + " for " + warned[dis]['reason'] or '')
                            else:
                                row.append("")
                    else:
                        for dis in warned:
                            row+=["Warning Ordered on "+ str(dis.order_date)+" for "+dis.reason]
                    # row.extend(discipline)
                else:
                    row.append(" ")
                awards=frappe.get_all("Rewards And Recognition",{"parent":chart.name},['*'],limit=3)
                if awards:
                    if len(awards)!=3:
                        award=3-len(awards)
                        awd=len(awards)+int(award)
                        for aw in range(awd):
                            if aw < len(awards):
                                row.append(awards[aw]['name_of_reward'] +" on "+str(awards[aw]["date_and_year"])+" for "+awards[aw]['reason'])
                            else:
                                row.append("")
                    else:
                        for aw in awards:
                            row+=[aw.name_of_reward+" on "+str(aw.date_and_year)+" for "+aw.reason]
                else:
                    row.append(" ")
                

                row.append(e.overseas_training_details or '')
                row.append(e.committee_member or '')
                # row.append(" ")
                # org=frappe.db.get_all("Organization Chart",{'employee':e.name},['*'])
                # if org:
                #     for o in org:
                row.append(chart.p1 or '')
                row.append(chart.p2 or '')
                row.append(chart.p3 or '')
                row.append(chart.p4 or '')
                row.append(chart.p5 or '')
                row.append(chart.p6 or '')
                row.append(chart.p7 or '')
                row.append(chart.p8 or '')
                row.append(chart.p9 or '')
                row.append(chart.p10 or '')
                row.append(chart.p11 or '')
                row.append(chart.p12 or '')
                row.append(chart.p13 or '')
                row.append(chart.p14 or '')
                row.append(" ")
                row.append(chart.p15 or '')
                row.append(chart.p16 or '')
                row.append(chart.p17 or '')
                row.append(chart.p18 or '')
                row.append(chart.p19 or '')
                row.append(chart.p20 or '')
                row.append(chart.p21 or '')
                row.append(chart.p22 or '')
                row.append(chart.p23 or '')
                row.append(chart.p24 or '')
                row.append(chart.p25 or '')
                row.append(chart.p26 or '')
                row.append(chart.p27 or '')
                row.append(chart.p28 or '')
                # else:
                #     row.append(e.p1 or '')
                #     row.append(e.p2 or '')
                #     row.append(e.p3 or '')
                #     row.append(e.p4 or '')
                #     row.append(e.p5 or '')
                #     row.append(e.p6 or '')
                #     row.append(e.p7 or '')
                #     row.append(e.p8 or '')
                #     row.append(e.p9 or '')
                #     row.append(e.p10 or '')
                #     row.append(e.p11 or '')
                #     row.append(e.p12 or '')
                #     row.append(e.p13 or '')
                #     row.append(e.p14 or '')
                #     row.append(" ")
                #     row.append(e.p15 or '')
                #     row.append(e.p16 or '')
                #     row.append(e.p17 or '')
                #     row.append(e.p18 or '')
                #     row.append(e.p19 or '')
                #     row.append(e.p20 or '')
                #     row.append(e.p21 or '')
                #     row.append(e.p22 or '')
                #     row.append(e.p23 or '')
                #     row.append(e.p24 or '')
                #     row.append(e.p25 or '')
                #     row.append(e.p26 or '')
                #     row.append(e.p27 or '')
                #     row.append(e.p28 or '')


            ind+=1
            
            data.append(row)
        
    return data
