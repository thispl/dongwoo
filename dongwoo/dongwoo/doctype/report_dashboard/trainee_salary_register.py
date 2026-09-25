import frappe
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from six import BytesIO
from datetime import datetime, timedelta
import calendar
from frappe.utils import getdate

@frappe.whitelist()
def download():
    filename = 'STIPEND GET DT'
    test = build_xlsx_response(filename)
def format_currency(value):
    if value is None:
        return "0"  
    number_str = str(int(value))   
    if len(number_str) > 3:
        last_three = number_str[-3:]  
        other_digits = number_str[:-3] 
        
        formatted_other = []
        while len(other_digits) > 2:
            formatted_other.append(other_digits[-2:])
            other_digits = other_digits[:-2]
        if other_digits:
            formatted_other.append(other_digits)        
        formatted_other.reverse()
        formatted_number = ','.join(formatted_other) + ',' + last_three
    else:
        formatted_number = number_str
    
    return formatted_number

def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()

    att_date = getdate(args.start_date)
    att_month = att_date.strftime("%B")
    att_year = att_date.year
    # Execute the SQL query to fetch employee details along with corresponding salary slip names
    employee_details = frappe.db.sql("""
    SELECT e.bank_ac_no AS 'AC No', e.employee, e.first_name AS 'Employee Name', s.employee_name AS 'Salary Slip Ename', 
           e.mop, e.employee_type AS 'Employee Type', e.designation AS 'Designation', e.employee_type AS 'Etype', 
           e.department AS 'Department', e.uan_number AS "UAN NO", e.esi_number AS "Esi", e.date_of_joining AS 'Date of Joining',
           s.payment_days AS 'Payment Days', s.leave_without_pay AS 'LOP', s.absent_days AS 'ABS', s.total_working_days AS 'Days in Month',
           e.basic AS 'Basic Pay', e.house_rent_allowance AS 'HRA', e.medical_allowance AS 'Medical Allowance', 
           e.conveyance_allowance AS 'Conveyance', e.education_allowance AS 'Education Allowance', 
           e.leave_and_travel_allowance AS 'LTA', e.dress_allowance AS 'Dress Allowance', e.gross_pay AS 'Fixed Gross', 
           s.name AS 'Name', s.total_deduction AS 'Total Deduction'
    FROM `tabEmployee` e
    INNER JOIN `tabSalary Slip` s ON e.name = s.employee
    INNER JOIN `tabDepartment` d ON e.department = d.name
    INNER JOIN `tabDesignation` de ON de.name = e.designation
    WHERE s.start_date <= %s 
      AND s.end_date >= %s
      AND e.employee_type IN ("D . Trainee", "GET")
    ORDER BY d.order_value, e.employee_type, de.order, e.date_of_joining
""", (args.start_date, args.end_date), as_dict=True)

    # frappe.errprint(args.end_date)
    start_date = datetime.strptime(args.start_date, '%Y-%m-%d')
    current_year = start_date.year
    current_month = start_date.month
    if current_month == 1:
        prev_month = 12
        prev_year = current_year - 1
    else:
        prev_month = current_month - 1
        prev_year = current_year
    if prev_month == 1:
        prev_prev_month = 12
        prev_prev_year = prev_year - 1
    else:
        prev_prev_month = prev_month - 1
        prev_prev_year = prev_year
    prev_month_name = calendar.month_name[prev_month]
    next_month_name = calendar.month_name[prev_prev_month]
    pmonth= 'OT'+'('+prev_month_name+')'
    nmonth= 'OT'+'('+next_month_name+')'
    if start_date.month == 1:
        prev_month = 12
        prev_year = start_date.year - 1
    else:
        prev_month = start_date.month - 1
        prev_year = start_date.year
    
    prev_month_start = datetime(prev_year, prev_month, 1)
    prev_month_end = (prev_month_start + timedelta(days=31)).replace(day=1) - timedelta(days=1)    
    if prev_month == 1:
        prev_prev_month = 12
        prev_prev_year = prev_year - 1
    else:
        prev_prev_month = prev_month - 1
        prev_prev_year = prev_year
    prev_prev_month_start = datetime(prev_prev_year, prev_prev_month, 1)
    prev_prev_month_end = (prev_prev_month_start + timedelta(days=31)).replace(day=1) - timedelta(days=1)
    # prev_month_start, prev_month_end, prev_prev_month_start, prev_prev_month_end
    ws = wb.create_sheet(sheet_name, 0)
    ws.append(['TRAINEES STIPEND STATEMENT FOR THE MONTH OF'  ''  +str(att_month)+ '-' +str(att_year)])
    
    # ws.append(['S.NO','DEPT','EMP NO.','NAMES','A/c No','MOP','DESIGNATION','ESI NO','DATE OF JOINING','SALARY CALCULATION','','','','','','','','','','','','DEDUCTION','','','','','','TAKE HOME'])
    ws.append([
    'S.NO','DEPT','EMP NO.','NAMES','A/c No','MOP','DESIGNATION','ESI NO','DATE OF JOINING',
    '','','','','','','','','','','','','','','','','','TAKE HOME'
])
    ws.append(['S.NO','DEPT','EMP NO.','NAMES','A/c No','MOP','DESIGNATION','ESI NO','DATE OF JOINING','No.of.Days Paid','Basic Pay','Gross','OT',pmonth,nmonth,'Fest.Allow','Arrears','Attendance Bonus','Shift Allow','Site Allow','Total','LWF','Advance','PF(12%)','ESI @.75%','Total Deduction','LOP'])
    serial_number = 1
    previous_department = None
    prev_category=None
    pre_category="GET"
    department_total = 0
    total=0
    sub=0
    count=0
    count_1=0
    total1=0
    total_pf_dec=0
    takehome=0

    grand_total_days=0
    grand_total_ebasic=0
    grand_total_elwf=0
    grand_total_ot=0
    grand_total_egross=0
    grand_total_bus=0
    grand_total_fes=0
    grand_total_arrear=0
    grand_total_att=0
    grand_total_supervisor=0
    grand_total_shift=0
    grand_total_spec=0
    grand_total_etotal=0
    grand_total_pf=0
    grand_total_vpf=0
    grand_total_totalpf=0
    grand_total_pot=0
    grand_total_ppot=0
    grand_total_advance=0
    grand_total_esi=0
    grand_total_totaldec=0
    grand_total_dlop=0
    grand_total_takehome=0
    grand_total_desig=0
    grand_total_desigw=0
    grand_total_desigs=0
    grand_total_desigt=0
    
    sgrand_total_days=0
    sgrand_total_month=0
    sgrand_total_ebasic=0
    sgrand_total_elwf=0
    sgrand_total_ot=0
    sgrand_total_egross=0
    sgrand_total_bus=0
    sgrand_total_fes=0
    sgrand_total_arrear=0
    sgrand_total_att=0
    sgrand_total_supervisor=0
    sgrand_total_shift=0
    sgrand_total_spec=0
    sgrand_total_etotal=0
    sgrand_total_pf=0
    sgrand_total_totalpf= 0
    sgrand_total_pot=0
    sgrand_total_ppot=0
    sgrand_total_advance=0
    sgrand_total_esi=0
    sgrand_total_totaldec=0
    sgrand_total_dlop=0
    sgrand_total_takehome=0
    tgrand_total_days=0
    tgrand_total_month=0
    tgrand_total_ebasic=0
    tgrand_total_elwf=0
    tgrand_total_ot=0
    tgrand_total_egross=0
    tgrand_total_bus=0
    tgrand_total_fes=0
    tgrand_total_arrear=0
    tgrand_total_att=0
    tgrand_total_supervisor=0
    tgrand_total_shift=0
    tgrand_total_spec=0
    tgrand_total_etotal=0
    tgrand_total_pf=0
    tgrand_total_vpf=0
    tgrand_total_totalpf= 0
    tgrand_total_pot=0
    tgrand_total_ppot=0
    tgrand_total_advance=0
    tgrand_total_esi=0
    tgrand_total_totaldec=0
    tgrand_total_dlop=0
    tgrand_total_takehome=0
    tgrand_total_desig=0
    tot_emp=0
    employees_by_category = {
    # 'Staff': [],
    # 'Worker': [],
    'GET':[],
    'D . Trainee': []
    }

    for row in employee_details: 
        employee_type = row['Etype']
        if employee_type in employees_by_category:
            employees_by_category[employee_type].append(row)
    category_order = ['GET', 'D . Trainee']
    for category in category_order:
        for row in employees_by_category[category]:
            tot_emp+=1
            frappe.errprint(row['employee'])
            prevslip = frappe.db.get_value("Salary Slip", {"employee": row['employee'] ,"start_date":prev_month_start,"end_date":prev_month_end},["name"]) 
            if prevslip:
                prevot = frappe.db.get_value("Salary Detail", {"parent": prevslip , "salary_component": "Overtime"},"amount") or 0.0
            else:
                prevot = 0.0
            pprevslip = frappe.db.get_value("Salary Slip", {"employee": row['employee'] ,"start_date":prev_prev_month_start,"end_date":prev_prev_month_end},["name"]) 
            if pprevslip:
                pprevot = frappe.db.get_value("Salary Detail", {"parent": pprevslip , "salary_component": "Overtime"},"amount") or 0.0
            else:
                pprevot = 0.0
            ebasic = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Stipend"},"amount") or 0.0
            eot = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Overtime"},"amount") or 0.0
            efes = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Festival Allowance"},"amount") or 0.0
            earrear = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Arrear"},"amount") or 0.0
            eattendance = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Attendance Bonus"},"amount") or 0.0
            esupervisor = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Supervisor Allowance"},"amount") or 0.0
            eshift = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Shift Allowance"},"amount") or 0.0
            esda = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Site Allowance"},"amount") or 0.0
            elwf = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Labour Welfare  Fund"},"amount") or 0.0
            dprovident = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Provident Fund"},"amount") or 0.0
            dadvance = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Advance"},"amount") or 0.0
            desi = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Employee State Insurance"},"amount") or 0.0		
            dlop = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Loss Of Pay"},"amount") or 0.0
            egross = ebasic
            total = egross +efes+earrear+eattendance+esupervisor+eshift+esda+eot
            takehome= total - row['Total Deduction']
            
            if row['Department'] == previous_department:
                print(previous_department)
            else:
                if previous_department and row['Etype']==pre_category:
                    ws.append(['','SUBTOTAL', 'SUBTOTAL','', '', *sub, '', '','',
                     sub_total_days, format_currency(sub_total_ebasic), 
                    format_currency(sub_total_egross),format_currency(sub_total_ot),format_currency(sub_total_pot),format_currency(sub_total_ppot), format_currency(sub_total_fes), format_currency(sub_total_arrear), format_currency(sub_total_att), format_currency(sub_total_shift),format_currency(sub_total_spec), format_currency(sub_total_etotal), 
                    format_currency(sub_total_elwf),format_currency(sub_total_advance),format_currency(sub_total_pf), format_currency(sub_total_esi),format_currency(sub_total_totaldec), format_currency(sub_total_dlop),format_currency(sub_total_takehome)])
                count=0
                pre_category=row['Etype']
                department_name = row['Department']
                previous_department = row['Department']
                sub_total_lop=0
                sub_total_elwf=0
                sub_total_days=0
                sub_total_month=0
                sub_total_ebasic=0
                sub_total_ot=0
                sub_total_egross=0
                sub_total_bus=0
                sub_total_fes=0
                sub_total_arrear=0
                sub_total_att=0
                sub_total_supervisor=0
                sub_total_shift=0
                sub_total_spec=0
                sub_total_etotal=0
                sub_total_pf=0
                sub_total_vpf=0
                sub_total_totalpf=0
                sub_total_it=0
                sub_total_pot=0
                sub_total_ppot=0
                sub_total_advance=0
                sub_total_esi=0
                sub_total_totaldec=0
                sub_total_dlop=0
                sub_total_takehome=0

            sub_total_days+=row['Payment Days']
            sub_total_ebasic+=ebasic
            sub_total_ot+=eot
            sub_total_elwf+=elwf
            sub_total_egross+=egross
            sub_total_fes+=efes
            sub_total_arrear+=earrear
            sub_total_att+=eattendance
            sub_total_supervisor+=esupervisor
            sub_total_shift+=eshift
            sub_total_spec+=esda
            sub_total_etotal+=total
            sub_total_pf+=dprovident
            sub_total_totalpf+= total_pf_dec
            sub_total_pot+=prevot
            sub_total_ppot+=pprevot
            sub_total_advance+=dadvance
            sub_total_esi+=desi
            sub_total_totaldec+=row['Total Deduction']
            sub_total_lop+=dlop
            sub_total_takehome+=takehome
            grand_total_days+=row['Payment Days']
            grand_total_ebasic+=ebasic           
            grand_total_elwf+=elwf            
            grand_total_ot+=eot
            grand_total_egross+=egross           
            grand_total_fes+=efes
            grand_total_arrear+=earrear
            grand_total_att+=eattendance
            grand_total_supervisor+=esupervisor
            grand_total_shift+=eshift
            grand_total_spec+=esda
            grand_total_etotal+=total
            grand_total_pf+=dprovident           
            grand_total_totalpf+= total_pf_dec           
            grand_total_pot+=prevot
            grand_total_ppot+=pprevot
            grand_total_advance+=dadvance
            grand_total_esi+=desi
            grand_total_totaldec+=row['Total Deduction']
            grand_total_dlop+=dlop
            grand_total_takehome+=takehome
            grand_total_desig=serial_number
            
            if row['Etype'] == 'GET':
                grand_total_desigs+=1
                sgrand_total_days+=row['Payment Days']              
                sgrand_total_ebasic+=ebasic            
                sgrand_total_elwf+=elwf                
                sgrand_total_ot+=eot
                sgrand_total_egross+=egross              
                sgrand_total_fes+=efes
                sgrand_total_arrear+=earrear
                sgrand_total_att+=eattendance
                sgrand_total_supervisor+=esupervisor
                sgrand_total_shift+=eshift
                sgrand_total_spec+=esda
                sgrand_total_etotal+=total
                sgrand_total_pf+=dprovident              
                sgrand_total_totalpf+= total_pf_dec               
                sgrand_total_pot+=prevot
                sgrand_total_ppot+=pprevot
                sgrand_total_advance+=dadvance
                sgrand_total_esi+=desi
                sgrand_total_totaldec+=row['Total Deduction']
                sgrand_total_dlop+=dlop
                sgrand_total_takehome+=takehome
                sgrand_total_desig=serial_number
            if row['Etype'] == 'D . Trainee':
                grand_total_desigt+=1              
                tgrand_total_days+=row['Payment Days']             
                tgrand_total_ebasic+=ebasic
                tgrand_total_elwf+=elwf           
                tgrand_total_ot+=eot
                tgrand_total_egross+=egross               
                tgrand_total_fes+=efes
                tgrand_total_arrear+=earrear
                tgrand_total_att+=eattendance
                tgrand_total_supervisor+=esupervisor
                tgrand_total_shift+=eshift
                tgrand_total_spec+=esda
                tgrand_total_etotal+=total
                tgrand_total_pf+=dprovident
                
                tgrand_total_totalpf+= total_pf_dec
                
                tgrand_total_pot+=prevot
                tgrand_total_ppot+=pprevot
                tgrand_total_advance+=dadvance
                tgrand_total_esi+=desi
                tgrand_total_totaldec+=row['Total Deduction']
                tgrand_total_dlop+=dlop
                tgrand_total_takehome+=takehome
                tgrand_total_desig=serial_number
            
            if row['Etype'] == 'D . Trainee':
                ws.append([serial_number, department_name, row['employee'], row['Employee Name'],row['AC No'],row['mop'], row['Designation'],row['Esi'],row['Date of Joining'], row['Payment Days'], 
                format_currency(ebasic),  format_currency(egross),format_currency(eot),format_currency(prevot),
                format_currency(pprevot), format_currency(efes), format_currency(earrear), format_currency(eattendance), 
                format_currency(eshift), format_currency(esda),format_currency(total),
                format_currency(elwf), format_currency(dadvance),format_currency(dprovident),
                 format_currency(desi), format_currency(row['Total Deduction']),format_currency(dlop),format_currency(takehome)])
            else:
                ws.append([serial_number, department_name, row['employee'], row['Employee Name'],row['AC No'],
            row['mop'], row['Designation'],row['Esi'],row['Date of Joining'], 
                row['Payment Days'],format_currency(ebasic),  format_currency(egross),
                format_currency(eot),format_currency(prevot),format_currency(pprevot), format_currency(efes), 
                format_currency(earrear), format_currency(eattendance), format_currency(eshift), 
                format_currency(esda),format_currency(total), 
                format_currency(elwf), format_currency(dadvance), format_currency(dprovident),
                format_currency(desi), format_currency(row['Total Deduction']),format_currency(dlop),
                format_currency(takehome)])
            serial_number += 1
            count+=1
            sub = [str(count) + " Members"]
        count=0
        
        
        mem= [str(grand_total_desig)+ " Members"] 
        wmem= [str(grand_total_desigw)+ " Members"] 
        smem= [str(grand_total_desigs)+ " Members"] 
        tmem= [str(grand_total_desigt)+ " Members"]
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=28)
        for i in range(1,10):
            ws.merge_cells(start_row=2,start_column=i,end_row=3,end_column=i)
            i=i+1
        # ws.merge_cells(start_row=2, start_column=11, end_row=2, end_column=21)
        # ws.merge_cells(start_row=2, start_column=22, end_row=2, end_column=27)
        # merge Salary Calculation (cols 11–21)
        ws.merge_cells(start_row=2, start_column=11, end_row=2, end_column=21)
        ws.cell(row=2, column=11).value = "SALARY CALCULATION"

        # merge Deduction (cols 22–27)
        ws.merge_cells(start_row=2, start_column=22, end_row=2, end_column=27)
        ws.cell(row=2, column=22).value = "DEDUCTION"

        # merge Take Home (col 28)
        ws.merge_cells(start_row=2, start_column=28, end_row=3, end_column=28)
        ws.cell(row=2, column=28).value = "TAKE HOME"
        # ws.cell(row=2, column=11).value = "Salary Calculation"
        # ws.merge_cells(start_row=2,start_column=19,end_row=2,end_column=21)
        # ws.merge_cells(start_row=2,start_column=22,end_row=2,end_column=41)
        # ws.merge_cells(start_row=2,start_column=42,end_row=2,end_column=52)
        ws.merge_cells(start_row=2,start_column=28,end_row=3,end_column=28)
        if tot_emp > 0:
        
            # ws.append(['','SUBTOTAL', 'SUBTOTAL','', '', *sub, '', '','', '',sub_total_days, format_currency(sub_total_ebasic), 
            # format_currency(sub_total_egross),format_currency(sub_total_ot),format_currency(sub_total_pot),format_currency(sub_total_ppot), format_currency(sub_total_fes), format_currency(sub_total_arrear), format_currency(sub_total_att),format_currency(sub_total_shift),format_currency(sub_total_spec), format_currency(sub_total_etotal), 
            # format_currency(sub_total_pf), format_currency(sub_total_elwf), format_currency(sub_total_advance),format_currency(sub_total_esi),format_currency(sub_total_totaldec), format_currency(sub_total_dlop),format_currency(sub_total_takehome)])
            ws.append(['','SUBTOTAL', 'SUBTOTAL','', '', *sub, '', '','', sub_total_days, format_currency(sub_total_ebasic), 
            format_currency(sub_total_egross),format_currency(sub_total_ot),format_currency(sub_total_pot),format_currency(sub_total_ppot),  format_currency(sub_total_fes), format_currency(sub_total_arrear), format_currency(sub_total_att),format_currency(sub_total_shift),format_currency(sub_total_spec), format_currency(sub_total_etotal), 
            format_currency(sub_total_elwf), format_currency(sub_total_advance),format_currency(sub_total_pf),format_currency(sub_total_esi),format_currency(sub_total_totaldec), format_currency(sub_total_dlop),format_currency(sub_total_takehome)])
            
            prev_category=row['Etype']
            pre_category=row['Etype']
            tot_emp=0       
            sub_total_lop=0
            sub_total_elwf=0
            sub_total_days=0
            sub_total_month=0
            sub_total_ebasic=0        
            sub_total_ot=0
            sub_total_egross=0
            sub_total_bus=0
            sub_total_fes=0
            sub_total_arrear=0
            sub_total_att=0
            sub_total_supervisor=0
            sub_total_shift=0
            sub_total_spec=0
            sub_total_etotal=0
            sub_total_pf=0
            sub_total_vpf=0
            sub_total_totalpf=0
            sub_total_it=0
            sub_total_pot=0
            sub_total_ppot=0
            sub_total_mis=0
            sub_total_pt=0
            sub_total_advance=0
            sub_total_esi=0
            sub_total_totaldec=0
            sub_total_dlop=0
            sub_total_takehome=0

            if row['Etype'] == 'GET':
                ws.append(['', 'TOTAL FOR GET','TOTAL FOR GET', '','',*smem,'','','',
                        sgrand_total_days,  format_currency(sgrand_total_ebasic),
                            format_currency(sgrand_total_egross), format_currency(sgrand_total_ot),format_currency(sgrand_total_pot),format_currency(sgrand_total_ppot), format_currency(sgrand_total_fes), format_currency(sgrand_total_arrear), format_currency(sgrand_total_att),format_currency(sgrand_total_shift),format_currency(sgrand_total_spec), format_currency(sgrand_total_etotal), 
                            format_currency(sgrand_total_elwf), format_currency(sgrand_total_advance), format_currency(sgrand_total_pf), format_currency(sgrand_total_esi), format_currency(sgrand_total_totaldec), format_currency(sgrand_total_dlop), format_currency(sgrand_total_takehome)])
            elif row['Etype'] == 'D . Trainee':
                ws.append(['', 'TOTAL FOR D.TRAINEE','TOTAL FOR TRAINEE', '','',*tmem,'','','',
                         tgrand_total_days, format_currency(tgrand_total_ebasic),
                            format_currency(tgrand_total_egross), format_currency(tgrand_total_ot),format_currency(tgrand_total_pot),format_currency(tgrand_total_ppot), format_currency(tgrand_total_fes), format_currency(tgrand_total_arrear), format_currency(tgrand_total_att), format_currency(tgrand_total_shift),format_currency(tgrand_total_spec), format_currency(tgrand_total_etotal), 
                            format_currency(tgrand_total_elwf), format_currency(tgrand_total_advance), format_currency(tgrand_total_pf), format_currency(tgrand_total_esi), format_currency(tgrand_total_totaldec), format_currency(tgrand_total_dlop), format_currency(tgrand_total_takehome)])

    ws.append(['','Grand Total','Grand Total','','',*mem,'','','',
             grand_total_days,  format_currency(grand_total_ebasic),
                format_currency(grand_total_egross), format_currency(grand_total_ot),format_currency(grand_total_pot),format_currency(grand_total_ppot),format_currency(grand_total_fes), format_currency(grand_total_arrear), format_currency(grand_total_att),format_currency(grand_total_shift),format_currency(grand_total_spec), format_currency(grand_total_etotal), 
                format_currency(grand_total_elwf), format_currency(grand_total_advance), format_currency(grand_total_pf),format_currency(grand_total_esi), format_currency(grand_total_totaldec), format_currency(grand_total_dlop), format_currency(grand_total_takehome)])
    align_center = Alignment(horizontal='center',vertical='center')

    for cell in ws["1:1"]:
        cell.font = Font(bold=True,size=14)
        cell.alignment = align_center

    for cell in ws["2:2"]:
        if cell.column > 9 and cell.column != 50:
            # cell.font = Font(bold=True)
        # 	cell.alignment = align_center 
        # else:
            cell.font = Font(bold=True,size=10)
            cell.alignment = align_center     
    for cell in ws["2:2"]:
        cell.alignment = align_center 
    ws['A1'].fill = PatternFill(fgColor="FFFFFF", fill_type = "solid")

    for header in ws.iter_rows(min_row=2, max_row=3,min_col=1, max_col=27):
        for cell in header:
            cell.fill = PatternFill(fgColor='FFFF00', fill_type="solid")
    for header in ws.iter_rows(min_row=2, max_row=3,min_col=28, max_col=28):
        for cell in header:
            cell.fill = PatternFill(fgColor='FFC0CB', fill_type="solid")
    for header in ws.iter_rows(min_row=3, max_row=3,min_col=14, max_col=15):
        for cell in header:
            cell.fill = PatternFill(fgColor='3e8004', fill_type="solid")
    pink_fill = PatternFill(start_color='FFC0CB', end_color='FFC0CB', fill_type='solid')
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=28, max_col=28):
        for cell in row:
            cell.fill = pink_fill
    yellow_fill = PatternFill(start_color='fff600', end_color='fff600', fill_type='solid')
    yellow = PatternFill(start_color='fff600',end_color='fff600',fill_type='solid')
    green = PatternFill(start_color='adff2f',end_color='adff2f',fill_type='solid')
    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=3, max_col=3):
        for cell in rows:
            if cell.value == 'SUBTOTAL':
                for cell in ws[cell.row]:
                    cell.fill = green
                    # cell.alignment = Alignment(horizontal="center", vertical="center") 
                    ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=4) 
                    # ws.cell(row=cell.row, column=2).alignment = Alignment(horizontal="center", vertical="center")
            if cell.value == 'Grand Total' or cell.value=='TOTAL FOR WORKERS' or cell.value=='TOTAL FOR TRAINEE' or cell.value=='TOTAL FOR GET':
                for cell in ws[cell.row]:
                    cell.fill = yellow
                    # cell.alignment = Alignment(horizontal="center", vertical="center") 
                    ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=4) 
                    # ws.cell(row=cell.row, column=2).alignment = Alignment(horizontal="center", vertical="center")

    max_row = ws.max_row
    previous_department = None
    merge_start_row = None

    for row in range(1, max_row + 1):
        cell_value = ws.cell(row=row, column=2).value

        # Apply center alignment to the current cell
        ws.cell(row=row, column=2).alignment = Alignment(horizontal='center', vertical='center')

        if cell_value is None or cell_value.strip() == '':
            continue  # Skip empty rows

        if cell_value == previous_department:
            # Continue merging if the department is the same
            continue
        else:
            # If a new department is found, merge the previous group
            if merge_start_row is not None and previous_department is not None:
                ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=row - 1, end_column=2)

            # Update tracking variables for the new department
            previous_department = cell_value
            merge_start_row = row

    # Merge the last department group at the end of the loop
    if merge_start_row is not None and previous_department is not None:
        ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=max_row, end_column=2)
    right_align = Alignment(horizontal="right")
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=10, max_col=ws.max_column):
        for cell in row:
            cell.alignment = right_align
    # Handle the last group if needed
    if merge_start_row is not None and previous_department is not None:
        end_row = max_row
        ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=end_row, end_column=2)
    for cell in ws["3:3"]:
        cell.alignment = align_center
    align_left = Alignment(horizontal='left', vertical='center')
    for row in range(4, ws.max_row + 1): 
        for col in [1, 9]:
            cell = ws.cell(row=row, column=col) 
            cell.alignment = align_center
    for row in range(4, ws.max_row + 1): 
        cell = ws.cell(row=row, column=2) 
        cell.alignment = align_left
    for col in [2,4,5,7,8,9]:
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 20 
    # Define thin border
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Apply border to the whole sheet (all used rows & columns)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
    ws.column_dimensions['C'].width = 15 
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'
    
