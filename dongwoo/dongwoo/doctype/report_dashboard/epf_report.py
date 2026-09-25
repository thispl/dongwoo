import frappe
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from six import BytesIO
from datetime import datetime, timedelta
import calendar
from frappe.utils import getdate

@frappe.whitelist()
def download():
    filename = 'EPF Salary Report'
    test = build_xlsx_response(filename)

def format_currency(value):
    if value is None:
        return "0"
    
    number_str = str(int(value))
    
    if len(number_str) > 3:
        last_three = number_str[-3:]  
        other_digits = number_str[:-3] 
        
        formatted_other = []
        while len(other_digits) > 2:
            formatted_other.append(other_digits[-2:])
            other_digits = other_digits[:-2]
        if other_digits:
            formatted_other.append(other_digits)        
        formatted_other.reverse()
        formatted_number = ','.join(formatted_other) + ',' + last_three
    else:
        formatted_number = number_str
    
    return formatted_number


    

def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()

    att_date = getdate(args.start_date)
    att_month = att_date.strftime("%B")
    att_year = att_date.year
    # Execute the SQL query to fetch employee details along with corresponding salary slip names
    employee_details = frappe.db.sql("""
    SELECT e.bank_ac_no AS 'AC No', e.employee, e.first_name AS 'Employee Name', s.employee_name AS 'Salary Slip Ename', e.mop, e.employee_type AS 'Employee Type',
        e.designation AS 'Designation', e.employee_type AS 'Etype', e.department AS 'Department', e.uan_number AS "UAN NO", e.esi_number AS "Esi", e.date_of_joining AS 'Date of Joining',
        s.payment_days AS 'Payment Days', s.leave_without_pay AS 'LOP',s.absent_days AS 'ABS', s.total_working_days AS 'Days in Month',
        e.basic AS 'Basic Pay', e.house_rent_allowance AS 'HRA', e.medical_allowance AS 'Medical Allowance', e.conveyance_allowance AS 'Conveyance',
        e.education_allowance AS 'Education Allowance', e.leave_and_travel_allowance AS 'LTA', e.dress_allowance AS 'Dress Allowance', e.gross_pay AS 'Fixed Gross', 
        s.name AS 'Name', s.total_deduction AS 'Total Deduction'
    FROM `tabEmployee` e
    INNER JOIN `tabSalary Slip` s ON e.name = s.employee
    INNER JOIN `tabDepartment` d ON e.department = d.name
    INNER JOIN `tabDesignation` de ON de.name = e.designation
    WHERE e.name = s.employee and s.start_date <= %s AND s.end_date >= %s
    ORDER BY d.order_value, e.employee_type, de.order, e.date_of_joining
""", (args.start_date, args.end_date), as_dict=True)
    ws = wb.create_sheet(sheet_name, 0)
    ws.append(['PF STATEMENT FOR THE MONTH OF' +' '  +str(att_month)+ '-' +str(att_year)])
    ws.append(['S.NO','DEPT','UAN','EMP NO.','NAMES','No.of.Days','WAGE','LIMIT/RESTRICTED WAGES','','Basic Pay','Employee Contribution','','','Employer Contribution','','','Administrative Charges','','PF(13.00%)(Rs)','PF(25.00%)(Rs)'])
    ws.append(['','','','','','','','','','',' Regular','Addition','Total'," A/c No. 10 ",' A/c No. 1 ',' A/c No. 21 ',' A/c No. 2',' A/c No. 22','',''])
    ws.append(['','','','','','','Gross Wages','EPF Wages','EPS Wages','','PF(12%)','VPF','PF(12%)+VPF',"Pension Fund(8.33%)",'EPF (3.67%)','EDLI(0.5%)','ADMIN CHR(0.5% / 500)','Administration Charges for EDLI(0.01% / 200)','',''])
    
    serial_number = 0
    previous_department = None
    # prev_category=None
    pre_category="Staff"
    
    # serial_number = 1
    previous_department = None
    prev_category=None
    pre_category="Staff"
    department_total = 0
    # total=0
    # sub=0
    # count=0
    # count_1=0
    # total1=0
    # total_pf_dec=0
    # takehome=0
    tot_emp=0
    sub_tot_days=0
    sub_tot_gross=0
    sub_tot_epf=0
    sub_tot_epd=0
    sub_tot_restrict=0
    sub_tot_pf=0
    sub_tot_vpf=0
    sub_tot_pfvpf=0
    sub_tot_pension=0
    sub_tot_pen_epf=0
    sub_tot_admin=0
    sub_tot_admin_edli=0
    sub_tot_pf_rs=0
    sub_tot_pf_prs=0
    grand_tot_days=0
    grand_tot_gross=0
    grand_tot_epf=0
    grand_tot_epd=0
    grand_tot_restrict=0
    grand_tot_pf=0
    grand_tot_vpf=0
    grand_tot_pfvpf=0
    grand_tot_pension=0
    grand_tot_pen_epf=0
    grand_tot_admin=0
    grand_tot_admin_edli=0
    grand_tot_pf_rs=0
    grand_tot_pf_prs=0
    worker_tot_days=0
    worker_tot_gross=0
    worker_tot_epf=0
    worker_tot_epd=0
    worker_tot_restrict=0
    worker_tot_pf=0
    worker_tot_vpf=0
    worker_tot_pfvpf=0
    worker_tot_pension=0
    worker_tot_pen_epf=0
    worker_tot_admin=0
    worker_tot_admin_edli=0
    worker_tot_pf_rs=0
    worker_tot_pf_prs=0
    staff_tot_days=0
    staff_tot_gross=0
    staff_tot_epf=0
    staff_tot_epd=0
    staff_tot_restrict=0
    staff_tot_pf=0
    staff_tot_vpf=0
    staff_tot_pfvpf=0
    staff_tot_pension=0
    staff_tot_pen_epf=0
    staff_tot_admin=0
    staff_tot_admin_edli=0
    staff_tot_pf_rs=0
    staff_tot_pf_prs=0
    trainee_tot_days=0
    trainee_tot_gross=0
    trainee_tot_epf=0
    trainee_tot_epd=0
    trainee_tot_restrict=0
    trainee_tot_pf=0
    trainee_tot_vpf=0
    trainee_tot_pfvpf=0
    trainee_tot_pension=0
    trainee_tot_pen_epf=0
    trainee_tot_admin=0
    trainee_tot_admin_edli=0
    trainee_tot_pf_rs=0
    trainee_tot_pf_prs=0
    employees_by_category = {
    'Staff': [],
    'Worker': [],
    'D . Trainee': []
    }

    for row in employee_details: 
        employee_type = row['Etype']
        if employee_type in employees_by_category:
            employees_by_category[employee_type].append(row)
        category_order = ['Staff', 'Worker', 'D . Trainee']
    for category in category_order:
        for row in employees_by_category[category]:
            tot_emp+=1
            dprovident = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Provident Fund"},"amount") or 0.0
            dvpf = frappe.db.get_value("Salary Detail", {"parent": row['Name'] , "salary_component": "Voluntary Provident Fund"},"amount") or 0.0
            act_gross=frappe.db.get_value('Employee',{'name':row['employee']},['actual_gross'])
            if not act_gross:
                act_gross=0
            epf_wages=(15000/row['Days in Month'])*float(row['Payment Days'])
            epf_wages=round(epf_wages,0)
            totpf=dprovident + dvpf
            pension=round(epf_wages * 0.0833, 0)
            pension_tot=dprovident - round(epf_wages * 0.0833, 0)
            admin_char= round(epf_wages * 0.005, 2)
            pf_fin= round(epf_wages * 0.0833, 0) + dprovident - round(epf_wages * 0.0833, 0) + round(epf_wages * 0.005, 2) + round(epf_wages * 0.005, 2)
            pf_final=totpf+pf_fin
            serial_number += 1
   
            if row['Department'] == previous_department:
                print(previous_department)
            else:
                if previous_department and row['Etype']==pre_category:
                    ws.append(['','SUBTOTAL', 'SUBTOTAL','', '',sub_tot_days,format_currency(sub_tot_gross),format_currency(sub_tot_epf),format_currency(sub_tot_epd),format_currency(sub_tot_restrict),format_currency(sub_tot_pf),
            format_currency(sub_tot_vpf),format_currency(sub_tot_pfvpf),format_currency(sub_tot_pension),format_currency(sub_tot_pen_epf),format_currency(sub_tot_admin),format_currency(sub_tot_admin_edli),
            '-',format_currency(sub_tot_pf_rs),format_currency(sub_tot_pf_prs)])
                count=0
                pre_category=row['Etype']
                department_name = row['Department']
                previous_department = row['Department']
                sub_tot_days=0
                sub_tot_gross=0
                sub_tot_epf=0
                sub_tot_epd=0
                sub_tot_restrict=0
                sub_tot_pf=0
                sub_tot_vpf=0
                sub_tot_pfvpf=0
                sub_tot_pension=0
                sub_tot_pen_epf=0
                sub_tot_admin=0
                sub_tot_admin_edli=0
                sub_tot_pf_rs=0
                sub_tot_pf_prs=0
            sub_tot_days += float(row['Payment Days'])
            sub_tot_gross+=act_gross
            sub_tot_epf+=epf_wages
            sub_tot_epd+=epf_wages
            sub_tot_restrict+=epf_wages
            sub_tot_pf+=dprovident
            sub_tot_vpf+=dvpf
            sub_tot_pfvpf+=totpf
            sub_tot_pension+=pension
            sub_tot_pen_epf+=pension_tot
            sub_tot_admin+=admin_char
            sub_tot_admin_edli+=admin_char
            sub_tot_pf_rs+=pf_fin
            sub_tot_pf_prs+=pf_final
            
            grand_tot_days+= float(row['Payment Days'])
            grand_tot_gross+=act_gross
            grand_tot_epf+=epf_wages
            grand_tot_epd+=epf_wages
            grand_tot_restrict+=epf_wages
            grand_tot_pf+=dprovident
            grand_tot_vpf+=dvpf
            grand_tot_pfvpf+=totpf
            grand_tot_pension+=pension
            grand_tot_pen_epf+=pension_tot
            grand_tot_admin+=admin_char
            grand_tot_admin_edli+=admin_char
            grand_tot_pf_rs+=pf_fin
            grand_tot_pf_prs+=pf_final
            
            if row['Etype'] == 'Worker':
                worker_tot_days+= float(row['Payment Days'])
                worker_tot_gross+=act_gross
                worker_tot_epf+=epf_wages
                worker_tot_epd+=epf_wages
                worker_tot_restrict+=epf_wages
                worker_tot_pf+=dprovident
                worker_tot_vpf+=dvpf
                worker_tot_pfvpf+=totpf
                worker_tot_pension+=pension
                worker_tot_pen_epf+=pension_tot
                worker_tot_admin+=admin_char
                worker_tot_admin_edli+=admin_char
                worker_tot_pf_rs+=pf_fin
                worker_tot_pf_prs+=pf_final
                
            if row['Etype'] == 'Staff':
                staff_tot_days+= float(row['Payment Days'])
                staff_tot_gross+=act_gross
                staff_tot_epf+=epf_wages
                staff_tot_epd+=epf_wages
                staff_tot_restrict+=epf_wages
                staff_tot_pf+=dprovident
                staff_tot_vpf+=dvpf
                staff_tot_pfvpf+=totpf
                staff_tot_pension+=pension
                staff_tot_pen_epf+=pension_tot
                staff_tot_admin+=admin_char
                staff_tot_admin_edli+=admin_char
                staff_tot_pf_rs+=pf_fin
                staff_tot_pf_prs+=pf_final
            if row['Etype'] == 'D . Trainee':
                trainee_tot_days+= float(row['Payment Days'])
                trainee_tot_gross+=act_gross
                trainee_tot_epf+=epf_wages
                trainee_tot_epd+=epf_wages
                trainee_tot_restrict+=epf_wages
                trainee_tot_pf+=dprovident
                trainee_tot_vpf+=dvpf
                trainee_tot_pfvpf+=totpf
                trainee_tot_pension+=pension
                trainee_tot_pen_epf+=pension_tot
                trainee_tot_admin+=admin_char
                trainee_tot_admin_edli+=admin_char
                trainee_tot_pf_rs+=pf_fin
                trainee_tot_pf_prs+=pf_final
  
            ws.append([
                serial_number,
                row['Department'],
                row['UAN NO'],
                row['employee'],
                row['Employee Name'],
                row['Payment Days'],
                format_currency(act_gross),
                format_currency(epf_wages),
                format_currency(epf_wages),
                format_currency(epf_wages),
                format_currency(dprovident),
                format_currency(dvpf),
                format_currency(totpf),
                format_currency(pension), 
                format_currency(pension_tot), 
                format_currency(admin_char), 
                format_currency(admin_char), 
                '-',
                format_currency(pf_fin),format_currency(pf_final)
            ])
            
        if tot_emp > 0:
            ws.append(['','SUBTOTAL', 'SUBTOTAL','', '',sub_tot_days,format_currency(sub_tot_gross),format_currency(sub_tot_epf),format_currency(sub_tot_epd),format_currency(sub_tot_restrict),format_currency(sub_tot_pf),
            format_currency(sub_tot_vpf),format_currency(sub_tot_pfvpf),format_currency(sub_tot_pension),format_currency(sub_tot_pen_epf),format_currency(sub_tot_admin),format_currency(sub_tot_admin_edli),
            '-',format_currency(sub_tot_pf_rs),format_currency(sub_tot_pf_prs)])
            prev_category=row['Etype']
            pre_category=row['Etype']
            tot_emp=0


            if row['Etype'] == 'Worker':
                ws.append(['', 'TOTAL FOR WORKERS','TOTAL FOR WORKERS','','',worker_tot_days,format_currency(worker_tot_gross),format_currency(worker_tot_epf),format_currency(worker_tot_epd),format_currency(worker_tot_restrict),format_currency(worker_tot_pf),
                format_currency(worker_tot_vpf),format_currency(worker_tot_pfvpf),format_currency(worker_tot_pension),format_currency(worker_tot_pen_epf),format_currency(worker_tot_admin),format_currency(worker_tot_admin_edli),
                '-',format_currency(worker_tot_pf_rs),format_currency(worker_tot_pf_prs)])
            elif row['Etype'] == 'Staff':
                ws.append(['', 'TOTAL FOR STAFF','TOTAL FOR STAFF', '','',staff_tot_days,format_currency(staff_tot_gross),format_currency(staff_tot_epf),format_currency(staff_tot_epd),format_currency(staff_tot_restrict),format_currency(staff_tot_pf),
                format_currency(staff_tot_vpf),format_currency(staff_tot_pfvpf),format_currency(staff_tot_pension),format_currency(staff_tot_pen_epf),format_currency(staff_tot_admin),format_currency(staff_tot_admin_edli),
                '-',format_currency(staff_tot_pf_rs),format_currency(staff_tot_pf_prs)])
            elif row['Etype'] == 'D . Trainee':
                ws.append(['', 'TOTAL FOR TRAINEE','TOTAL FOR TRAINEE', '','',trainee_tot_days,format_currency(trainee_tot_gross),format_currency(trainee_tot_epf),format_currency(trainee_tot_epd),format_currency(trainee_tot_restrict),format_currency(trainee_tot_pf),
                format_currency(trainee_tot_vpf),format_currency(trainee_tot_pfvpf),format_currency(trainee_tot_pension),format_currency(trainee_tot_pen_epf),format_currency(trainee_tot_admin),format_currency(trainee_tot_admin_edli),
                '-',format_currency(trainee_tot_pf_rs),format_currency(trainee_tot_pf_prs)])
    ws.append(['','Grand Total','Grand Total','','',grand_tot_days,format_currency(grand_tot_gross),format_currency(grand_tot_epf),format_currency(grand_tot_epd),format_currency(grand_tot_restrict),format_currency(grand_tot_pf),
                format_currency(grand_tot_vpf),format_currency(grand_tot_pfvpf),format_currency(grand_tot_pension),format_currency(grand_tot_pen_epf),format_currency(grand_tot_admin),format_currency(grand_tot_admin_edli),
                '-',format_currency(grand_tot_pf_rs),format_currency(grand_tot_pf_prs)])
    align_center = Alignment(horizontal='center',vertical='center')
    max_row = ws.max_row    
    border = Border(left=Side(border_style='thin', color='000000'),
                 right=Side(border_style='thin', color='000000'),
                 top=Side(border_style='thin', color='000000'),
                 bottom=Side(border_style='thin', color='000000'))
    for cell in ws["1:1"]:
        cell.font = Font(bold=True,size=14)
        cell.alignment = align_center
    for cell in ws["2:2"]:
        cell.font = Font(bold=True,size=10)
        cell.alignment = align_center
    for cell in ws["3:3"]:
        cell.font = Font(bold=True,size=10)
        cell.alignment = align_center
    for cell in ws["4:4"]:
        cell.font = Font(bold=True,size=10)
        cell.alignment = align_center
    for row in ws.iter_rows(min_row=1, max_row=4, min_col=1, max_col=20):
        for cell in row:
            # Set wrap text and center alignment
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=20) 
    ws.merge_cells(start_row=2, start_column=1, end_row=4, end_column=1)
    ws.merge_cells(start_row=2, start_column=2, end_row=4, end_column=2)
    ws.merge_cells(start_row=2, start_column=3, end_row=4, end_column=3)
    ws.merge_cells(start_row=2, start_column=4, end_row=4, end_column=4)
    ws.merge_cells(start_row=2, start_column=5, end_row=4, end_column=5)
    ws.merge_cells(start_row=2, start_column=6, end_row=4, end_column=6)
    ws.merge_cells(start_row=2, start_column=7, end_row=3, end_column=7)
    ws.merge_cells(start_row=2, start_column=8, end_row=3, end_column=9)
    ws.merge_cells(start_row=2, start_column=10, end_row=4, end_column=10)
    ws.merge_cells(start_row=2, start_column=11, end_row=2, end_column=13)
    ws.merge_cells(start_row=2, start_column=14, end_row=2, end_column=16)
    ws.merge_cells(start_row=2, start_column=17, end_row=2, end_column=18)
    ws.merge_cells(start_row=2, start_column=19, end_row=4, end_column=19)
    ws.merge_cells(start_row=2, start_column=20, end_row=4, end_column=20)
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['Q'].width = 15
    ws.column_dimensions['R'].width = 15
    ws.column_dimensions['C'].width = 13
    ws.column_dimensions['E'].width = 13
    ws.column_dimensions['N'].width = 11
    border = Border(left=Side(border_style='thin', color='000000'),
                right=Side(border_style='thin', color='000000'),
                top=Side(border_style='thin', color='000000'),
                bottom=Side(border_style='thin', color='000000'))
    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=20):
        for cell in rows:
            cell.border = border
    for header in ws.iter_rows(min_row=4, max_row=4,min_col=6, max_col=9):
        for cell in header:
            cell.fill = PatternFill(fgColor='0b12db', fill_type="solid")
            cell.font = Font(color='FFFFFF')
    for rows in ws.iter_rows(min_row=3, max_row=3, min_col=14, max_col=18):
        for cell in rows:
            cell.font = Font(color='FF0000')
    ws.row_dimensions[3].height = 15  
    ws.row_dimensions[4].height = 35
    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=20):
        for cell in rows:
            cell.border = border
    
    yellow = PatternFill(start_color='fff600',end_color='fff600',fill_type='solid')
    green = PatternFill(start_color='98c91c',end_color='98c91c',fill_type='solid')
    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=3, max_col=3):
        for cell in rows:
            if cell.value == 'SUBTOTAL':
                for cell in ws[cell.row]:
                    cell.fill = green
                    ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=5) 
            if cell.value == 'Grand Total' or cell.value=='TOTAL FOR WORKERS' or cell.value=='TOTAL FOR TRAINEE' or cell.value=='TOTAL FOR STAFF':
                for cell in ws[cell.row]:
                    cell.fill = yellow
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=4) 
    max_row = ws.max_row
    previous_department = None
    merge_start_row = None

    for row in range(1, max_row + 1):
        cell_value = ws.cell(row=row, column=2).value

        # Apply center alignment to the current cell
        ws.cell(row=row, column=2).alignment = Alignment(horizontal='center', vertical='center')
        
        if cell_value is None or cell_value.strip() == '':
            previous_department = None
            merge_start_row = None
        elif cell_value == previous_department:
            if merge_start_row is not None:
                end_row = row
                ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=end_row, end_column=2)
        else:
            if merge_start_row is not None and previous_department is not None:
                end_row = row - 1
                ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=end_row, end_column=2)
            
            previous_department = cell_value
            merge_start_row = row 

    # Handle the last group if needed
    if merge_start_row is not None and previous_department is not None:
        end_row = max_row
        ws.merge_cells(start_row=merge_start_row, start_column=2, end_row=end_row, end_column=2)
    for rows in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=20):
        for cell in rows:
            cell.alignment = align_center
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'
    
