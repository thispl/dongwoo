from __future__ import unicode_literals
import frappe
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.data import format_date
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue
from six import BytesIO, string_types
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
import re
from frappe import _
import frappe
from frappe.model.document import Document
from datetime import date, timedelta, datetime,time
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime,today, format_date)
import math
from frappe.utils import add_months, cint, flt, getdate, time_diff_in_hours,time_diff_in_seconds
import locale
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
MONTH_MAP = {
    "Jan": 1,
    "Feb": 2,
    "Mar": 3,
    "Apr": 4,
    "May": 5,
    "June": 6,
    "July": 7,
    "Aug": 8,
    "Sep": 9,
    "Oct": 10,
    "Nov": 11,
    "Dec": 12
}


@frappe.whitelist()
def download():
    filename = 'Overtime Report'
    build_xlsx_response(filename)


def build_xlsx_response(filename):
    args = frappe.local.form_dict
    xlsx_file = make_xlsx(args, filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'


def make_xlsx(args, sheet_name=None, wb=None, column_widths=None):
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  
    ws = wb.create_sheet(sheet_name or "Sheet1", 0)
    ws.column_dimensions['A'].width = 50
    columns = ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']
    for c in columns:
        ws.column_dimensions[c].width = 12
    period = f'{args["month"]}-{args["year"]}'
    ws.append(['OVERTIME ANALYSIS']+ [''] * 8)
    ws.append(['Reason wise', period] + [''] * 7)
    ws.append(['', 'ADMIN', 'PROD', 'MAINT', 'QC', 'PE', 'HMI-PMT', 'PMT', 'Total'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=9)
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    data = get_data(args)
    for row in data:
        ws.append(row)
    align_center = Alignment(horizontal='center',vertical='center')
    for cell in ws["2:2"]:
        cell.alignment = align_center
    for cell in ws["3:3"]:
        cell.alignment = align_center
    align_right = Alignment(horizontal='right',vertical='center')
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=3, max_col=ws.max_column):
        for cell in row:
            cell.alignment = align_right
    green_fill = PatternFill(start_color='B7D996', end_color='B7D996', fill_type='solid')
    for row in ws.iter_rows(min_row=2, max_row=3, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.fill = green_fill
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    bold_font = Font(bold=True)
    target_row = ws.max_row - 5
    for cell in ws[target_row]:
        cell.fill = yellow_fill
        cell.font = bold_font
    blue_fill = PatternFill(start_color='C3DCE6', end_color='C3DCE6', fill_type='solid')
    target_row = ws.max_row - 3
    for cell in ws[target_row]:
        cell.fill = blue_fill
    for row in ws.iter_rows(min_row=ws.max_row - 3, max_row=ws.max_row):
        for cell in row:
            cell.font = bold_font
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
            if cell.row == 1:
                cell.font = Font(size=14, bold=True)  
    ws.row_dimensions[1].height = 25 
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file


def get_data(args):
    data = []
    departments = ['HR & Admin', 'Production', 'Maintenance', 'QC', 'PE', 'H.PMT', 'D.PMT']
    month_name = args["month"]
    month = MONTH_MAP.get(month_name)
    year = cint(args["year"])
    reasons = frappe.get_all("OT Reasons", filters={'reason': ['!=', '']}, fields=['name'], order_by='name')
    holiday_dates = frappe.db.sql("""
        SELECT holiday_date
        FROM `tabHoliday`
        WHERE parent = 'DWSI Holiday List'
            AND MONTH(holiday_date) = %s
            AND YEAR(holiday_date) = %s
            AND (national_holiday = 1 OR festival_holiday = 1)
    """, (month, year), as_dict=True)
    holiday_date_list = [d['holiday_date'] for d in holiday_dates]
    for r in reasons:
        row = [r.name]
        total_ot = 0
        for dept in departments:
            ot_hrs = frappe.db.sql("""
                SELECT SUM(total_hours) AS ot_total
                FROM `tabOvertime Request`
                WHERE MONTH(ot_date) = %s
                  AND YEAR(ot_date) = %s
                  AND is_considered_as != 'Compensatory Off'
                  AND department = %s
                  AND ot_reasons = %s
                  AND employee_type = 'Worker'
                  AND workflow_state = 'Approved'
            """, (month, year, dept, r.name), as_dict=True)
            ot = ot_hrs[0]['ot_total'] or 0 if ot_hrs else 0
            row.append(ot)
            total_ot += ot
        row.append(total_ot)
        data.append(row)
    row_ot = ['Total OT Hours']
    row_holiday = ['Worked on Paid Holidays']
    row_combined = ['Grand Total']
    row_employee_count = ['No. of Workers in Dept.']
    row_per_person = ['Per person OT without Holiday']
    row_per_person_incl = ['Per person OT with Holiday']
    total_ot = 0
    total_holiday_ot = 0
    total_employees = 0
    ot_values_by_dept = []
    combined_ot_by_dept = []
    emp_counts_by_dept = []
    for dept in departments:
        ot_hrs = frappe.db.sql("""
            SELECT SUM(total_hours) AS ot_total
            FROM `tabOvertime Request`
            WHERE MONTH(ot_date) = %s
              AND YEAR(ot_date) = %s
              AND is_considered_as != 'Compensatory Off'
              AND department = %s
              AND ot_reasons != ''
              AND employee_type = 'Worker'
              AND workflow_state = 'Approved'
        """, (month, year, dept), as_dict=True)
        ot = ot_hrs[0]['ot_total'] or 0
        row_ot.append(ot)
        ot_values_by_dept.append(ot)
        total_ot += ot
        ot_holiday = 0
        if holiday_date_list:
            ot_hrs_holiday = frappe.db.sql("""
                SELECT SUM(total_hours) AS ot_total
                FROM `tabOvertime Request`
                WHERE ot_date IN %(holiday_dates)s
                  AND is_considered_as != 'Compensatory Off'
                  AND department = %(department)s
                  AND ot_reasons != ''
                  AND employee_type = 'Worker'
                  AND workflow_state = 'Approved'
            """, {
                "holiday_dates": tuple(holiday_date_list),
                "department": dept
            }, as_dict=True)
            ot_holiday = ot_hrs_holiday[0]['ot_total'] or 0
        row_holiday.append(ot_holiday)
        total_holiday_ot += ot_holiday
        combined_ot = ot + ot_holiday
        row_combined.append(combined_ot)
        combined_ot_by_dept.append(combined_ot)
        emp_count = frappe.db.sql("""
            SELECT COUNT(DISTINCT employee) AS emp_count
            FROM `tabOvertime Request`
            WHERE MONTH(ot_date) = %s
              AND YEAR(ot_date) = %s
              AND department = %s
              AND ot_reasons != ''
              AND is_considered_as != 'Compensatory Off'
              AND employee_type = 'Worker'
              AND workflow_state = 'Approved'
        """, (month, year, dept), as_dict=True)
        emp_total = emp_count[0]['emp_count'] or 0
        row_employee_count.append(emp_total)
        emp_counts_by_dept.append(emp_total)
        total_employees += emp_total
    row_ot.append(total_ot)
    row_holiday.append(total_holiday_ot)
    row_combined.append(total_ot + total_holiday_ot)
    row_employee_count.append(total_employees)
    for i in range(len(departments)):
        ot_val = ot_values_by_dept[i]
        emp = emp_counts_by_dept[i]
        per_person = round(ot_val / emp, 2) if emp else 0
        row_per_person.append(per_person)
    total_per_person = round(total_ot / total_employees, 2) if total_employees else 0
    row_per_person.append(total_per_person)
    for i in range(len(departments)):
        combined_val = combined_ot_by_dept[i]
        emp = emp_counts_by_dept[i]
        per_person_incl = round(combined_val / emp, 2) if emp else 0
        row_per_person_incl.append(per_person_incl)
    total_combined = total_ot + total_holiday_ot
    total_per_person_incl = round(total_combined / total_employees, 2) if total_employees else 0
    row_per_person_incl.append(total_per_person_incl)
    data.append(row_ot)
    data.append(row_holiday)
    data.append(row_combined)
    data.append(row_employee_count)
    data.append(row_per_person)
    data.append(row_per_person_incl)
    return data
